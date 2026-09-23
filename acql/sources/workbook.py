"""Parser for the ACQL Dashboard.xlsx workbook.

Sheet layouts, established from the real workbook:

  Season    A4 header; A5:A39 players; B..S = weeks 1..18; T season wins;
            U avg; V best; W worst; X rank. Row 3 B..S = games played that week.
  Winnings  B3 "Buy in $40"; A5:A39 players; B..S weekly winnings;
            T total won; U net (won - buy-in).
  Week N    row 4  F..U home teams      row 5  F..U away teams
            row 6  F..U result/winner   row 7  F..U picked-home %
            row 9  headers              rows 10..44 one row per player
            B total wins (incl. big loser)   C regular wins   D rank
            E player   F..U the player's CORRECT picks (blank = missed)
            V,W,X big-loser picks (a 1 marks a correct one)   Y suicide pick
            AA games behind   AB win %

Note: the Week sheet's AA/AB headers are transposed relative to their data -
AA holds games behind and AB holds win percentage. The data is followed, not
the labels.
"""

from __future__ import annotations

import re
from ..config import GAMES_PER_WEEK, WEEKS_IN_SEASON
from ..models import Game, Player, PlayerWeek, SourceFile, Week
from ..names import AliasTable
from ..predictions import team_code

WEEK_SHEET = re.compile(r"^week\s*(\d{1,2})$", re.IGNORECASE)
# The suicide column doubles as an elimination marker once a player is out.
SUICIDE_OUT_MARKERS = {"dead", "out", "x", "eliminated", "knocked out"}

# Week sheet geometry (1-based column numbers).
COL_WINNINGS, COL_TOTAL_WINS, COL_REGULAR_WINS, COL_RANK, COL_PLAYER = 1, 2, 3, 4, 5
COL_FIRST_GAME = 6                       # column F
COL_LAST_GAME = COL_FIRST_GAME + GAMES_PER_WEEK - 1   # column U
COL_LOSER_1, COL_LOSER_2, COL_LOSER_3 = 22, 23, 24    # V, W, X
COL_SUICIDE = 25                                       # Y
COL_GAMES_BEHIND, COL_WIN_PCT = 27, 28                 # AA, AB
ROW_HOME, ROW_AWAY, ROW_RESULT = 4, 5, 6
ROW_FIRST_PLAYER, ROW_LAST_PLAYER = 10, 44

# Season / Winnings geometry.
SEASON_FIRST_ROW, SEASON_LAST_ROW = 5, 39
SEASON_FIRST_WEEK_COL = 2                # column B
SEASON_GAMES_ROW = 3
WINNINGS_TOTAL_COL = 20                  # column T


def _num(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text or text in {"-", "--"}:
        return None
    try:
        return float(text.replace(",", "").replace("$", "").replace("%", ""))
    except ValueError:
        return None


def _int(value: object) -> int | None:
    n = _num(value)
    return int(round(n)) if n is not None else None


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).replace(" ", " ").strip()
    # Formula errors leak through as "#VALUE!" etc; treat them as blank.
    return "" if text.startswith("#") else text


def parse_workbook(
    source: SourceFile,
    aliases: AliasTable,
    players: dict[str, Player] | None = None,
) -> dict:
    """Read the dashboard workbook: weekly picks, season grid and winnings."""
    import openpyxl

    result: dict = {
        "players": players if players is not None else {},
        "weeks": {},
        "buy_in": None,
        "season_title": source.season,
        "current_week": source.week,
        "warnings": [],
    }
    people: dict[str, Player] = result["players"]
    weeks: dict[int, Week] = result["weeks"]

    wb = openpyxl.load_workbook(source.path, data_only=True)
    try:
        def player_for(raw: object) -> Player | None:
            display = _text(raw)
            if not display or display.lower() in {"player", "coach", "rank"}:
                return None
            # A coach is never a number, and never a spelled-out football team.
            # Both land in the player column when a raw pick sheet is pasted a
            # row or two out of place - the sheet's own team rows and its
            # picked-home percentages come down into the player block - and
            # each one becomes a phantom player that then matches nothing in
            # stats.xls. Short codes are deliberately left alone: "tb" is a
            # coach in this pool as well as a football team.
            if _num(display) is not None or (len(display) > 3 and team_code(display)):
                result["warnings"].append(
                    f"{source.path.name}: ignored \"{display}\" in the player "
                    f"column - that is a team or a number, not a coach. Check "
                    f"that the week sheet has coach names in E10:E44 and "
                    f"nothing else."
                )
                return None
            key = aliases.key(display)
            if not key:
                return None
            person = people.get(key)
            if person is None:
                person = Player(key=key, display=display)
                people[key] = person
            person.sources.add(source.path.name)
            return person

        # ---- Week sheets ------------------------------------------------
        for name in wb.sheetnames:
            m = WEEK_SHEET.match(name.strip())
            if not m:
                continue
            number = int(m.group(1))
            if not 1 <= number <= WEEKS_IN_SEASON:
                continue
            ws = wb[name]
            grid = {
                (c.row, c.column): c.value
                for row in ws.iter_rows(
                    min_row=1, max_row=ROW_LAST_PLAYER, max_col=COL_WIN_PCT
                )
                for c in row
            }

            def at(r: int, c: int) -> object:
                return grid.get((r, c))

            week = Week(number=number, source=source.path)
            for i in range(GAMES_PER_WEEK):
                col = COL_FIRST_GAME + i
                home, away = _text(at(ROW_HOME, col)), _text(at(ROW_AWAY, col))
                winner = _text(at(ROW_RESULT, col))
                if home or away or winner:
                    week.games.append(
                        Game(index=i + 1, home=home, away=away, winner=winner)
                    )
            week.game_count = len(week.games)

            for r in range(ROW_FIRST_PLAYER, ROW_LAST_PLAYER + 1):
                person = player_for(at(r, COL_PLAYER))
                if person is None:
                    continue
                line = PlayerWeek(player=person.display, week=number)
                for i in range(GAMES_PER_WEEK):
                    pick = _text(at(r, COL_FIRST_GAME + i))
                    if pick:
                        line.correct_picks[i + 1] = pick

                total = _int(at(r, COL_TOTAL_WINS))
                regular = _int(at(r, COL_REGULAR_WINS))
                line.total_wins = total or 0
                line.regular_wins = regular if regular is not None else line.total_wins

                for slot, col in enumerate((COL_LOSER_1, COL_LOSER_2, COL_LOSER_3)):
                    raw = at(r, col)
                    line.big_loser_raw[slot] = _text(raw)
                    # A correct big-loser pick is overwritten with 1 in place
                    # of the team name.
                    if _num(raw) == 1 and not _text(raw).isalpha():
                        line.big_loser_wins += 1
                    elif _text(raw):
                        line.big_loser_picks.append(_text(raw))

                suicide = _text(at(r, COL_SUICIDE))
                if suicide.casefold() in SUICIDE_OUT_MARKERS:
                    line.suicide_out = True
                else:
                    line.suicide_pick = suicide
                line.rank = _int(at(r, COL_RANK))
                line.games_behind = _num(at(r, COL_GAMES_BEHIND)) or 0.0
                line.win_pct = _num(at(r, COL_WIN_PCT)) or 0.0
                line.winnings = _num(at(r, COL_WINNINGS)) or 0.0

                week.lines[person.key] = line
                if line.suicide_pick:
                    person.suicide_picks[number] = line.suicide_pick
                if line.rank:
                    person.weekly_rank[number] = line.rank

            if week.games or week.lines:
                weeks[number] = week

        # ---- Season sheet -------------------------------------------------
        if "Season" in wb.sheetnames:
            ws = wb["Season"]
            games_per_week: dict[int, int] = {}
            for w in range(1, WEEKS_IN_SEASON + 1):
                n = _int(ws.cell(SEASON_GAMES_ROW, SEASON_FIRST_WEEK_COL + w - 1).value)
                if n:
                    games_per_week[w] = n
            for r in range(SEASON_FIRST_ROW, SEASON_LAST_ROW + 1):
                person = player_for(ws.cell(r, 1).value)
                if person is None:
                    continue
                for w in range(1, WEEKS_IN_SEASON + 1):
                    wins = _int(ws.cell(r, SEASON_FIRST_WEEK_COL + w - 1).value)
                    # stats.xls is authoritative, so only fill gaps here.
                    if wins is not None and w not in person.weekly_wins:
                        person.weekly_wins[w] = wins
            result["games_per_week"] = games_per_week

        # ---- Winnings sheet ------------------------------------------------
        if "Winnings" in wb.sheetnames:
            ws = wb["Winnings"]
            buy_in = None
            for r in range(1, 6):
                for c in range(1, 6):
                    m = re.search(r"buy\s*in\s*\$?\s*([\d.]+)", _text(ws.cell(r, c).value), re.I)
                    if m:
                        buy_in = float(m.group(1))
                        break
                if buy_in:
                    break
            result["buy_in"] = buy_in

            for r in range(SEASON_FIRST_ROW, SEASON_LAST_ROW + 1):
                person = player_for(ws.cell(r, 1).value)
                if person is None:
                    continue
                for w in range(1, WEEKS_IN_SEASON + 1):
                    amount = _num(ws.cell(r, SEASON_FIRST_WEEK_COL + w - 1).value)
                    if amount is not None:
                        person.weekly_winnings[w] = amount

        # ---- Dashboard sheet (title / selected week) ------------------------
        if "Dashboard" in wb.sheetnames:
            ws = wb["Dashboard"]
            title = _text(ws["C1"].value)
            if title:
                result["season_title"] = title
            m = WEEK_SHEET.match(_text(ws["A1"].value))
            if m:
                result["current_week"] = int(m.group(1))
    finally:
        wb.close()

    if not people:
        result["warnings"].append(
            f"{source.path.name}: no players found - the sheet layout may have changed."
        )
    return result