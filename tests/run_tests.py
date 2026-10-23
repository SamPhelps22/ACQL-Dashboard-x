#!/usr/bin/env python3
"""Tests for the ACQL Dashboard's data layer.

Run from the project root:

    python tests/run_tests.py

These exercise the parsers, the merge rules and the write-back guards against
the real spreadsheets in data/. No Qt is imported, so they run headless.
"""

from __future__ import annotations

import shutil
import sys
import contextlib
import tempfile
import time
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

# Fetches in tests keep their raw payloads in a scratch folder, never the real archive.
from acql import rawarchive as _rawarchive  # noqa: E402
_rawarchive.use(Path(tempfile.mkdtemp(prefix="acql-raw-test-")))

from acql.pool import awards
from acql.config import Settings
from acql.models import Player, Season
from acql.names import AliasTable, normalize
from acql.repository import load_season
from acql.sources import classify, discover
from acql.writeback import WorkbookEditor, WriteBackError


def find_sources() -> tuple[Path | None, Path | None]:
    """Locate a stats export and a workbook among the watched folders."""
    stats = workbook = None
    for source in discover(Settings().search_paths()):
        if source.kind == "stats" and stats is None:
            stats = source.path
        elif source.kind == "workbook" and workbook is None and _is_dashboard(source.path):
            workbook = source.path
    return stats, workbook


def _is_dashboard(path: Path) -> bool:
    """Whether a spreadsheet is the pool's own workbook - the one with a
    Dashboard sheet, which the write-back tests are about. Another .xlsx in
    the folder (a standings export saved as one) is read as a workbook too,
    and those tests would fail on it rather than be skipped."""
    try:
        import openpyxl
        book = openpyxl.load_workbook(path, read_only=True)
        try:
            return "Dashboard" in book.sheetnames
        finally:
            book.close()
    except Exception:  # noqa: BLE001 - unreadable is not the workbook
        return False


STATS_PATH, WORKBOOK_PATH = find_sources()
needs_stats = unittest.skipIf(STATS_PATH is None, "no stats.xls available")
needs_workbook = unittest.skipIf(WORKBOOK_PATH is None, "no workbook available")


class TestNames(unittest.TestCase):
    def test_normalize_collapses_incidental_differences(self):
        self.assertEqual(normalize("SaleeNmd "), normalize("saleenmd"))
        self.assertEqual(normalize("MayeDay! MayeDay!"), "mayeday mayeday")
        self.assertEqual(normalize("  Bob   h "), "bob h")
        self.assertEqual(normalize(None), "")

    def test_distinct_players_stay_distinct(self):
        self.assertNotEqual(normalize("Phelpy"), normalize("Phelps Phamily"))

    def test_aliases_map_to_a_canonical_key(self):
        table = AliasTable({"Zach N": "Zachary Nelson"})
        self.assertEqual(table.key("zach n"), normalize("Zachary Nelson"))
        # An unmapped name is returned normalized, not dropped.
        self.assertEqual(table.key("Phelpy"), "phelpy")

    def test_alias_does_not_chain_indefinitely(self):
        table = AliasTable({"a": "b", "b": "c"})
        self.assertEqual(table.key("a"), "b")


class TestDiscovery(unittest.TestCase):
    @needs_stats
    def test_stats_file_declares_its_own_week(self):
        source = classify(STATS_PATH)
        self.assertIsNotNone(source)
        self.assertEqual(source.kind, "stats")
        self.assertIsInstance(source.week, int)
        self.assertEqual(source.season, "ACQL 2026")

    @needs_workbook
    def test_workbook_is_identified(self):
        source = classify(WORKBOOK_PATH)
        self.assertEqual(source.kind, "workbook")

    def test_unrelated_files_are_ignored(self):
        with tempfile.TemporaryDirectory() as tmp:
            junk = Path(tmp) / "notes.txt"
            junk.write_text("not a spreadsheet", encoding="utf-8")
            self.assertIsNone(classify(junk))

    @needs_stats
    def test_identical_content_is_collapsed(self):
        """A re-download under a new name must not be counted twice."""
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            shutil.copy2(STATS_PATH, folder / "stats.xls")
            shutil.copy2(STATS_PATH, folder / "stats (1).xls")
            found = discover([folder])
            self.assertEqual(len(found), 1, "byte-identical copies should collapse")


class TestSeasonLoad(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.season = load_season(Settings())

    def test_players_are_loaded(self):
        self.assertGreater(len(self.season.players), 20)

    def test_the_marker_row_is_not_a_player(self):
        """stats.xls trails a 'Big Losers week N' row in the coach column."""
        for name in (p.display.lower() for p in self.season.players.values()):
            self.assertNotIn("big loser", name)

    def test_standings_are_ordered_by_position(self):
        order = self.season.ordered_players()
        positions = [p.pos for p in order if p.pos]
        self.assertEqual(positions, sorted(positions))

    def test_wins_include_big_loser_points(self):
        """W* covers regular + big-loser wins, so W + L is the full slate:
        the same for every coach, and more than the regular games alone."""
        totals = {p.wins + p.losses for p in self.season.players.values() if p.wins and p.losses}
        if not totals:
            self.skipTest("no standings loaded")
        self.assertEqual(len(totals), 1, totals)
        (total,) = totals
        weeks = len(self.season.final_weeks())
        self.assertGreater(total, weeks * self.season.max_regular_points)
        # The standings export can run ahead of the pick sheets on hand, so the
        # ceiling is the export's own week, not the weeks with sheets.
        covered = max(weeks, self.season.current_week or 0)
        self.assertLessEqual(
            total, covered * (self.season.max_regular_points + self.season.max_big_loser_points))

    def test_provisional_weeks_are_excluded_from_averages(self):
        for week in self.season.provisional_weeks:
            self.assertGreater(week, self.season.current_week)
            for p in self.season.players.values():
                self.assertNotIn(week, p.weekly_wins)

    def test_conflicts_are_reported_not_swallowed(self):
        for c in self.season.conflicts:
            self.assertNotEqual(c.authoritative, c.other)
            self.assertIn(c.player, {p.display for p in self.season.players.values()})

    def test_buy_in_comes_from_the_workbook(self):
        self.assertGreater(self.season.buy_in, 0)

    def test_weeks_carry_slates_and_lines(self):
        played = self.season.played_weeks()
        self.assertTrue(played)
        week = self.season.weeks[played[0]]
        self.assertTrue(week.games)
        self.assertTrue(week.lines)

    def test_correct_pick_counts_are_within_the_field(self):
        for week in self.season.weeks.values():
            for game in week.games:
                self.assertLessEqual(week.correct_count(game.index), len(week.lines))

    def test_big_loser_slots_round_trip(self):
        """The raw slots keep their position; the readable view drops markers."""
        for week in self.season.weeks.values():
            for line in week.lines.values():
                self.assertEqual(len(line.big_loser_raw), 3)
                markers = sum(1 for v in line.big_loser_raw if v == "1")
                if markers:
                    self.assertEqual(markers, line.big_loser_wins, line.player)
                # A team named in the readable view but not in the slots is one
                # that came in: a week built from the pick sheet knows which
                # team a "1" was, where the graded workbook only keeps the "1".
                named = [p for p in line.big_loser_picks if p not in line.big_loser_raw]
                self.assertLessEqual(len(named), markers, line.player)

    def test_suicide_markers_are_not_treated_as_picks(self):
        for p in self.season.players.values():
            for pick in p.suicide_picks.values():
                self.assertNotIn(pick.casefold(), {"dead", "out", "eliminated"})

    def test_eliminated_players_are_not_alive(self):
        for p in self.season.players.values():
            if p.suicide_out_week:
                self.assertFalse(p.suicide_alive, p.display)


class TestDerivedStats(unittest.TestCase):
    def setUp(self):
        self.player = Player(key="x", display="X")
        self.player.weekly_wins = {1: 10, 2: 14, 3: 6}

    def test_season_aggregates(self):
        self.assertEqual(self.player.season_wins, 30)
        self.assertEqual(self.player.best_week, 14)
        self.assertEqual(self.player.worst_week, 6)
        self.assertAlmostEqual(self.player.average_wins, 10.0)

    def test_trend_uses_the_most_recent_weeks(self):
        self.assertAlmostEqual(self.player.trend(2), 10.0)   # weeks 2 and 3
        self.assertAlmostEqual(self.player.trend(3), 10.0)

    def test_rank_movement_is_positive_when_climbing(self):
        self.player.prev_pos, self.player.pos = 10, 4
        self.assertEqual(self.player.rank_movement, 6)
        self.player.prev_pos, self.player.pos = 2, 5
        self.assertEqual(self.player.rank_movement, -3)

    def test_empty_player_does_not_divide_by_zero(self):
        blank = Player(key="y", display="Y")
        self.assertEqual(blank.average_wins, 0.0)
        self.assertEqual(blank.consistency, 0.0)
        self.assertEqual(blank.trend(), 0.0)


class TestStandingsWithoutExport(unittest.TestCase):
    """The standings stats.xls holds, worked out from the weekly scores."""

    @staticmethod
    def _season(scores, *, picks=None, suicide=None):
        from acql.models import Game, PlayerWeek, Week
        season = Season(buy_in=40)
        for key in scores:
            season.players[key] = Player(key=key, display=key.title())
        for n in range(1, len(next(iter(scores.values()))) + 1):
            games = [Game(index=1, home="Chicago", away="Detroit", winner="Chicago",
                          home_points=20, away_points=17)]
            week = Week(number=n, games=games, source=None)
            for key, row in scores.items():
                line = PlayerWeek(player=key, week=n, regular_wins=row[n - 1], total_wins=row[n - 1])
                if suicide:
                    line.suicide_pick = suicide[key][n - 1]
                week.lines[key] = line
            season.weeks[n] = week
        return season

    def test_points_pot_and_ties(self):
        from acql.pool import standings
        s = self._season({"a": [10, 8], "b": [10, 12], "c": [6, 12]})
        self.assertEqual(standings.roll_forward(s, 0), 2)
        a, b, c = (s.players[k] for k in "abc")
        # week 1: a and b split c's 4 short; week 2: b and c split a's 4 short
        self.assertEqual((a.points_won, b.points_won, c.points_won), (2.0, 4.0, 2.0))
        self.assertEqual((a.points_paid, b.points_paid, c.points_paid), (-44.0, -40.0, -44.0))
        self.assertEqual(b.net_points, -36.0)
        self.assertEqual((b.pos, a.pos, c.pos), (1, 2, 3))   # a and c level: a was higher
        self.assertEqual(c.games_behind, 4.0)
        self.assertEqual(s.current_week, 2)

    def test_export_weeks_are_left_alone(self):
        from acql.pool import standings
        s = self._season({"a": [10], "b": [9]})
        self.assertEqual(standings.roll_forward(s, 1), 0)
        self.assertEqual(s.players["a"].wins, 0)

    def test_suicide_pool_is_won_and_starts_again(self):
        from acql.pool import standings
        s = self._season({"a": [1, 1], "b": [1, 1], "c": [1, 1]},
                         suicide={"a": ["Chicago", "Chicago"], "b": ["Detroit", "Detroit"],
                                  "c": ["Detroit", "Chicago"]})
        self.assertTrue(standings.replay_suicide(s, 0))
        self.assertEqual(s.suicide_winners, ["A"])
        self.assertEqual(s.suicide_restart, 2)
        self.assertEqual({k for k, p in s.players.items() if p.suicide_alive}, {"a", "c"})

    def test_ungraded_workbook_row_is_graded(self):
        from pathlib import Path as P
        from acql.pool import standings
        s = self._season({"a": [0]})
        week = s.weeks[1]
        week.source = P("ACQL Dashboard.xlsx")
        line = week.lines["a"]
        line.correct_picks = {1: "detroit"}     # picked the loser, never deleted
        line.regular_wins = line.total_wins = 1
        s.results_year = 2026
        standings.regrade(s)
        self.assertEqual((line.regular_wins, line.correct_picks), (0, {}))


class TestAwards(unittest.TestCase):
    def test_awards_never_raise_on_an_empty_season(self):
        self.assertEqual(awards.compute(Season()), [])

    def test_awards_have_real_winners(self):
        season = load_season(Settings())
        names = {p.display for p in season.players.values()}
        for award in awards.compute(season):
            self.assertIn(award.winner, names, award.title)
            self.assertTrue(award.value)


@needs_workbook
class TestWriteBack(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.book = self.tmp / WORKBOOK_PATH.name
        shutil.copy2(WORKBOOK_PATH, self.book)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def census(self) -> dict:
        import openpyxl
        from openpyxl.worksheet.formula import ArrayFormula

        wb = openpyxl.load_workbook(self.book, data_only=False)
        try:
            return {
                "formulas": sum(
                    1 for ws in wb.worksheets for row in ws.iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.startswith("=")
                ),
                "arrays": sum(
                    1 for ws in wb.worksheets for row in ws.iter_rows() for c in row
                    if isinstance(c.value, ArrayFormula)
                ),
                "tables": sum(len(ws.tables) for ws in wb.worksheets),
                "cf": sum(len(ws.conditional_formatting._cf_rules) for ws in wb.worksheets),
                "validation": sum(len(ws.data_validations.dataValidation) for ws in wb.worksheets),
                "merged": sum(len(ws.merged_cells.ranges) for ws in wb.worksheets),
            }
        finally:
            wb.close()

    def test_writing_results_preserves_every_formula(self):
        before = self.census()
        editor = WorkbookEditor(self.book)
        editor.set_result(1, 1, "Seattle")
        editor.set_result(1, 2, "San Francisco")
        editor.set_matchup(1, 3, "Cincinnati", "Tampa Bay")
        result = editor.apply(backup=False)
        self.assertTrue(result.ok, [r for _, r in result.refused])
        self.assertEqual(before, self.census())

    def test_values_land_where_they_were_addressed(self):
        import openpyxl

        editor = WorkbookEditor(self.book)
        editor.set_result(1, 4, "Detroit")
        editor.apply(backup=False)
        wb = openpyxl.load_workbook(self.book, data_only=False)
        try:
            # Game 4 sits in column I on row 6 of the week sheet.
            self.assertEqual(wb["Week 1"]["I6"].value, "Detroit")
        finally:
            wb.close()

    def test_a_formula_cell_is_refused(self):
        editor = WorkbookEditor(self.book)
        # Dashboard!B7 is a SORTBY formula driving the whole leaderboard.
        editor._queue("Dashboard", 7, 2, "overwritten", "test")
        result = editor.apply(backup=False)
        self.assertFalse(result.ok)
        self.assertEqual(result.applied, [])
        self.assertIn("formula", result.refused[0][1])

    def test_a_refused_edit_abandons_the_whole_change_set(self):
        import openpyxl

        editor = WorkbookEditor(self.book)
        editor.set_result(1, 5, "Tennessee")            # valid
        editor._queue("Dashboard", 7, 2, "nope", "test")  # invalid
        result = editor.apply(backup=False)
        self.assertFalse(result.ok)
        wb = openpyxl.load_workbook(self.book, data_only=False)
        try:
            self.assertNotEqual(wb["Week 1"]["J6"].value, "Tennessee")
        finally:
            wb.close()

    def test_unknown_player_is_rejected(self):
        editor = WorkbookEditor(self.book)
        with self.assertRaises(WriteBackError):
            editor.set_pick(1, "Nobody McGee", 1, "Seattle")

    def test_out_of_range_week_and_game_are_rejected(self):
        editor = WorkbookEditor(self.book)
        with self.assertRaises(WriteBackError):
            editor.set_result(99, 1, "Seattle")
        with self.assertRaises(WriteBackError):
            editor.set_result(1, 99, "Seattle")

    def test_dry_run_changes_nothing(self):
        before = self.book.read_bytes()
        editor = WorkbookEditor(self.book)
        editor.set_result(1, 6, "Baltimore")
        result = editor.apply(dry_run=True)
        self.assertTrue(result.applied)
        self.assertEqual(before, self.book.read_bytes())

    def test_a_backup_is_written_before_saving(self):
        editor = WorkbookEditor(self.book)
        editor.set_result(1, 7, "Pittsburgh")
        result = editor.apply(backup=True)
        self.assertIsNotNone(result.backup)
        self.assertTrue(result.backup.is_file())
        result.backup.unlink(missing_ok=True)

    def test_big_loser_win_markers_survive_editing_a_neighbour(self):
        """A slot reading 1 marks a win; editing slot 3 must not clear 1 and 2."""
        import openpyxl

        editor = WorkbookEditor(self.book)
        editor.set_big_loser_picks(1, "cuetop", ["1", "1", "Denver"])
        self.assertTrue(editor.apply(backup=False).ok)
        wb = openpyxl.load_workbook(self.book, data_only=False)
        try:
            ws = wb["Week 1"]
            row = next(
                r for r in range(10, 45)
                if str(ws.cell(r, 5).value).strip() == "cuetop"
            )
            self.assertEqual(
                [ws.cell(row, c).value for c in (22, 23, 24)], [1, 1, "Denver"]
            )
        finally:
            wb.close()

    def test_a_numeric_entry_is_stored_as_a_number(self):
        """Excel counts 1 but not "1", so a win marker must not become text."""
        import openpyxl

        editor = WorkbookEditor(self.book)
        editor.set_big_loser_picks(1, "Phelpy", ["1", "Arizona", "Miami"])
        editor.apply(backup=False)
        wb = openpyxl.load_workbook(self.book, data_only=False)
        try:
            ws = wb["Week 1"]
            row = next(
                r for r in range(10, 45)
                if str(ws.cell(r, 5).value).strip() == "Phelpy"
            )
            self.assertIsInstance(ws.cell(row, 22).value, int)
            self.assertEqual(ws.cell(row, 22).value, 1)
        finally:
            wb.close()

    def test_picks_and_suicide_write_to_the_right_columns(self):
        import openpyxl

        editor = WorkbookEditor(self.book)
        editor.set_pick(1, "Phelpy", 2, "Los Angeles Rams")
        editor.set_suicide_pick(1, "Phelpy", "Buffalo")
        editor.set_game_count(1, 15)
        self.assertTrue(editor.apply(backup=False).ok)
        wb = openpyxl.load_workbook(self.book, data_only=False)
        try:
            ws = wb["Week 1"]
            row = next(
                r for r in range(10, 45)
                if str(ws.cell(r, 5).value).strip() == "Phelpy"
            )
            self.assertEqual(ws.cell(row, 7).value, "Los Angeles Rams")   # game 2 -> G
            self.assertEqual(ws.cell(row, 25).value, "Buffalo")           # suicide -> Y
            self.assertEqual(wb["Season"]["B3"].value, 15)
        finally:
            wb.close()

    def test_row_lookups_are_cached_within_a_session(self):
        editor = WorkbookEditor(self.book)
        editor.set_pick(1, "Phelpy", 1, "Seattle")
        self.assertTrue(editor._row_cache)
        editor.set_pick(1, "DPSOG", 1, "Seattle")
        # A second player must not have reopened the workbook.
        self.assertEqual(len(editor._row_cache), 1)

    def test_repeated_edits_to_one_cell_collapse(self):
        editor = WorkbookEditor(self.book)
        editor.set_result(1, 8, "Carolina")
        editor.set_result(1, 8, "Chicago")
        self.assertEqual(len(editor), 1)
        self.assertEqual(editor.edits[0].value, "Chicago")


@needs_workbook
class TestHeatmapTooltips(unittest.TestCase):
    """Hovering a heat map cell asks the page for its tooltip with (row label, column label, value)."""

    @classmethod
    def setUpClass(cls):
        import os

        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        try:
            from PySide6.QtWidgets import QApplication
        except ImportError as exc:  # pragma: no cover - Qt absent
            raise unittest.SkipTest(f"PySide6 unavailable: {exc}")
        cls.app = QApplication.instance() or QApplication([])

    def test_hovering_the_favourites_grid_on_insights(self):
        # It raised TypeError on every mouse move (Oct 23): its tooltip took (row, column) positions.
        from types import SimpleNamespace
        from unittest import mock
        from acql.pool import crowd
        from acql.ui.pages.insights import InsightsPage
        from acql.ui.theme import DARK

        page = InsightsPage(DARK)
        grid = (["Under 3", "3 to 6.5"], [6, 7], [[60.0, None], [75.0, 80.0]], [[5, 0], [4, 2]])
        with mock.patch.object(crowd, "observations", return_value=[]), \
                mock.patch.object(crowd, "favourite_grid", return_value=grid):
            page._update_favourites(None)
        text, _anchor = page.fav_chart._hover_at(SimpleNamespace(xdata=1.5, ydata=1.5))
        self.assertIn("Week 7", text)
        self.assertIn("80% of the pool on the favourite (2 games)", text)
        text, _anchor = page.fav_chart._hover_at(SimpleNamespace(xdata=1.5, ydata=0.5))
        self.assertIn("no games", text)


class TestEditorState(unittest.TestCase):
    """The editor widget's change tracking, driven headlessly."""

    @classmethod
    def setUpClass(cls):
        import os

        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        try:
            from PySide6.QtWidgets import QApplication
        except ImportError as exc:  # pragma: no cover - Qt absent
            raise unittest.SkipTest(f"PySide6 unavailable: {exc}")
        cls.app = QApplication.instance() or QApplication([])

    def _editor(self):
        from acql.ui.editor import WeekEditor
        from acql.ui.theme import DARK

        editor = WeekEditor(DARK)
        editor.set_season(load_season(Settings()))
        return editor

    def test_a_freshly_loaded_week_has_no_pending_changes(self):
        editor = self._editor()
        self.assertEqual(editor._collect(), [])

    def test_edits_survive_switching_players_and_back(self):
        editor = self._editor()
        first = editor.player_picker.itemText(0)
        second = editor.player_picker.itemText(1)

        editor.player_picker.setCurrentIndex(0)
        editor.suicide_field.setText("Buffalo")
        self.assertEqual(len(editor._collect()), 1)

        editor.player_picker.setCurrentIndex(1)
        self.assertIn(first, editor._pending_players)
        editor._loser_fields[2].setText("Denver")
        changes = editor._collect()
        self.assertEqual(len(changes), 2)
        self.assertEqual({c[2] for c in changes}, {first, second})

        # Returning shows the held edit rather than the workbook value.
        editor.player_picker.setCurrentIndex(0)
        self.assertEqual(editor.suicide_field.text(), "Buffalo")
        self.assertEqual(len(editor._collect()), 2)

    def test_discarding_clears_every_held_edit(self):
        editor = self._editor()
        editor.player_picker.setCurrentIndex(0)
        editor.suicide_field.setText("Buffalo")
        editor.player_picker.setCurrentIndex(1)
        editor._loser_fields[0].setText("Denver")
        self.assertTrue(editor._collect())
        editor.load_week()
        self.assertEqual(editor._collect(), [])
        self.assertEqual(editor._pending_players, {})

    def test_selecting_the_stored_pick_is_not_a_change(self):
        """Picks stored shouted ("CINCINNATI") must not look edited."""
        editor = self._editor()
        for index in range(min(6, editor.player_picker.count())):
            editor.player_picker.setCurrentIndex(index)
            self.assertEqual(
                [c for c in editor._collect() if c[0] == "pick"], [],
                editor.player_picker.currentText(),
            )

    def test_a_large_change_set_previews_without_flooding_the_dialog(self):
        from acql.ui.editor import PREVIEW_LIMIT, _describe

        editor = self._editor()
        for index in range(min(6, editor.player_picker.count())):
            editor.player_picker.setCurrentIndex(index)
            for row in range(16):
                combo = editor._pick_fields[row]
                if combo.count() > 1:
                    combo.setCurrentIndex(1)
        workbook_editor, changes = editor._build_editor()
        self.assertGreater(len(changes), PREVIEW_LIMIT)

        summary, preview, detail = _describe(workbook_editor.apply(dry_run=True))
        self.assertIn("would be written", summary)
        # The shortened view is capped; the full listing stays complete.
        self.assertLessEqual(len(preview.splitlines()), PREVIEW_LIMIT + 1)
        self.assertIn("more cell(s)", preview)
        self.assertGreater(len(detail.splitlines()), PREVIEW_LIMIT)

    def test_retyping_a_team_keeps_a_pending_result(self):
        editor = self._editor()
        editor._winner_fields[0].setCurrentIndex(0)
        before = [c for c in editor._collect() if c[0] == "result"]
        self.assertTrue(before)
        editor._home_fields[0].setText("Seattle Seahawks")
        editor._refresh_winner_options(0)
        after = editor._collect()
        self.assertTrue([c for c in after if c[0] == "result"])
        self.assertTrue([c for c in after if c[0] == "matchup"])


class _ScratchStore(unittest.TestCase):
    """Each test on its own empty database, so nothing real is read or written."""

    def setUp(self):
        from acql import store
        self.folder = Path(tempfile.mkdtemp(prefix="acql-test-"))
        self.previous = store._OVERRIDE
        store.use(self.folder / "test.db")
        # A new database brings in the old JSON files beside the real one;
        # these tests start from nothing.
        for kind in ("odds", "forecasts", "picks"):
            store.delete(kind)

    def tearDown(self):
        from acql import store
        store.use(self.previous)
        shutil.rmtree(self.folder, ignore_errors=True)


class TestSealedSecrets(unittest.TestCase):
    """Passwords and keys are kept sealed, never as readable text (a made-up value here)."""

    def test_round_trip_and_nothing_readable(self):
        import base64
        from acql.site import mailer
        made_up = "not-a-real-secret-0123"
        sealed = mailer.seal(made_up)
        self.assertEqual(mailer.unseal(sealed), made_up)
        self.assertNotIn(made_up, sealed)
        if sys.platform == "win32":
            self.assertTrue(sealed.startswith("dpapi:"), sealed[:6])
            self.assertNotIn(made_up.encode(), base64.b64decode(sealed[6:]))
        self.assertEqual(mailer.unseal(""), "")


class TestFieldSize(unittest.TestCase):
    """How many other coaches the week planner plays against: one number, and only it moves."""

    CHOICES = [(0.80, 0.85), (0.72, 0.80), (0.66, 0.70), (0.60, 0.75), (0.55, 0.65), (0.52, 0.55),
               (0.70, 0.90), (0.58, 0.60)]

    def _choices(self):
        from acql.pool import weekplan
        return [weekplan.Choice(f"g{i}", f"P{i}", f"O{i}", p, c) for i, (p, c) in enumerate(self.CHOICES)]

    def test_one_default_everywhere(self):
        from acql.model import backtest
        from acql.pool import moneyplan, weekplan
        self.assertEqual(weekplan.OTHER_COACHES, 34)          # 35 coaches in 2026
        self.assertEqual(backtest.OTHER_COACHES, weekplan.OTHER_COACHES)
        self.assertEqual(moneyplan.COACHES, weekplan.OTHER_COACHES)

    def test_field_size_moves_the_week_odds_not_the_game_chances(self):
        from acql.pool import weekplan
        small = weekplan.plan_week(self._choices(), coaches=34, sims=4000)
        large = weekplan.plan_week(self._choices(), coaches=49, sims=4000)
        self.assertEqual((small.coaches, large.coaches), (34, 49))
        self.assertAlmostEqual(small.fair_share, 1 / 35)
        self.assertAlmostEqual(large.fair_share, 1 / 50)
        self.assertGreater(small.plain_win_odds, large.plain_win_odds)   # more coaches, harder to be first
        self.assertAlmostEqual(small.plain_hits, large.plain_hits, places=1)   # same games, same chances
        again = weekplan.plan_week(self._choices(), coaches=49, sims=4000)
        self.assertEqual((again.take, again.win_odds), (large.take, large.win_odds))   # seeded


class TestFirstSeenIsNotTheOpener(_ScratchStore):
    """The first line the app fetched is never passed off as the books' opening line."""

    KICKOFF = 1_790_000_000.0

    def _quote(self, margin):
        from datetime import datetime, timezone
        from acql.feeds import odds
        iso = datetime.fromtimestamp(self.KICKOFF, timezone.utc).isoformat().replace("+00:00", "Z")
        return odds.Quote("phi", "chi", iso, margin, 5)

    def test_first_fetch_is_kept_and_later_fetches_only_move_the_latest(self):
        from acql.feeds import odds
        odds.update(3, [self._quote(3.0)], [("phi", "chi")], now=self.KICKOFF - 5 * 86400)
        odds.update(3, [self._quote(5.5)], [("phi", "chi")], now=self.KICKOFF - 86400)
        kept = odds.stored(3)
        self.assertEqual(kept.first_seen, {"phi@chi": 3.0})
        self.assertEqual(kept.latest, {"phi@chi": 5.5})
        self.assertEqual([(t, n) for _, t, n in odds.movements(3)], [(3.0, 5.5)])

    def test_the_files_opener_stands_and_a_fetched_one_is_marked_first_seen(self):
        from acql.feeds import odds
        from acql.model.predictions import Forecast
        odds.update(3, [self._quote(5.5)], [("phi", "chi"), ("atl", "gb")], now=self.KICKOFF - 86400)
        with_opener = Forecast("Philadelphia", "Chicago", market=4.0, opening=2.5)
        without = Forecast("Philadelphia", "Chicago", market=4.0)
        self.assertEqual(odds.overlay(3, [with_opener], 0)[0].opening, 2.5)
        self.assertFalse(odds.overlay(3, [with_opener], 0)[0].opening_seen)
        seen = odds.overlay(3, [without], 0)[0]
        self.assertEqual((seen.opening, seen.market, seen.opening_seen), (4.0, 5.5, True))
        fetched_only = odds.overlay(3, [], 0)[0]
        self.assertEqual((fetched_only.opening, fetched_only.opening_seen), (5.5, True))

    def test_the_evidence_column_shows_models_and_top_reason(self):
        from acql.model import evidence as ev
        case = ev.Case(1, "Indianapolis @ Minnesota", "Minnesota", "Indianapolis", 0.51, "", False, (
            ev.Signal("Form", ev.FOR, "", "", "form +4.0 a game vs -1.0"),
            ev.Signal("Line move", ev.FOR, "", "", "line moved 1.5 toward Minnesota"),
            ev.Signal("Computer models", ev.AGAINST, "", "", "3 of 9 models agree"),
        ))
        self.assertEqual(case.lines(), ["3 of 9 models agree", "Line moved 1.5 toward Minnesota", "2 for, 1 against"])
        alone = ev.Case(2, "", "X", "Y", 0.5, "", False, (ev.Signal("Rest", ev.AGAINST, "", "", "3 fewer days' rest"),))
        self.assertEqual(alone.lines(), ["Against: 3 fewer days' rest", "0 for, 1 against"])

    def test_the_evidence_card_says_first_seen(self):
        from types import SimpleNamespace
        from acql.model import evidence
        from acql.model.predictions import Forecast
        game = SimpleNamespace(index=1, home="Chicago", away="Philadelphia", line=None)
        for flag, word in ((True, "first seen"), (False, "opened")):
            f = Forecast("Philadelphia", "Chicago", market=5.5, opening=3.0, opening_seen=flag)
            sig = evidence._move(game, "Chicago", True, f)
            self.assertTrue(sig.headline.startswith(word), sig.headline)


class TestOddsKickoff(_ScratchStore):
    """A line fetched once a game is under way is a live line, not one anyone picked against."""

    KICKOFF = 1_790_000_000.0

    def _iso(self, when):
        from datetime import datetime, timezone
        return datetime.fromtimestamp(when, timezone.utc).isoformat().replace("+00:00", "Z")

    def test_a_started_game_is_left_out_of_an_update(self):
        from acql.feeds import odds
        quotes = [
            odds.Quote("atl", "gb", self._iso(self.KICKOFF), -17.5, 5),          # under way
            odds.Quote("phi", "chi", self._iso(self.KICKOFF + 86400), 4.5, 5),   # tomorrow
        ]
        done = odds.update(3, quotes, [("atl", "gb"), ("phi", "chi")], now=self.KICKOFF + 3 * 3600)
        kept = odds.stored(3)
        self.assertEqual(kept.latest, {"phi@chi": 4.5})
        self.assertNotIn("Atlanta @ Green Bay", " ".join(done.missing))

    def test_a_line_kept_after_kickoff_is_set_aside(self):
        from acql.feeds import odds
        before, after = self.KICKOFF - 3600, self.KICKOFF + 3 * 3600
        odds._write({"weeks": {"3": {
            "opening": {"atl@gb": 6.5, "phi@chi": -17.0},
            "latest": {"atl@gb": -17.5, "phi@chi": -17.0},
            "kickoff": {"atl@gb": self._iso(self.KICKOFF), "phi@chi": self._iso(self.KICKOFF)},
            "history": [
                {"at": before, "lines": {"atl@gb": 4.5}},
                {"at": after, "lines": {"atl@gb": -17.5, "phi@chi": -17.0}},
            ],
            "at": after,
        }}})
        kept = odds.stored(3)
        self.assertEqual(kept.latest, {"atl@gb": 4.5})      # the last line before kickoff
        self.assertEqual(kept.opening.get("atl@gb"), 6.5)
        self.assertNotIn("phi@chi", kept.opening)           # only ever seen live
        self.assertEqual(odds.lines_at(3, after + 1), {"atl@gb": 4.5})


class TestDeadline(_ScratchStore):
    """Cards are due Thursday at 6:00 pm - or at the first kickoff, when that is sooner."""

    def test_an_ordinary_week_locks_at_the_deadline(self):
        from acql.feeds import oddslock
        year, week = 2026, 6
        at = time.localtime(oddslock.lock_time(year, week))
        self.assertEqual((at.tm_wday, at.tm_hour, at.tm_min), (3, 18, 0))
        self.assertTrue(oddslock.at_deadline(year, week))
        self.assertTrue(oddslock.locked(year, week, oddslock.deadline(year, week) + 1))
        self.assertFalse(oddslock.locked(year, week, oddslock.deadline(year, week) - 60))

    def test_thanksgiving_locks_at_its_early_kickoff(self):
        from acql.feeds import oddslock
        year, week = 2026, 12
        self.assertFalse(oddslock.at_deadline(year, week))
        self.assertEqual(oddslock.lock_time(year, week), oddslock.kickoff_time(year, week))

    def test_a_kept_deadline_is_used(self):
        from acql.feeds import oddslock
        from acql import store
        store.put("settings", "deadline", {"weekday": 3, "hour": 17, "minute": 30})
        at = time.localtime(oddslock.deadline(2026, 6))
        self.assertEqual((at.tm_wday, at.tm_hour, at.tm_min), (3, 17, 30))


class TestMoneylinesAndTotals(unittest.TestCase):
    """The moneyline answers "who wins"; the total says how widely results run."""

    EVENT = {"home_team": "Kansas City Chiefs", "away_team": "Denver Broncos",
             "commence_time": "2026-10-18T17:00:00Z", "bookmakers": [
                 {"markets": [
                     {"key": "spreads", "outcomes": [{"name": "Kansas City Chiefs", "point": -3.0},
                                                     {"name": "Denver Broncos", "point": 3.0}]},
                     {"key": "h2h", "outcomes": [{"name": "Kansas City Chiefs", "price": -160},
                                                 {"name": "Denver Broncos", "price": 135}]},
                     {"key": "totals", "outcomes": [{"name": "Over", "point": 48.5},
                                                    {"name": "Under", "point": 48.5}]},
                 ]}]}

    def test_a_quote_carries_all_three(self):
        from acql.feeds import odds
        quote = odds.parse([self.EVENT])[0]
        self.assertEqual(quote.home_margin, 3.0)
        self.assertEqual(quote.total, 48.5)
        # -160 is 61.5% and +135 is 42.6%: 104.1% between them, so 59.1% without the book's margin
        self.assertAlmostEqual(quote.home_win, 0.5912, places=3)

    def test_a_spread_and_the_chance_it_wins_go_both_ways(self):
        from acql import analytics
        for chance in (0.55, 0.61, 0.75):
            spread = analytics.spread_for_win_chance(chance)
            self.assertAlmostEqual(analytics.win_chance(spread), chance, delta=0.02)
        self.assertLess(analytics.spread_for_win_chance(0.4), 0)

    def _price(self, market_win=None, total=None, line=None):
        from acql.model import predictions, pricing
        from acql.models import Game
        game = Game(1, "Kansas City", "Denver", line=line, line_favourite="Kansas City" if line else "")
        forecast = predictions.Forecast("Denver", "Kansas City", market=3.0, opening=3.0,
                                        market_win=market_win, total=total)
        ctx = pricing.Context(1, {1: forecast}, 0.0, 1.0, predictions.CONSENSUS_WEIGHT, typical_total=44.0)
        return pricing.price(game, ctx)

    def test_a_straight_game_is_priced_off_the_moneyline(self):
        self.assertEqual(self._price().margin, 3.0)
        self.assertLess(self._price(market_win=0.55).margin, 3.0)      # the moneyline says closer
        self.assertEqual(self._price(market_win=0.55).base, 3.0)       # the spread shown stays the spread

    def test_a_line_game_keeps_the_spread(self):
        self.assertEqual(self._price(market_win=0.55, line=7).margin, 3.0)

    def test_a_moneyline_far_from_its_spread_is_set_aside(self):
        self.assertEqual(self._price(market_win=0.90).margin, 3.0)

    def test_a_high_total_widens_the_results(self):
        from acql import analytics
        self.assertGreater(self._price(total=54.0).sd, analytics.SPREAD_SD)
        self.assertLess(self._price(total=36.0).sd, analytics.SPREAD_SD)
        self.assertEqual(self._price(total=44.0).sd, analytics.SPREAD_SD)


class TestInjuries(_ScratchStore):
    """ESPN's report, kept to the players who matter to this week's game."""

    NOW = 1_791_000_000.0

    def _report(self):
        from datetime import datetime, timezone
        stamp = lambda days: datetime.fromtimestamp(self.NOW - days * 86400, timezone.utc).isoformat()
        def player(name, pos, status, what, days=1):
            return {"status": status, "date": stamp(days), "details": {"type": what},
                    "athlete": {"displayName": name, "position": {"abbreviation": pos}}}
        return {"injuries": [{"displayName": "Baltimore Ravens", "injuries": [
            player("Lamar Jackson", "QB", "Out", "Ankle"),
            player("Third String", "QB", "Out", "Coach's Decision"),
            player("A Receiver", "WR", "Out", "Hamstring"),
            player("A Corner", "CB", "Questionable", "Knee"),
            player("Long Gone", "G", "Injured Reserve", "Knee - ACL", days=40),
            player("Just Hurt", "S", "Injured Reserve", "Calf", days=3),
            player("Fine Now", "TE", "Active", ""),
        ]}]}

    def test_only_the_hurt_and_the_recent_are_kept(self):
        from acql.feeds import injuries
        kept = injuries.parse(self._report(), now=self.NOW)["bal"]
        self.assertEqual([i["name"] for i in kept], ["Lamar Jackson", "A Receiver", "Just Hurt", "A Corner"])

    def test_the_summary_names_the_quarterback_and_counts_the_rest(self):
        from acql.feeds import injuries
        from acql import store
        store.put("injuries", "latest", {"at": self.NOW, "teams": injuries.parse(self._report(), now=self.NOW)})
        self.assertEqual(injuries.summary("bal"),
                         "QB Lamar Jackson out (ankle); 2 more out (WR, S); 1 questionable")
        self.assertEqual(injuries.summary("atl"), "")
        self.assertTrue(injuries.due(None, self.NOW + 7 * 3600))
        self.assertFalse(injuries.due(None, self.NOW + 3600))
        self.assertTrue(injuries.due(3 * 3600, self.NOW + 2 * 3600))      # hourly on deadline day


class TestBigLoser(unittest.TestCase):
    """What a Big Loser pick is worth: the blowout, and how few cards share it."""

    def test_a_pick_everyone_has_gains_less_than_one_few_do(self):
        from acql.pool import bigloser
        popular = bigloser.Option("mia", "Miami", "Buffalo", -10.0, 0.36, 0.70)
        quiet = bigloser.Option("lv", "Las Vegas", "Denver", -5.0, 0.19, 0.16)
        self.assertGreater(popular.chance, quiet.chance)
        self.assertGreater(quiet.gain, popular.gain)
        board = bigloser.Board(6, [popular, quiet])
        self.assertEqual(board.likeliest[0].code, "mia")
        self.assertEqual(board.best_gain[0].code, "lv")

    def test_the_pool_s_share_is_counted_off_the_sheet(self):
        from types import SimpleNamespace
        from acql.pool import bigloser
        sheet = SimpleNamespace(losers={"A": ["Miami", "Arizona", "Las Vegas"],
                                        "B": ["Miami", "Arizona", "Tennessee"]})
        held, counted = bigloser._held(sheet, {"mia": 0.3, "lv": 0.2, "den": 0.1})
        self.assertTrue(counted)
        self.assertEqual(held, {"mia": 1.0, "lv": 0.5, "den": 0.0})

    def test_a_record_is_judged_against_what_its_picks_deserved(self):
        from acql.pool import bigloser
        record = bigloser.Record("a", "A", weeks=4, picks=12, hits=3, expected=3.6, priced=12, held=6.0)
        self.assertAlmostEqual(record.rate, 0.25)
        self.assertAlmostEqual(record.luck, -0.6)
        self.assertAlmostEqual(record.crowd, 0.5)


class TestBigLoserChance(unittest.TestCase):
    """A Big Loser's chance is how often teams like it have been routed, not the margin curve's guess."""

    def test_it_is_history_s_number(self):
        from acql import analytics
        # seven-point underdogs have been routed three times in ten; the curve said a quarter
        self.assertAlmostEqual(analytics.rout_chance(-7), 0.296, places=3)
        self.assertAlmostEqual(analytics.curve_rout_chance(-7), 0.243, places=3)
        self.assertAlmostEqual(analytics.rout_chance(0), 0.126, places=3)
        self.assertAlmostEqual(analytics.rout_chance(10), 0.028, places=3)
        self.assertEqual(analytics.big_loser_chance(-7), analytics.rout_chance(-7))

    def test_a_weaker_team_is_always_likelier(self):
        from acql import analytics
        chances = [analytics.rout_chance(half / 2) for half in range(-60, 61)]
        self.assertTrue(all(worse > better for worse, better in zip(chances, chances[1:])))
        self.assertTrue(all(0.0 < c < 1.0 for c in chances))

    def test_a_wide_open_game_is_likelier_still(self):
        from acql import analytics
        plain = analytics.big_loser_chance(-7)
        wide = analytics.big_loser_chance(-7, analytics.SPREAD_SD * 1.07)
        self.assertTrue(plain < wide < plain + 0.03)
        self.assertEqual(analytics.big_loser_chance(-7, analytics.SPREAD_SD * 0.94), plain)    # the curve only widens

    def test_the_likeliest_come_first(self):
        from acql import analytics
        found = analytics.big_losers([("Buffalo", "Miami", 10.0), ("Dallas", "Houston", -3.0)])
        self.assertEqual([b.team for b in found], ["Miami", "Dallas", "Houston", "Buffalo"])
        self.assertAlmostEqual(found[0].chance, 0.364, places=3)

    def test_in_play_the_score_takes_over(self):
        from acql import analytics
        width = analytics.SPREAD_SD
        at_kickoff = analytics.big_loser_chance_in_play(analytics.curve_rout_chance(-7), -7, width, 1.0)
        self.assertAlmostEqual(at_kickoff, analytics.big_loser_chance(-7))
        self.assertTrue(0.52 < analytics.big_loser_chance_in_play(0.5, -7, width, 0.5) < 0.55)
        self.assertAlmostEqual(analytics.big_loser_chance_in_play(0.5, -7, width, 0.0), 0.5)
        for settled in (0.0, 1.0):
            self.assertEqual(analytics.big_loser_chance_in_play(settled, -7, width, 0.3), settled)

    def test_the_played_out_weeks_carry_the_same_chances(self):
        from types import SimpleNamespace
        import numpy
        from acql.pool import gameday
        draws = numpy.rint(numpy.random.default_rng(3).normal(-7, 13, size=(4000, 2))).astype(int)
        draws[:, 1] = -20                                                   # a game that is over
        states = [SimpleNamespace(state="pre", blowout_home=0.30, blowout_away=0.04),
                  SimpleNamespace(state="final", blowout_home=1.0, blowout_away=0.0)]
        blown = gameday._routs(draws, states, 1)
        home, away = blown[gameday.HOME][:, 0], blown[gameday.AWAY][:, 0]
        self.assertAlmostEqual(home.mean(), 0.30, places=3)
        self.assertAlmostEqual(away.mean(), 0.04, places=3)
        self.assertFalse((home & away).any())
        self.assertLessEqual(draws[home, 0].max(), draws[~home, 0].min())   # the most lopsided draws are the routs
        self.assertTrue(blown[gameday.HOME][:, 1].all())
        self.assertFalse(blown[gameday.AWAY][:, 1].any())

    def test_the_long_replay_keeps_the_curve_for_its_rougher_margins(self):
        from acql import analytics
        from acql.model import backtest
        # a rating's margin is not a betting line: history's number would read it as one
        self.assertEqual(backtest.routed(-8.0), analytics.curve_rout_chance(-8.0, backtest._sd()))
        self.assertLess(backtest.routed(-8.0), analytics.big_loser_chance(-8.0, backtest._sd()))

    def test_it_can_be_taught_again_from_the_files(self):
        from acql import analytics
        from acql.feeds import nflverse
        if not nflverse.available():
            self.skipTest("the nflverse files are not in data/nflverse")
        margins, routed = nflverse.routs()
        self.assertEqual(len(margins), 2 * 6967)
        for kept, taught in zip(analytics.ROUT_WEIGHTS, analytics.rout_fit(margins, routed)):
            self.assertAlmostEqual(kept, taught, places=3)


class TestSignalScoreboard(unittest.TestCase):
    """A signal with a wrong-way record keeps its place and loses its vote."""

    def test_a_distrusted_signal_is_shown_but_not_counted(self):
        from acql.model import evidence
        form = evidence.Signal("Form", evidence.FOR, "Dallas 3-1", "Dallas beat Houston.")
        record = evidence.SignalRecord("Form", games=42, right=15, split=27, split_right=9)
        shown = evidence._discounted(form, {"Form": record})
        self.assertEqual(shown.verdict, evidence.INFO)
        self.assertIn("9 of 27", shown.detail)
        self.assertEqual(evidence._discounted(form, {}).verdict, evidence.FOR)
        self.assertAlmostEqual(record.rate, 15 / 42)

    def test_an_empty_season_scores_nothing(self):
        from acql.model import evidence, power
        board = evidence.scoreboard(Season())
        self.assertEqual([r.games for r in board], [0] * len(evidence.SCORED))
        self.assertEqual(evidence.distrusted(Season()), {})


class TestDeadeyeBigLosers(_ScratchStore):
    """Deadeye's three Big Losers are part of its card: kept with it, and scored from it."""

    def _brief(self):
        from types import SimpleNamespace
        from acql import analytics
        choice = SimpleNamespace(index=1, pick="Kansas City", other="Denver", chance=0.7, matchup="Denver @ Kansas City",
                                 home_margin=6.5, turned=False)
        losers = [analytics.BigLoser("Arizona", "Los Angeles Rams", -10.5, 0.36),
                  analytics.BigLoser("Washington", "San Francisco", -8.0, 0.28),
                  analytics.BigLoser("Los Angeles Chargers", "Kansas City", -7.5, 0.27)]
        return SimpleNamespace(choices=[choice], losers=losers, plan=None, win_odds=None)

    def test_the_kept_card_holds_its_three(self):
        from acql.pool import cards
        snap = cards.Snapshot.of(self._brief())
        self.assertEqual(snap.losers, ["ari", "was", "lac"])
        self.assertAlmostEqual(snap.loser_chance["ari"], 0.36)
        cards._put(6, "model", snap)
        self.assertEqual(cards.model(6).losers, ["ari", "was", "lac"])

    def test_a_card_kept_before_they_were_still_loads(self):
        from acql.pool import cards
        old = cards.Snapshot.load({"at": 1.0, "take": {"1": "Dallas"}, "chance": {"1": 0.6}, "label": {"1": "A @ B"}})
        self.assertEqual(old.losers, [])

    def test_the_ghost_is_scored_on_the_ones_it_kept(self):
        from acql.pool import cards
        from acql.model import ghost
        from acql.models import Week
        snap = cards.Snapshot.of(self._brief())
        self.assertEqual(ghost._losers(Season(), 6, Week(6), snap), ["ari", "was", "lac"])


class TestCardResults(_ScratchStore):
    """Your card marks each pick right or wrong once its game is in, says which
    games are line games, and shows how far each Big Loser lost."""

    def _week(self):
        from acql.models import Game, Week
        return Week(6, games=[
            Game(1, home="Denver", away="Seattle", winner="Seattle", home_points=17, away_points=24),
            Game(2, home="Kansas City", away="Los Angeles Chargers", winner="Los Angeles Chargers",
                 line=7, line_favourite="Kansas City", home_points=27, away_points=21),
            Game(3, home="Green Bay", away="Dallas", winner="Tie", home_points=20, away_points=20),
            Game(4, home="Indianapolis", away="Tennessee", winner="Indianapolis", home_points=38, away_points=13),
            Game(5, home="New England", away="New York Jets"),
        ])

    def test_a_pick_is_graded_against_the_credited_side(self):
        from acql.pool import briefing
        g = {x.index: x for x in self._week().games}
        self.assertEqual(briefing.graded(g[1], "Seattle"), briefing.RIGHT)
        self.assertEqual(briefing.graded(g[1], "Denver"), briefing.WRONG)
        # Kansas City won, but by less than the pool's line: the Chargers are credited.
        self.assertEqual(briefing.graded(g[2], "Kansas City"), briefing.WRONG)
        self.assertEqual(briefing.graded(g[2], "LAC"), briefing.RIGHT)
        self.assertEqual(briefing.graded(g[3], "Dallas"), briefing.TIED)
        self.assertEqual(briefing.graded(g[5], "New England"), "")
        self.assertEqual(briefing.graded(None, "New England"), "")

    def test_a_big_loss_is_read_from_the_score(self):
        from acql.pool import briefing
        games = self._week().games
        self.assertEqual(briefing.big_loss(games, "Tennessee"), 25)
        self.assertEqual(briefing.big_loss(games, "Indianapolis"), -25)
        self.assertIsNone(briefing.big_loss(games, "New York Jets"))      # not played
        self.assertIsNone(briefing.big_loss(games, "Buffalo"))            # not on the slate

    def test_the_glance_shows_the_rule_and_the_results(self):
        import os
        from types import SimpleNamespace
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        try:
            from PySide6.QtWidgets import QApplication
        except ImportError as exc:  # pragma: no cover - Qt absent
            self.skipTest(f"PySide6 unavailable: {exc}")
        # Kept on the class: a QApplication nobody holds is freed under its widgets.
        type(self).app = QApplication.instance() or QApplication([])
        from acql import analytics
        from acql.ui.pages.thisweek import BriefingPage
        from acql.ui.theme import PALETTES

        def choice(index, matchup, pick, chance, line=""):
            return SimpleNamespace(index=index, matchup=matchup, pick=pick, chance=chance,
                                   lined=bool(line), line=line, why="")
        brief = SimpleNamespace(week=6, expected=2.9, for_money=False, choices=[
            choice(1, "Seattle @ Denver", "Seattle", 0.59),
            choice(2, "Los Angeles Chargers @ Kansas City", "Kansas City", 0.55, "Kansas City by 7"),
            choice(5, "New York Jets @ New England", "New England", 0.87),
        ], losers=[analytics.BigLoser("Tennessee", "Indianapolis", -15.1489, 0.49),
                   analytics.BigLoser("New York Jets", "New England", -12.5818, 0.43)])
        season = Season()
        season.weeks[6] = self._week()
        page = BriefingPage(next(iter(PALETTES.values())))
        page.season = season
        page._update_glance(season, brief, [])
        chips = [page.glance_grid.itemAt(i).widget().text() for i in range(page.glance_grid.count())]
        self.assertIn("Right", chips[0])
        self.assertIn("Straight up", chips[0])
        self.assertIn("LINE GAME", chips[1])
        self.assertIn("Kansas City by 7", chips[1])
        self.assertIn("Wrong", chips[1])
        self.assertNotIn("Right", chips[2])
        self.assertNotIn("Wrong", chips[2])
        self.assertIn("1 of 2 right so far", page.glance_card.title_label.text())
        self.assertIn("1 line game", page.glance_note.text())
        page._update_losers(season, brief)
        losers = [page.losers_grid.itemAt(i).widget().text() for i in range(page.losers_grid.count())]
        self.assertIn("15.1-point underdogs", losers[0])
        self.assertNotIn("15.1489", losers[0])
        self.assertIn("Hit", losers[0])
        self.assertIn("lost by 25", losers[0])
        self.assertNotIn("lost by", losers[1])
        page.deleteLater()


class TestCardEmail(_ScratchStore):
    """Deadeye's card goes to whoever runs the dashboard - never to the pool's list."""

    def _settings(self):
        from acql.site import mailer
        settings = mailer.MailSettings(sender="me@example.com",
                                       recipients=["pool-a@example.com", "pool-b@example.com"])
        settings.sealed_password = "plain:eA=="
        settings.save()
        return settings

    def test_it_goes_to_the_sender_and_not_the_recap_list(self):
        from acql.site import cardmail
        settings = self._settings()
        self.assertEqual(cardmail.recipients(settings), ["me@example.com"])
        card = cardmail.CardMail(6, "Deadeye week 6: 2 picks", "text", "<p>html</p>")
        message = cardmail.message(card, settings)
        self.assertEqual(message["To"], "me@example.com")
        self.assertNotIn("pool-a", str(message))

    def test_a_kept_address_is_used(self):
        from acql.site import cardmail
        from acql import store
        settings = self._settings()
        store.put("settings", "cardmail", {"to": ["phone@example.com", "not an address"]})
        self.assertEqual(cardmail.recipients(settings), ["phone@example.com"])

    def test_it_says_what_changed_since_the_last_one(self):
        from acql.site import cardmail
        before = {"picks": {"1": ["Dallas", 0.51], "2": ["Chicago", 0.60], "3": ["Miami", 0.57]}}
        now = {"1": ["Tampa Bay", 0.55], "2": ["Chicago", 0.68], "3": ["Miami", 0.58], "4": ["Denver", 0.61]}
        games = {"1": "Tampa Bay @ Dallas", "2": "Chicago @ Green Bay", "3": "Cincinnati @ Miami", "4": "Denver @ LAC"}
        changes = cardmail._changes(now, games, before)
        self.assertEqual(len(changes), 3)
        self.assertTrue(changes[0].startswith("Changed: Tampa Bay @ Dallas - now Tampa Bay"))
        self.assertTrue(changes[1].startswith("Moved: Chicago"))
        self.assertTrue(changes[2].startswith("New: Denver @ LAC"))
        self.assertEqual(cardmail._changes(now, games, {}), [])        # a first email has nothing to compare

    def test_sending_keeps_the_card_for_next_time(self):
        from acql.site import cardmail
        self._settings()
        sent = []

        class Server:
            def __enter__(self): return self
            def __exit__(self, *_): return False
            def login(self, user, password): sent.append(("login", user))
            def send_message(self, message): sent.append(("to", message["To"]))

        card = cardmail.CardMail(6, "Deadeye week 6", "text", "<p>html</p>", picks={"1": ["Seattle", 0.55]})
        self.assertEqual(cardmail.send(card, smtp_factory=Server), ["me@example.com"])
        self.assertEqual(sent, [("login", "me@example.com"), ("to", "me@example.com")])
        self.assertEqual(cardmail._last(6)["picks"], {"1": ["Seattle", 0.55]})


class TestNewestNumberWins(unittest.TestCase):
    """A spread typed before the latest line gives way to it; one typed since wins."""

    def _price(self, typed_at):
        from acql.feeds import lines
        from acql.model import predictions, pricing
        from acql.models import Game
        game = Game(1, "Kansas City", "Denver")
        forecast = predictions.Forecast("Denver", "Kansas City", market=6.5, opening=5.5)
        ctx = pricing.Context(1, {1: forecast}, 0.0, 1.0, predictions.CONSENSUS_WEIGHT, market_at=1000.0)
        return pricing.price(game, ctx, lines.Spread("home", 2.5, typed_at), typed=True)

    def test_an_older_typed_spread_gives_way_to_the_market(self):
        stale = self._price(500.0)
        self.assertEqual(stale.base, 6.5)
        self.assertTrue(stale.overruled)

    def test_a_spread_typed_since_the_latest_line_wins(self):
        fresh = self._price(2000.0)
        self.assertEqual(fresh.base, 2.5)
        self.assertFalse(fresh.overruled)

    def test_a_number_being_typed_now_wins(self):
        self.assertEqual(self._price(None).base, 2.5)


class TestPatterns(_ScratchStore):
    """Blind spots, the line move and chaos, on a small made-up pool."""

    MATCHUPS = [("Atlanta", "Green Bay"), ("Philadelphia", "Chicago"), ("Dallas", "New York Giants"),
                ("Denver", "Kansas City"), ("Seattle", "Washington"), ("Detroit", "Carolina"),
                ("Miami", "Buffalo"), ("Houston", "Tennessee")]

    def _season(self, home_wins=8):
        from acql.models import Game, PlayerWeek, Week
        games = [Game(i, home=h, away=a, winner=h if i <= home_wins else a)
                 for i, (a, h) in enumerate(self.MATCHUPS, 1)]
        week = Week(1, games=games)
        season = Season()
        # Four coaches take every home team; "contra" takes every road team.
        for key in ("a", "b", "c", "d", "contra"):
            season.players[key] = Player(key=key, display=key.title())
            side = (lambda g: g.away) if key == "contra" else (lambda g: g.home)
            week.lines[key] = PlayerWeek(key, 1, correct_picks={g.index: side(g) for g in games})
        season.weeks[1] = week
        return season

    def _market(self, home_margin=7.0, opening=None):
        from acql.pool import patterns
        return {1: {i: patterns.Market(opening, home_margin) for i in range(1, 9)}}

    def test_a_graded_line_puts_the_misses_back(self):
        from acql.pool import patterns
        season = self._season(home_wins=6)
        line = season.weeks[1].lines["a"]
        line.correct_picks = {i: t for i, t in line.correct_picks.items() if i <= 6}   # graded: misses wiped
        card = patterns.week_cards(season, 1)["a"]
        self.assertEqual(len(card), 8)
        self.assertEqual(sum(p.right for p in card), 6)
        self.assertEqual({p.side for p in card if not p.right}, {"home"})   # the side the road team beat

    def test_blind_spots_are_measured_against_the_same_games(self):
        from acql.pool import patterns
        season = self._season(home_wins=6)
        spots = patterns.blind_spots(season, market=self._market())
        # The rest of the pool (three on the home side, one on the road) went
        # 3 in 4 on the six home wins and 1 in 4 on the two road wins.
        self.assertAlmostEqual(spots["a"].overall.expected, 6 * 0.75 + 2 * 0.25)
        self.assertAlmostEqual(spots["a"].overall.gap, 1.0)
        self.assertEqual(spots["a"].worst, [])
        contra = spots["contra"]
        self.assertEqual(contra.overall.record, "2-6")
        self.assertAlmostEqual(contra.overall.field_rate, 6 / 8)

    def test_the_line_move_credits_the_side_it_moved_to(self):
        from acql.pool import patterns
        season = self._season()
        moves = {m.key: m for m in patterns.line_moves(season, market=self._market(7.0, opening=3.0))}
        self.assertEqual((moves["a"].with_move, moves["a"].moved), (8, 8))
        self.assertGreater(moves["a"].gained, 0)
        self.assertEqual(moves["contra"].with_move, 0)
        self.assertLess(moves["contra"].gained, 0)
        # No opening line, no move to read.
        self.assertEqual(patterns.line_moves(season, market=self._market(7.0)), [])

    def test_chaos_rates_a_chalk_week_low_and_a_week_of_upsets_high(self):
        from acql.pool import patterns
        calm = patterns.week_chaos(self._season(home_wins=8), 1, self._market())
        wild = patterns.week_chaos(self._season(home_wins=0), 1, self._market())
        self.assertLess(calm.rating, 2.0)
        self.assertEqual((calm.label, calm.upsets), ("Chalk", 0))
        self.assertGreater(wild.rating, 9.0)
        self.assertEqual((wild.label, wild.upsets), ("Chaos", 8))
        self.assertIn(" over ", wild.biggest_text)


class TestClosingLines(_ScratchStore):
    """Each game's line is kept as it moves, and a pick is judged by where its line ended."""

    KICKOFF = 1_790_000_000.0

    def _score(self, spread, state="pre"):
        from acql.feeds import results
        return results.Score("phi", "chi", state=state, kickoff=self.KICKOFF, home_spread=spread)

    def test_a_line_is_kept_as_stretches_and_never_after_kickoff(self):
        from acql.model import closing
        early, later, last = self.KICKOFF - 3 * 86400, self.KICKOFF - 86400, self.KICKOFF - 3600
        self.assertEqual(closing.capture(2026, 6, [self._score(3.0)], early), 1)
        closing.capture(2026, 6, [self._score(3.0)], later)         # the same number: one stretch, longer
        closing.capture(2026, 6, [self._score(4.5)], last)
        self.assertEqual(closing.capture(2026, 6, [self._score(9.5, "in")], self.KICKOFF + 3600), 0)
        self.assertEqual(closing.capture(2026, 6, [self._score(9.5)], self.KICKOFF + 3600), 0)
        kept = closing.lines(2026, 6)[frozenset(("phi", "chi"))]
        self.assertEqual(kept["home"], "chi")
        self.assertEqual(kept["runs"], [(early, later, 3.0), (last, last, 4.5)])

    def test_keeping_the_scores_keeps_the_lines(self):
        from acql.model import closing
        from acql.feeds import results
        results.save(2026, 6, [self._score(3.0)], at=self.KICKOFF - 86400)
        self.assertEqual(closing.lines(2026, 6)[frozenset(("phi", "chi"))]["runs"][-1][2], 3.0)

    def test_a_pick_is_judged_by_where_its_line_ended(self):
        from acql.model import closing
        toward = closing.PickValue(6, "Philadelphia @ Chicago", "Chicago", 3.0, 4.5, False)
        away = closing.PickValue(6, "Dallas @ Denver", "Dallas", -2.0, -3.0, False)
        still = closing.PickValue(6, "Miami @ Buffalo", "Buffalo", 6.5, 6.5, False)
        two_books = closing.PickValue(6, "Seattle @ Arizona", "Seattle", 1.0, 1.5, True)
        self.assertEqual([p.verdict for p in (toward, away, still, two_books)], [1, -1, 0, 0])
        value = closing.Value([toward, away, still, two_books])
        self.assertEqual((value.toward, value.away, value.weeks), (1, 1, [6]))
        self.assertAlmostEqual(value.points, (1.5 - 1.0 + 0.0 + 0.5) / 4)


class TestWeather(_ScratchStore):
    """Weather is said only when there is some, and only where there is no roof."""

    KICKOFF = 1_790_010_000.0

    def _opener(self, calls):
        import io
        import json
        hours = [time.strftime("%Y-%m-%dT%H:00", time.gmtime(self.KICKOFF + h * 3600)) for h in range(5)]

        def open_url(request):
            calls.append(request.full_url)
            if "geocoding" in request.full_url:
                body = {"results": [{"latitude": 51.5, "longitude": -0.12, "admin1": "England"}]}
            else:
                body = {"hourly": {"time": hours, "temperature_2m": [50, 52, 51, 51, 40],
                                   "precipitation_probability": [20, 70, 40, 10, 0],
                                   "wind_speed_10m": [12, 18, 16, 9, 5],
                                   "wind_gusts_10m": [20, 31, 25, 15, 10]}}
            return io.BytesIO(json.dumps(body).encode("utf-8"))

        return open_url

    def _slate(self, now):
        from acql.feeds import results
        results.save(2026, 6, [
            results.Score("den", "nyj", kickoff=self.KICKOFF, city="London, England", indoor=False),
            results.Score("ari", "lar", kickoff=self.KICKOFF, city="Inglewood, CA", indoor=True),
        ], at=now)

    def test_only_the_outdoor_game_is_forecast(self):
        from acql.feeds import weather
        now, calls = self.KICKOFF - 2 * 86400, []
        self._slate(now)
        self.assertTrue(weather.due(2026, 6, now))
        self.assertEqual(weather.refresh(2026, 6, opener=self._opener(calls), now=now), 1)
        found = weather.game(2026, 6, "Denver", "New York Jets")
        self.assertEqual((found.wind, found.gust, found.wet, found.temperature), (18.0, 31.0, 70, 51.0))
        self.assertTrue(found.notable)
        self.assertIn("18 mph wind, gusts to 31, 70% chance of rain", found.words())
        self.assertEqual(weather.game(2026, 6, "New York Jets", "Denver"), found)     # either way round
        self.assertIsNone(weather.game(2026, 6, "Arizona", "Los Angeles Rams"))
        self.assertTrue(weather.indoors(2026, 6, "Arizona", "Los Angeles Rams"))
        self.assertFalse(weather.due(2026, 6, now + 60))
        self.assertTrue(weather.due(2026, 6, now + weather.EVERY))
        # The city is looked up once and kept.
        weather.refresh(2026, 6, opener=self._opener(calls), now=now + weather.EVERY)
        self.assertEqual(sum("geocoding" in url for url in calls), 1)

    def test_it_is_shown_under_the_pick_and_never_counted(self):
        from acql.model import evidence
        from acql.feeds import weather
        from acql.models import Game
        now = self.KICKOFF - 2 * 86400
        self._slate(now)
        weather.refresh(2026, 6, opener=self._opener([]), now=now)
        signal = evidence._weather(2026, 6, Game(1, home="New York Jets", away="Denver"))
        self.assertEqual((signal.name, signal.verdict), ("Weather", evidence.INFO))
        self.assertTrue(signal.headline.endswith("in London"))
        self.assertIsNone(evidence._weather(2026, 6, Game(2, home="Los Angeles Rams", away="Arizona")))

    def test_a_calm_day_is_not_worth_a_line(self):
        from acql.feeds import weather
        self.assertFalse(weather.Forecast(60.0, 5.0, 10.0, 10).notable)
        self.assertTrue(weather.Forecast(20.0, 5.0, 10.0, 10).notable)       # cold alone is
        self.assertIn("chance of snow", weather.Forecast(28.0, 5.0, 10.0, 80).words())


class TestTeamStrength(unittest.TestCase):
    """The dashboard's own ratings: from final scores, and only the ones before the week."""

    def _games(self, season=2025):
        # Dallas beats Miami by 14 every week, home and away.
        return [(season, w, "dal", "mia", 14.0) if w % 2 else (season, w, "mia", "dal", -14.0)
                for w in range(1, 9)]

    def test_a_week_is_rated_on_what_came_before_it(self):
        from acql.model import backtest, teamstrength
        rated, starts, final, last = teamstrength._walk(self._games())
        self.assertEqual((starts[(2025, 1)].teams, starts[(2025, 1)].games), ({}, 0))
        self.assertEqual(starts[(2025, 3)].games, 2)
        # The first game was priced level but for the home edge - before it was played.
        self.assertAlmostEqual(rated[(2025, 1, frozenset(("dal", "mia")))], backtest.HOME_EDGE_START)
        self.assertGreater(final.teams["dal"], starts[(2025, 3)].teams["dal"])
        self.assertGreater(starts[(2025, 3)].teams["dal"], 0)
        self.assertAlmostEqual(final.teams["dal"], -final.teams["mia"])
        self.assertEqual((final.games, last), (8, 2025))

    def test_a_new_season_starts_part_of_the_way_back_to_average(self):
        from acql.model import backtest, teamstrength
        _rated, starts, _final, last = teamstrength._walk(self._games() + self._games(2026))
        ended = teamstrength._walk(self._games())[2].teams["dal"]
        self.assertAlmostEqual(starts[(2026, 1)].teams["dal"], ended * (1 - backtest.PRESEASON_REGRESSION))
        self.assertEqual(last, 2026)

    def test_a_matchup_is_priced_from_the_home_side(self):
        from acql.model import teamstrength
        strength = teamstrength.Strength({"dal": 3.0, "mia": -3.0}, 2.0, 8)
        self.assertEqual(strength.margin("Dallas", "Miami"), 8.0)
        self.assertEqual(strength.margin("Miami", "Dallas"), -4.0)
        self.assertIsNone(strength.margin("Dallas", "Seattle"))        # a team never seen


class TestLockedPage(unittest.TestCase):
    """The card on the website is scrambled before it leaves, and only then."""

    @staticmethod
    @contextlib.contextmanager
    def _scratch_store():
        from acql import store
        previous = store._OVERRIDE
        with tempfile.TemporaryDirectory() as scratch:
            store.use(Path(scratch) / "test.db")
            try:
                yield
            finally:
                store.use(previous)

    def test_the_cipher_is_the_standards(self):
        from acql.site import lockedpage as L
        h = bytes.fromhex
        # NIST's own AES-GCM test cases 2, 13, 14 and 16
        self.assertEqual(L.encrypt(bytes(16), bytes(12), bytes(16)).hex(),
                         "0388dace60b6a392f328c2b971b2fe78ab6e47d42cec13bdf53a67b21257bddf")
        self.assertEqual(L.encrypt(bytes(32), bytes(12), b"").hex(), "530f8afbc74536b9a963b4f1c4cb738b")
        self.assertEqual(L.encrypt(bytes(32), bytes(12), bytes(16)).hex(),
                         "cea7403d4d606b6e074ec5d3baf39d18d0d1c8a799996bf0265b98b5d48ab919")
        key = h("feffe9928665731c6d6a8f9467308308feffe9928665731c6d6a8f9467308308")
        plain = h("d9313225f88406e5a55909c5aff5269a86a7a9531534f7da2e4c303d8a318a721c3c0c95956809532fcf0e2449a6b525"
                  "b16aedf5aa0de657ba637b39")
        out = L.encrypt(key, h("cafebabefacedbaddecaf888"), plain, h("feedfacedeadbeeffeedfacedeadbeefabaddad2"))
        self.assertEqual(out.hex(), "522dc1f099567d07f47f37a32a84427d643a8cdcbfe5c0c97598a2bd2555d1aa8cb08e48590dbb3d"
                                    "a7b08b1056828838c5f61e6393ba7a0abcc9f66276fc6ece0f4e1768cddf8853bb2d551b")

    def test_the_page_says_nothing_of_what_it_holds(self):
        from acql.site import lockedpage as L
        from acql.version import MODEL_NAME
        page = L.page(f"<h2>{MODEL_NAME}: Week 6</h2><p>Kansas City 71%</p>", "plum anchor violin seventeen", rounds=1000)
        for told in (MODEL_NAME, "Week 6", "Kansas City", "71%", "plum anchor"):
            self.assertNotIn(told, page)
        self.assertIn("noindex", page)
        one = L.lock(b"the same card", "plum anchor violin seventeen", rounds=1000)
        two = L.lock(b"the same card", "plum anchor violin seventeen", rounds=1000)
        self.assertNotEqual(one["data"], two["data"])               # a fresh salt and iv every time

    def test_a_short_passphrase_is_refused(self):
        from acql.site import lockedpage as L
        self.assertTrue(L.weak("football"))
        self.assertTrue(L.weak("aaaaaaaaaaaaaaaaaaaa"))
        self.assertEqual(L.weak("plum anchor violin seventeen"), "")

    def test_nothing_goes_up_without_a_passphrase_and_nothing_twice(self):
        from acql.site import lockedpage as L, webpublish
        sent = []
        real_put = webpublish.put_file
        webpublish.put_file = lambda web, name, body, **k: sent.append((name, body)) or "https://x.github.io/pool/"
        try:
            with self._scratch_store():
                web = webpublish.WebSettings(user="x", repo="pool")
                web.set_token("t")
                web.save()
                self.assertIsNone(L.put("<p>zebra</p>"))                         # no passphrase: nothing
                self.assertEqual(sent, [])
                settings = L.LockedSettings.load()
                settings.set_passphrase("plum anchor violin seventeen")
                settings.save()
                real_rounds, L.ROUNDS = L.ROUNDS, 1000
                try:
                    self.assertEqual(L.put("<p>zebra</p>"), "https://x.github.io/pool/" + L.FILE)
                    self.assertIsNone(L.put("<p>zebra</p>"))                     # the same card: not again
                    self.assertIsNotNone(L.put("<p>zebra, changed</p>"))
                finally:
                    L.ROUNDS = real_rounds
                self.assertEqual([name for name, _ in sent], [L.FILE, L.FILE])
                self.assertNotIn(b"zebra", sent[0][1])                          # scrambled before it left
                settings = L.LockedSettings.load()
                settings.on = False
                settings.save()
                self.assertIsNone(L.put("<p>another</p>"))                      # switched off: nothing
        finally:
            webpublish.put_file = real_put


class TestModelLab(unittest.TestCase):
    """A challenger takes Deadeye's place only on enough games, and clearly better."""

    def test_the_training_grounds_answer_is_read_or_is_nothing(self):
        import json
        from acql.model import challengers
        clear = {"games": 900, "edge": 0.004, "low": 0.001, "high": 0.007}
        body = {"call": "A has a case.", "in_use": "Deadeye at 25%", "best": "A", "made": 100.0,
                "archive": {"games": 900, "seasons": [2002, 2026], "held_back": [2024, 2025, 2026], "challengers": {
                    "A": {"all": clear, "held_back": {"games": 300, "edge": 0.002}},
                    "B": {"all": {**clear, "low": -0.001}},
                    "C": {"all": {"games": 900, "edge": -0.004, "low": -0.007, "high": -0.001}},
                    "D": {"all": clear, "held_back": {"games": 300, "edge": -0.002}}}},
                "this_season": {"games": 14, "challengers": {"A": {"games": 14, "edge": 0.01, "weeks": 1, "ahead": 1}}}}
        with tempfile.TemporaryDirectory() as scratch:
            path = Path(scratch) / "verdict.json"
            self.assertIsNone(challengers.read(path))                       # no file: no answer
            path.write_text("{not json", encoding="utf-8")
            self.assertIsNone(challengers.read(path))                       # a damaged one: no answer
            path.write_text(json.dumps(body), encoding="utf-8")
            found = challengers.read(path)
        self.assertEqual([row[0] for row in found.rows], ["A", "B", "C", "D"])
        self.assertEqual((found.cases, found.behind), ({0}, {2}))           # D led, but not on the held-back seasons
        self.assertEqual(found.rows[0][4:6], ["+0.0100 on 14", "1 of 1"])
        self.assertEqual(found.rows[1][3:6], ["-", "-", "-"])
        self.assertTrue(found.stale(now=100.0 + challengers.STALE_S + 1))
        self.assertFalse(found.stale(now=101.0))

    def _entry(self, shared, right, deadeye_right, logloss, deadeye_logloss):
        from acql.model import modellab
        return modellab.Entry("Challenger", "", games=shared, right=right, shared=shared,
                              shared_right=right, deadeye_right=deadeye_right,
                              shared_logloss=shared * logloss, deadeye_logloss=shared * deadeye_logloss)

    def test_a_good_month_is_not_enough(self):
        from acql.model import modellab
        hot = self._entry(60, 45, 38, 0.60, 0.66)
        self.assertAlmostEqual(hot.edge, 0.06)
        self.assertEqual(hot.verdict, f"ahead, 60 of {modellab.ENOUGH_GAMES} games")
        self.assertIsNone(modellab.Lab([hot], 60, (1, 2, 3, 4)).challenger)

    def test_enough_games_and_clearly_better_earns_the_switch(self):
        from acql.model import modellab
        n = modellab.ENOUGH_GAMES
        earned = self._entry(n, 100, 96, 0.62, 0.64)
        close = self._entry(n, 100, 96, 0.635, 0.64)        # better, but not clearly
        fewer = self._entry(n, 95, 96, 0.62, 0.64)          # sharper chances, fewer winners
        self.assertEqual(earned.verdict, "has earned the switch")
        self.assertNotEqual(close.verdict, "has earned the switch")
        self.assertNotEqual(fewer.verdict, "has earned the switch")
        self.assertIs(modellab.Lab([close, earned], n, (1,)).challenger, earned)

    def test_the_one_in_service_is_not_its_own_challenger(self):
        from acql.model import modellab
        service = modellab.Entry("Deadeye", "", games=73, right=49, incumbent=True)
        self.assertEqual((service.verdict, service.edge), ("in service", None))
        self.assertAlmostEqual(service.rate, 49 / 73)
        self.assertEqual(modellab._blend(3.0, 7.0, 0.25), 4.0)
        self.assertIsNone(modellab._blend(None, 7.0, 0.25))


class TestPlaces(unittest.TestCase):
    """A page remembered under its old name still opens where it lives now."""

    def test_every_old_name_leads_to_a_place(self):
        from acql import places
        sidebar = [places.SEASON, places.THIS_WEEK, places.GAME_DAY, places.RESULTS, places.SIDE_POOLS,
                   places.PLAYER, places.INSIGHTS, places.MODEL, places.DATA]
        self.assertEqual(len(set(sidebar)), len(sidebar))
        self.assertEqual(set(places.FORMERLY.values()), set(sidebar))
        self.assertFalse(set(places.FORMERLY) & set(sidebar))          # no old name is a new one


class TestTheme(unittest.TestCase):
    """The dashboard's colours read on their own surfaces; the pool's recap keeps the league's."""

    def test_text_reads_on_every_surface(self):
        from acql.ui import theme
        for p in theme.PALETTES.values():
            self.assertGreaterEqual(theme.contrast(p.ink, p.plane), 12)
            self.assertGreaterEqual(theme.contrast(p.ink_muted, p.surface), 4.0)
            self.assertGreaterEqual(theme.contrast(p.brand, p.surface), 4.0)
            written = "#05070a" if theme._bright(p.accent) else "#ffffff"
            self.assertGreaterEqual(theme.contrast(written, p.accent), 4.5)    # a primary button's label

    def test_the_recap_keeps_the_league_colours(self):
        from acql.ui import theme
        for name, p in theme.PALETTES.items():
            league = theme.league(p)
            self.assertEqual(league.name, name)
            self.assertNotEqual(league.accent, p.accent)
            self.assertIn(p.accent, theme.stylesheet(p))        # the stylesheet builds, in its own colours
            self.assertNotIn(league.accent, theme.stylesheet(p))


class TestHitsAndMisses(unittest.TestCase):
    """Deadeye's record cut by confidence and situation, and each miss by its kind."""

    def _game(self, **over):
        from acql.model import autopsy
        made = dict(week=1, index=1, matchup="Miami @ Buffalo", pick="Buffalo", chance=0.70, right=True,
                    margin=6.0, line=6.5, line_pick="Buffalo", line_right=True, actual=10.0, home_pick=True)
        made.update(over)
        return autopsy.Graded(**made)

    def test_a_row_counts_what_its_own_chances_called_for(self):
        from acql.model import autopsy
        games = [self._game(index=i, right=i <= 6, line_right=i <= 6) for i in range(1, 11)]
        row = autopsy.tally("Every pick", games)
        self.assertEqual((row.picks, row.right, row.line_right, row.line_games), (10, 6, 6, 10))
        self.assertAlmostEqual(row.expected, 7.0)
        self.assertAlmostEqual(row.expected_rate, 0.70)
        self.assertEqual(row.verdict, "too few to say")          # ten picks prove nothing

    def test_confidence_is_only_called_wrong_on_enough_picks_and_a_real_gap(self):
        from acql.model import autopsy
        cold = autopsy.tally("Lock", [self._game(index=i, chance=0.80, right=i <= 12) for i in range(1, 31)])
        fair = autopsy.tally("Lock", [self._game(index=i, chance=0.80, right=i <= 23) for i in range(1, 31)])
        hot = autopsy.tally("Lean", [self._game(index=i, chance=0.56, right=i <= 27) for i in range(1, 31)])
        self.assertEqual(cold.verdict, "surer than it should be")
        self.assertEqual(fair.verdict, "as good as it said")
        self.assertEqual(hot.verdict, "better than it said")

    def test_picks_are_sorted_by_how_sure_it_was(self):
        from acql.model import autopsy
        games = [self._game(index=1, chance=0.80), self._game(index=2, chance=0.66),
                 self._game(index=3, chance=0.58), self._game(index=4, chance=0.52)]
        self.assertEqual([(t.name, t.picks) for t in autopsy.tiers(games)],
                         [("Lock", 1), ("Strong", 1), ("Lean", 1), ("Coin flip", 1)])

    def test_a_miss_is_named_for_what_kind_it_was(self):
        from acql.model import autopsy
        lost = dict(right=False, actual=-4.0)
        toss_up = self._game(index=1, chance=0.52, **lost)
        shared = self._game(index=2, line_right=False, **lost)
        against = self._game(index=3, line_pick="Miami", line_right=True, **lost)
        no_cover = self._game(index=4, right=False, actual=3.0, pool_line=7.0, pool_favourite="Buffalo",
                              line_right=False)
        routed = self._game(index=5, line_right=False, right=False, actual=-20.0)
        found = {m.game.index: m for m in autopsy.misses([toss_up, shared, against, no_cover, routed, self._game(index=6)])}
        self.assertEqual({i: m.kind for i, m in found.items()}, {
            1: autopsy.TOSS_UP, 2: autopsy.SHARED, 3: autopsy.AGAINST, 4: autopsy.COVER, 5: autopsy.SHARED})
        self.assertEqual([i for i, m in found.items() if m.never_close], [5])
        self.assertNotIn(6, found)                               # a hit is not a miss

    def test_margin_error_and_runs(self):
        from acql.model import autopsy
        games = [self._game(index=1, margin=6.0, line=7.0, actual=10.0),            # 4 off, the line 3
                 self._game(index=2, margin=3.0, line=2.0, actual=-3.0, right=False),   # 6 off, the line 5
                 self._game(index=3, actual=None)]                                    # no score kept
        error = autopsy.margin_error(games)
        self.assertEqual((error.games, error.deadeye, error.line), (2, 5.0, 4.0))
        run = [self._game(index=i, right=r) for i, r in enumerate([True, True, False, True, True, True, False, False], 1)]
        self.assertEqual(autopsy.streaks(run), (3, 2))
        self.assertEqual(autopsy.margin_text(-3.0), "lost by 3")
        self.assertEqual(autopsy.margin_text(6.5), "by 6.5")

    def test_a_cut_with_nothing_to_compare_is_left_out(self):
        from acql.model import autopsy
        games = [self._game(index=i, home_pick=i % 2 == 0) for i in range(1, 9)]
        titles = [title for title, _rows in autopsy.situations(games)]
        self.assertIn("Home or road", titles)
        self.assertNotIn("Favourite or underdog", titles)        # every pick was the favourite
        self.assertNotIn("Time of season", titles)               # and all in week 1


class TestRadar(unittest.TestCase):
    """How a game could go, read off the same margin the pick is made from."""

    def test_the_spread_of_results_adds_up(self):
        from acql import analytics
        from acql.model import radar
        win, upset, close, rout = radar._spread_of_results(7.0, analytics.SPREAD_SD)
        level = analytics.margin_distribution(7.0, analytics.SPREAD_SD)
        tie = sum(p for d, p in level if d == 0)
        self.assertAlmostEqual(win + upset + tie, 1.0)
        self.assertAlmostEqual(win, analytics.win_chance(7.0))
        self.assertGreater(win, 0.70)
        # A bigger favourite is upset less, wins big more, and is in fewer close ones.
        w2, u2, c2, r2 = radar._spread_of_results(13.0, analytics.SPREAD_SD)
        self.assertTrue(u2 < upset and r2 > rout and c2 < close)

    def test_hard_to_call_runs_from_a_lock_to_a_coin_flip_in_a_gale(self):
        from acql.model import radar
        self.assertEqual(radar.hard_to_call(1.0, 0.0, 0.0, False), 0)
        self.assertEqual(radar.hard_to_call(0.5, 2 * radar.SPLIT, 3.0, True), 100)
        calm = radar.hard_to_call(0.60, 0.0, None, False)
        self.assertEqual(calm, 48)                                # 0.6 x 0.8 of the way to a coin flip
        self.assertGreater(radar.hard_to_call(0.60, 6.0, 1.5, True), calm)
        self.assertGreater(radar.hard_to_call(0.52, 0.0, 0.0, False), radar.hard_to_call(0.80, 0.0, 0.0, False))

    def test_the_slate_names_its_extremes(self):
        from acql.model import radar

        def shape(index, **over):
            made = dict(index=index, matchup=f"Game {index}", pick="A", chance=0.6, favourite="A", underdog="B",
                        margin=3.0, line=3.0, win=0.6, upset=0.39, close=0.27, rout=0.27, gap=0.0, split=0.0,
                        hard=48.0)
            made.update(over)
            return radar.Shape(**made)

        slate = radar.Slate(6, [
            shape(1, hard=61.0, flags=("a toss-up at 52%",)),
            shape(2, line=1.5, upset=0.46),                      # likelier, but no real underdog
            shape(3, line=3.5, upset=0.42),
            shape(4, rout=0.47, gap=0.4),
            shape(5, gap=-1.0, hard=30.0, flags=("weather",)),
        ])
        self.assertEqual((slate.hardest.index, slate.upset.index, slate.rout.index, slate.edge.index), (1, 3, 4, 5))
        self.assertEqual(slate.doubts, ["Game 1: a toss-up at 52%", "Game 5: weather"])
        quiet = radar.Slate(6, [shape(1)])
        self.assertIsNone(quiet.edge)                            # within half a point of the line everywhere
        self.assertEqual(radar.Slate(7).doubts, [])


class TestBlindSpotsAndFear(_ScratchStore):
    """Where the record sits under its own chances - and whether this week has games like it."""

    def _game(self, **over):
        from acql.model import autopsy
        made = dict(week=2, index=1, matchup="Miami @ Buffalo", pick="Buffalo", chance=0.65, right=True,
                    margin=5.0, line=5.5, line_pick="Buffalo", line_right=True, actual=7.0, home_pick=True)
        made.update(over)
        return autopsy.Graded(**made)

    def _season(self):
        # Fine outside the division (16 of 20), poor inside it (4 of 14).
        outside = [self._game(index=i, right=i <= 16) for i in range(1, 21)]
        inside = [self._game(index=100 + i, divisional=True, right=i <= 4, actual=7.0 if i <= 4 else -9.0)
                  for i in range(1, 15)]
        return outside + inside

    def test_a_game_falls_in_its_situations(self):
        from acql.model import autopsy
        names = autopsy.matching(self._game(divisional=True, night=True, pool_line=7.0))
        for name in ("The betting favourite", "The line's side", "The home team", "Line 3.5 to 6.5",
                     "Pool-line games", "Division games", "Night games", "Weeks 1-6"):
            self.assertIn(name, names)
        self.assertNotIn("The road team", names)

    def test_the_worst_situation_is_found_and_says_how_sure_it_is(self):
        from acql.model import autopsy
        spots = autopsy.blind_spots(self._season())
        self.assertEqual(spots[0].name, "Division games")
        self.assertEqual((spots[0].right, spots[0].picks), (4, 14))
        self.assertEqual(spots[0].verdict, "too few to say")              # fourteen picks prove nothing
        self.assertAlmostEqual(spots[0].error, (4 * 2.0 + 10 * 14.0) / 14)
        self.assertNotIn("Outside the division", [s.name for s in spots])
        self.assertEqual(autopsy.blind_spots([self._game(index=i) for i in range(1, 40)]), [])

    def test_the_fear_is_a_blind_spot_this_week_has_games_in(self):
        from acql.model import autopsy
        history = self._season()
        coming = {"Dallas @ Philadelphia": ("Division games", "Night games"), "Denver @ Miami": ("Day games",)}
        found = autopsy.fear(history, coming)
        self.assertEqual((found.row.name, found.games), ("Division games", ("Dallas @ Philadelphia",)))
        self.assertIsNone(autopsy.fear(history, {"Denver @ Miami": ("Day games",)}))

    def test_a_finished_week_is_read_back_as_it_was_written(self):
        from dataclasses import asdict
        from acql.model import autopsy
        from acql import store
        game = self._game(pool_line=7.0, pool_favourite="Buffalo", night=True)
        self.assertIsNone(autopsy.kept(2026, 2))
        store.put(autopsy.KIND, autopsy._key(2026, 2), {"at": 1.0, "version": "x", "games": [asdict(game)]})
        self.assertEqual(autopsy.kept(2026, 2), [game])


class TestScoutingAndPulse(unittest.TestCase):
    """A projected score from the margin and the total, and the week's pulse from its games."""

    def test_the_projected_score_splits_the_total_by_the_margin(self):
        from acql.model import scouting
        self.assertEqual(scouting.projected_score(7.0, 45.0), (26, 19))
        self.assertEqual(scouting.projected_score(-3.0, 41.0), (19, 22))
        self.assertIsNone(scouting.projected_score(7.0, None))             # no over/under, no score
        home, away = scouting.projected_score(1.0, 44.0)                   # 22.5-21.5 must not round to a tie
        self.assertGreater(home, away)
        self.assertEqual(scouting._ordinal(1), "1st")
        self.assertEqual(scouting._ordinal(12), "12th")
        self.assertEqual(scouting._ordinal(23), "23rd")

    def test_the_pulse_is_the_average_game(self):
        from acql.model import radar

        def shape(index, hard, chance=0.6):
            return radar.Shape(index=index, matchup=f"Game {index}", pick="A", chance=chance, favourite="A",
                               underdog="B", margin=3.0, line=3.0, win=0.6, upset=0.39, close=0.27, rout=0.27,
                               gap=0.0, split=0.0, hard=hard)

        slate = radar.Slate(6, [shape(1, 60.0, 0.52), shape(2, 40.0), shape(3, 20.0, 0.8)])
        self.assertEqual((slate.pulse, slate.mood, slate.toss_ups), (40.0, "Ordinary", 1))
        self.assertEqual(radar.Slate(6, [shape(1, 70.0), shape(2, 60.0)]).mood, "Wide open")
        self.assertEqual((radar.Slate(7).pulse, radar.Slate(7).mood), (0.0, "Chalk"))


class TestContenders(unittest.TestCase):
    """The Analyst and the Machine: taught on one set of games, asked about another."""

    def _games(self, n=900, seed=3):
        import numpy as np
        draw = np.random.default_rng(seed)
        X = draw.normal(size=(n, 3))
        # one straight-line input, one that only matters past a threshold, one that is noise
        y = 4.0 * X[:, 0] + np.where(X[:, 1] > 0.3, 6.0, -2.0) + draw.normal(scale=0.5, size=n)
        return X, y

    def test_the_regression_finds_the_straight_line(self):
        import numpy as np
        from acql.model import contenders
        X, y = self._games()
        mean, spread, weights, constant = contenders.ridge(X, y)
        guess = ((X - mean) / spread) @ weights + constant
        self.assertAlmostEqual(float(weights[0] / spread[0]), 4.0, delta=0.3)
        self.assertLess(abs(float(weights[2])), 0.3)                       # the noise gets next to nothing
        self.assertLess(float(np.abs(guess - y).mean()), float(np.abs(y - y.mean()).mean()))

    def test_the_trees_find_the_step_the_line_cannot(self):
        import numpy as np
        from acql.model import contenders
        X, y = self._games()
        Xt, yt = self._games(seed=4)                                        # games neither was taught on
        mean, spread, weights, constant = contenders.ridge(X, y)
        start, trees = contenders.boost(X, y)
        line = ((Xt - mean) / spread) @ weights + constant
        boosted = start + contenders.LEARNING_RATE * sum(contenders._leaves(t, Xt) for t in trees)
        self.assertLess(float(np.abs(boosted - yt).mean()), float(np.abs(line - yt).mean()))
        one = start + contenders.LEARNING_RATE * sum(contenders._leaf(t, Xt[0]) for t in trees)
        self.assertAlmostEqual(one, float(boosted[0]))                      # one game at a time agrees
        again = contenders.boost(X, y)
        self.assertEqual(again[1], trees)                                   # the same Machine every time

    def test_every_game_is_described_by_the_same_inputs(self):
        from acql.model import contenders
        inputs = contenders._inputs(3.5, 2.0, -1.5, [7.0, -3.0, 10.0], [], 6)
        self.assertEqual(len(inputs), len(contenders.INPUTS))
        self.assertEqual(inputs[:3], [3.5, 2.0, -1.5])
        self.assertAlmostEqual(inputs[3], 14.0 / 3)                         # its last three: all it has played
        self.assertEqual((inputs[4], inputs[5], inputs[6], inputs[7], inputs[8]), (0.0, 10.0, 0.0, 3.0, 0.0))
        untaught = contenders.Taught(2026)
        self.assertEqual(untaught.margins(inputs), (None, None))
        self.assertEqual(untaught.recent(3), {})


class TestGrades(unittest.TestCase):
    """The report card's letters: one measurement, one yardstick, and no mark without enough to go on."""

    def _report(self, weeks=5, model=10, mine=8, pool=7.5, hits=49, said=47.0):
        from acql.model import reportcard
        report = reportcard.Report(coach="Sam")
        report.weeks = [reportcard.WeekLine(week=w, games=16, model=model, mine=mine, pool_average=pool, best=12,
                                            model_place=4, model_tied=0, my_place=12, entrants=35, rebuilt=False)
                        for w in range(1, weeks + 1)]
        report.expected, report.graded_picks, report.graded_hits, report.luck = said, 80, hits, 4.3
        return report

    def _hits(self, right=49, line_right=48, picks=73, deadeye=10.0, line=10.0):
        from acql.model import autopsy
        return autopsy.Report([], autopsy.Row("Every pick", picks, right, 43.0, 17.0, picks, line_right), [], [], [],
                              autopsy.MarginError(picks, deadeye, line), (0, 0))

    def test_letters_and_steps(self):
        from acql.model import grades
        self.assertEqual((grades.letter(4.3), grades.letter(3.24), grades.letter(0.2)), ("A+", "B+", "F"))
        self.assertEqual(grades.by_steps(0.0, grades.BY_LINE), "B")              # level with the market
        self.assertEqual(grades.by_steps(3.0, grades.BY_LINE), "A")
        self.assertEqual(grades.by_steps(-2.0, grades.BY_LINE), "C+")
        self.assertEqual([grades.tone(g) for g in ("A-", "B+", "C", "F", "")], ["good", "", "warning", "bad", "muted"])

    def test_the_card_takes_its_place_in_the_field(self):
        from acql.model import grades
        from acql.models import PlayerWeek, Week
        season = Season()
        totals = {"ann": 12, "bob": 10, "cy": 10, "di": 7}
        for key in totals:
            season.players[key] = Player(key=key, display=key.title())
        week = Week(1)
        for key, wins in totals.items():
            week.lines[key] = PlayerWeek(key, 1, regular_wins=wins)
        season.weeks[1] = week
        stand = grades.standing(season, [1], 11, "di")
        self.assertEqual((stand.rank, stand.of, stand.tied, stand.leader), (2, 5, 0, ("Ann", 12)))
        self.assertEqual(stand.mine, (5, 7))                                   # the card is ahead of her too
        self.assertEqual([row[:2] for row in stand.table],
                         [(1, "Ann"), (2, "Deadeye's card"), (3, "Bob"), (3, "Cy"), (5, "Di")])
        self.assertAlmostEqual(stand.share, 0.75)
        self.assertIsNone(grades.standing(Season(), [1], 11, None))

    def test_each_subject_is_marked_against_its_own_yardstick(self):
        from acql.model import grades
        stand = grades.Standing(rank=2, of=36, total=50, leader=("Line-Pony", 51))
        card = grades.mark(self._report(), self._hits(), None, stand, pool_losers=None)
        marks = {s.name: s.grade for s in card.subjects}
        self.assertEqual(marks["Picking games"], "A+")              # ahead of 97% of the pool
        self.assertEqual(marks["Against the line"], "B+")           # one game up in 73
        self.assertEqual(marks["Honest chances"], "A")              # two games off, where luck moves it four
        self.assertEqual(marks["Margins"], "B")                     # no closer than the line, no further
        self.assertEqual(marks["Against you"], "A")                 # two games a week ahead
        self.assertEqual(marks["Winning the week"], "A")            # averaged 4th of 36
        self.assertEqual((marks["Big Losers"], marks["Line movement"]), ("", ""))   # nothing to mark them on
        self.assertEqual(len(card.marked), 6)
        self.assertEqual(card.overall, grades.letter(sum(s.points for s in card.marked) / 6))
        self.assertTrue(card.provisional)
        self.assertEqual((card.best.name, card.worst.name), ("Picking games", "Margins"))

    def test_a_poor_season_gets_poor_marks_and_a_long_one_is_no_longer_provisional(self):
        from acql.model import grades
        stand = grades.Standing(rank=30, of=36, total=60, leader=("Line-Pony", 90))
        report = self._report(weeks=9, model=7, mine=8, hits=40, said=50.0)
        card = grades.mark(report, self._hits(right=44, line_right=48, deadeye=11.0, line=10.0), None, stand)
        marks = {s.name: s.grade for s in card.subjects}
        self.assertEqual(marks["Picking games"], "C-")
        self.assertEqual(marks["Against the line"], "D")            # four games behind the market in 73
        self.assertEqual(marks["Honest chances"], "D")              # ten games short, where luck moves it four
        self.assertEqual(marks["Margins"], "D")                     # a point further off than the line
        self.assertEqual(marks["Against you"], "C")                 # a game a week behind
        self.assertFalse(card.provisional)
        self.assertEqual(self._hits(right=40, line_right=48).overall.right, 40)
        self.assertEqual(grades.mark(report, self._hits(right=40, line_right=48), None, stand).subjects[1].grade, "F")
        self.assertIn(card.overall, ("C-", "C", "D"))

    def test_nothing_is_marked_without_a_graded_week(self):
        from acql.model import grades, reportcard
        card = grades.mark(reportcard.Report(coach="Sam"))
        self.assertEqual((card.marked, card.overall, card.points), ([], "", None))
        self.assertEqual(len(card.subjects), 8)
        self.assertEqual(card.by_week(), {})


class TestModelLabSideBySide(unittest.TestCase):
    """Two recipes are told apart only by the games they split - and only on games both priced."""

    def _entry(self, name, plays, incumbent=False):
        from acql.model import modellab
        entry = modellab.Entry(name, "", incumbent=incumbent)
        for (week, index), (chance, hit, winner, lined) in plays.items():
            loss = -__import__("math").log(chance if hit == (chance > 0.5) else 1 - chance)
            entry.log[(week, index)] = modellab.Play(week, index, chance, hit, loss, winner, lined)
            entry.games += 1
            entry.right += hit
            entry.logloss += loss
        return entry

    def _pair(self):
        # ten games both priced: both right in 4, both wrong in 1, A alone in 4, B alone in 1 - and one only A priced
        a, b = {}, {}
        for i in range(1, 11):
            a[(1, i)] = (0.70, i <= 8, i <= 8, i <= 3)
            b[(1, i)] = (0.60, i <= 4 or i == 9, i <= 4 or i == 9, i <= 3)
        a[(2, 1)] = (0.70, True, None, False)
        return self._entry("A", a), self._entry("B", b)

    def test_the_split_games_are_counted_both_ways(self):
        from acql.model import modellab
        a, b = self._pair()
        duel = modellab.pair(a, b)
        self.assertEqual((duel.games, duel.both, duel.neither, duel.first, duel.second), (10, 4, 1, 4, 1))
        self.assertEqual((duel.split, duel.net), (5, 3))
        self.assertGreater(duel.edge, 0)                         # A's chances were the better ones too
        back = modellab.pair(b, a)
        self.assertEqual((back.first, back.second, back.net), (1, 4, -3))
        self.assertAlmostEqual(back.edge, -duel.edge)
        self.assertEqual(modellab.pair(a, modellab.Entry("C", "")).games, 0)

    def test_luck_is_only_asked_about_once_there_are_enough_split_games(self):
        from acql.model import modellab
        self.assertIsNone(modellab.Pair(games=76, first=3, second=2).luck)
        self.assertEqual(modellab.Pair(games=76, first=3, second=2).reading, "too few split games (5)")
        even = modellab.Pair(games=76, first=6, second=6)
        self.assertEqual(even.luck, 1.0)
        self.assertTrue(even.reading.startswith("within luck"))
        lopsided = modellab.Pair(games=150, first=9, second=0)
        self.assertAlmostEqual(lopsided.luck, 2 / 2 ** 9)        # nine split games, all one way
        self.assertTrue(lopsided.reading.startswith("a real gap"))
        self.assertAlmostEqual(modellab.Pair(first=8, second=13).luck, 0.3833, places=3)
        self.assertEqual(modellab.Pair(games=10, both=7, neither=3).reading, "never split")

    def test_winners_and_covering_the_line_are_kept_apart(self):
        a, _b = self._pair()
        self.assertEqual(a.winners, (8, 10))                     # the game with no score kept says nothing
        self.assertEqual(a.lined, (3, 3))
        self.assertEqual(a.on([(1, 1), (1, 9), (2, 1), (9, 9)]), (2, 3))

    def test_recipes_are_compared_on_the_games_all_of_them_priced(self):
        from acql.model import modellab
        a, b = self._pair()
        lab = modellab.Lab([a, b], 11, (1, 2), decided=12)
        self.assertEqual(len(lab.common), 10)
        self.assertEqual((a.on(lab.common), b.on(lab.common)), ((8, 10), (5, 10)))
        self.assertIsNone(lab.entry("nobody"))

    def test_chances_are_banded_by_how_sure_it_was_of_its_own_side(self):
        from acql.model import modellab
        entry = self._entry("A", {(1, 1): (0.62, True, True, False), (1, 2): (0.38, False, False, False),
                                  (1, 3): (0.64, False, False, False), (1, 4): (0.90, True, True, False)})
        bands = {band.label: band for band in modellab.calibration(entry)}
        mid = bands["60%-65%"]
        self.assertEqual(mid.picks, 3)                           # 38% on the home side is 62% on its own
        self.assertAlmostEqual(mid.said, (0.62 + 0.62 + 0.64) / 3)
        self.assertAlmostEqual(mid.came, 1 / 3)
        self.assertEqual((bands["80% and up"].picks, bands["80% and up"].came), (1, 1.0))
        self.assertEqual(bands["50%-55%"].picks, 0)

    def test_the_front_runner_has_to_have_priced_nearly_the_same_games(self):
        from acql.model import modellab
        service = modellab.Entry("Deadeye", "", games=76, incumbent=True)
        half = modellab.Entry("Half", "", games=44, shared=44, shared_logloss=44 * 0.62, deadeye_logloss=44 * 0.65)
        full = modellab.Entry("Full", "", games=76, shared=76, shared_logloss=76 * 0.642, deadeye_logloss=76 * 0.645)
        behind = modellab.Entry("Behind", "", games=76, shared=76, shared_logloss=76 * 0.66, deadeye_logloss=76 * 0.645)
        self.assertEqual(modellab.Lab([service, half, full, behind], 76, (1,)).front_runner.name, "Full")
        self.assertEqual(modellab.Lab([service, half, behind], 76, (1,)).front_runner, None)
        self.assertIsNone(modellab.Lab([half], 44, (1,)).front_runner)


class TestShowcase(unittest.TestCase):
    """The numbers behind the model's charts, on a small made-up season."""

    def _game(self, n, **over):
        from acql.model import autopsy
        made = dict(week=1 + (n - 1) // 10, index=n, matchup=f"Away {n} @ Home {n}", pick=f"Home {n}",
                    chance=0.62, right=True, margin=4.0, line=3.5, line_pick=f"Home {n}", line_right=True,
                    actual=7.0, home_pick=True)
        made.update(over)
        return autopsy.Graded(**made)

    def _season(self):
        # thirty picks: the first twenty mostly right, the last ten mostly wrong
        games = []
        for n in range(1, 31):
            right = n % 4 != 0 if n <= 20 else n % 3 == 0
            games.append(self._game(n, right=right, line_right=n % 5 != 0,
                                    chance=0.52 + (n % 4) * 0.06, actual=7.0 if right else -6.0))
        return games

    def test_the_trajectory_counts_half_a_game_either_way_on_shared_games(self):
        from acql.model import showcase
        games = self._season()
        base = {(g.week, g.index): g.index % 2 == 0 for g in games}
        found = showcase.trajectory(games, base)
        deadeye, line, ratings = found.lines
        self.assertEqual([l.name for l in found.lines], ["Deadeye", showcase.LINE, showcase.BASELINE])
        self.assertEqual((deadeye.picks, line.right, ratings.right), (30, 24, 15))
        self.assertEqual(deadeye.above[:4], [0.5, 1.0, 1.5, 1.0])              # three in, then a miss
        self.assertAlmostEqual(deadeye.above[-1], deadeye.right - 15)          # hits less half the picks
        self.assertEqual(found.weeks, [(1, 1), (11, 2), (21, 3)])
        self.assertEqual(found.labels[0], "Week 1 - Away 1 @ Home 1")
        # A game one of them had no side in is left out for all three.
        fewer = showcase.trajectory(games, {k: v for k, v in base.items() if k != (1, 1)})
        self.assertEqual(fewer.picks, 29)
        self.assertEqual(len(showcase.trajectory(games).lines), 2)             # no baseline given

    def test_bands_hold_what_it_said_and_what_happened(self):
        from acql.model import showcase
        games = self._season()
        bands = showcase.buckets(games)
        self.assertEqual([b.label for b in bands], [label for label, _lo, _hi in showcase.BANDS])
        self.assertEqual(sum(b.picks for b in bands), 30)
        first = bands[0]                                                       # the 52% picks: every fourth game
        self.assertEqual(first.picks, 7)
        self.assertAlmostEqual(first.said, 0.52)
        self.assertAlmostEqual(first.luck, (7 * 0.52 * 0.48) ** 0.5 / 7)
        self.assertTrue(all(b.picks for b in showcase.calibration(games)))     # the empty ones are left out
        self.assertEqual(showcase.buckets([])[0].rate, 0.0)

    def test_the_matrix_leaves_thin_rows_unshaded(self):
        from acql.model import showcase
        games = self._season() + [self._game(40 + i, home_pick=False, week=3) for i in range(3)]
        rows = {k.name: k for k in showcase.matrix(games)}
        self.assertEqual(rows["The home team"].picks, 30)
        self.assertIsNotNone(rows["The home team"].rate)
        self.assertAlmostEqual(rows["The home team"].line, 24 / 30)
        self.assertEqual((rows["The road team"].picks, rows["The road team"].rate), (3, None))   # too few to shade
        self.assertNotIn("Night games", rows)                                  # no kickoff times: no such row

    def test_margins_runs_and_sizes(self):
        from acql.model import showcase
        games = self._season()
        points = showcase.scope(games + [self._game(99, actual=None)])
        self.assertEqual(len(points), 30)                                      # no score, no point
        self.assertAlmostEqual(points[0].error, 3.0)
        self.assertEqual(points[0].short, "Home 1, wk 1")
        self.assertAlmostEqual(showcase.mean_error(points), sum(p.error for p in points) / 30)
        sizes = dict((label, (mine, line)) for label, mine, line in showcase.miss_sizes(games))
        self.assertEqual(sum(mine for mine, _line in sizes.values()), 30)
        self.assertEqual(sizes["0-3"][0], sum(1 for g in games if g.right))    # a hit here is three points off
        roll = showcase.rolling(games, window=20)
        self.assertEqual((roll.picks[0], roll.picks[-1], len(roll.rate)), (20, 30, 11))
        self.assertAlmostEqual(roll.rate[0], 15 / 20)
        self.assertLess(roll.rate[-1], roll.rate[0])                           # it went cold
        self.assertEqual(showcase.rolling(games[:5]).picks, [])                # too few for a window

    def test_the_dashboard_starts_on_its_owner(self):
        from acql.pool import me
        season = Season()
        for key, name in (("lp", "Line-Pony"), ("ph", "  phelpy "), ("zn", "Zach N")):
            season.players[key] = Player(key=key, display=name)
        self.assertEqual(me.NAME, "Phelpy")
        self.assertEqual((me.key(season), me.display(season)), ("ph", "  phelpy "))
        self.assertEqual((me.key(Season()), me.display(Season()), me.key(None)), ("", "", ""))


class TestKnownBeforeKickoff(_ScratchStore):
    """A week is written down while it is open and frozen at the deadline - never with hindsight."""

    LOCK = 1_790_000_000.0

    def _build(self, said):
        return lambda season, week: {"1": {"pick": said}}

    def test_an_open_week_is_kept_up_to_date_then_frozen(self):
        from acql.model import known
        day = 86400
        self.assertEqual(known.keep(Season(), 2026, 6, self.LOCK, self.LOCK - 3 * day, self._build("Dallas")), "kept")
        self.assertEqual(known.kept(2026, 6)["games"]["1"]["pick"], "Dallas")
        # asked again a minute later: written recently enough
        self.assertEqual(known.keep(Season(), 2026, 6, self.LOCK, self.LOCK - 3 * day + 60, self._build("Miami")), "fresh")
        # the next day the news has changed, and so does what is kept
        self.assertEqual(known.keep(Season(), 2026, 6, self.LOCK, self.LOCK - 2 * day, self._build("Miami")), "kept")
        # close to the deadline every look is written
        self.assertEqual(known.keep(Season(), 2026, 6, self.LOCK, self.LOCK - 600, self._build("Denver")), "kept")
        self.assertEqual(known.keep(Season(), 2026, 6, self.LOCK, self.LOCK - 300, self._build("Buffalo")), "kept")
        # past the deadline it is frozen as it stood, and nothing rewrites it
        self.assertEqual(known.keep(Season(), 2026, 6, self.LOCK, self.LOCK + 60, self._build("Hindsight")), "frozen")
        self.assertEqual(known.keep(Season(), 2026, 6, self.LOCK, self.LOCK + day, self._build("Hindsight")), "closed")
        kept = known.kept(2026, 6)
        self.assertEqual((kept["games"]["1"]["pick"], kept["frozen"]), ("Buffalo", True))

    def test_a_week_never_written_is_not_written_after_the_deadline(self):
        from acql.model import known
        self.assertEqual(known.keep(Season(), 2026, 5, self.LOCK, self.LOCK + 60, self._build("Hindsight")), "closed")
        self.assertIsNone(known.kept(2026, 5))
        # and a week with nothing priced yet is left for the next look
        self.assertEqual(known.keep(Season(), 2026, 7, self.LOCK, self.LOCK - 86400, lambda s, w: {}), "fresh")
        self.assertIsNone(known.kept(2026, 7))


class TestKindsOfRecipe(unittest.TestCase):
    """The market, football alone, and the two together - compared on the same games."""

    def test_the_best_of_each_kind_is_found_on_the_games_all_priced(self):
        import math
        from acql.model import modellab

        def entry(name, kind, chances):
            made = modellab.Entry(name, "", kind=kind, incumbent=(name == "Deadeye"))
            for i, (chance, hit) in enumerate(chances, 1):
                loss = -math.log(chance if hit else 1 - chance)
                made.log[(1, i)] = modellab.Play(1, i, chance, hit, loss, hit, False)
                made.games += 1
                made.right += hit
            return made

        sure = [(0.7, True), (0.7, True), (0.7, False)]
        unsure = [(0.55, True), (0.55, True), (0.55, False)]
        lab = modellab.Lab([entry("Deadeye", modellab.HYBRID, sure), entry("Line", modellab.MARKET, unsure),
                            entry("Open", modellab.MARKET, sure + [(0.6, True)]),
                            entry("Ratings", modellab.FOOTBALL, [(0.6, False), (0.6, True), (0.6, False)])], 4, (1,))
        best = lab.best_of_each_kind()
        self.assertEqual([(kind, e.name, right, games) for kind, e, right, games in best],
                         [("Market", "Open", 2, 3), ("Football", "Ratings", 1, 3), ("Hybrid", "Deadeye", 2, 3)])
        self.assertEqual(modellab.Lab([], 0, ()).best_of_each_kind(), [])


class TestMarketFootballHybrid(unittest.TestCase):
    """A model is scored on winners, on its margins, and on the side it takes against the closing line."""

    def test_a_side_is_only_taken_when_the_model_stands_off_the_line(self):
        import numpy as np
        from acql.feeds import nflverse
        #                 covers  loses  too close  push   big cover
        guess = np.array([6.0,    1.0,   3.2,       5.0,   -2.0])
        spread = np.array([3.0,   3.0,   3.0,       3.0,    3.0])
        margin = np.array([10.0,  7.0,   1.0,       3.0,   -4.0])
        line = nflverse._score("Football", 2025, guess, margin, spread)
        self.assertEqual((line.bets, line.covers), (3, 2))           # too close and the push are left out
        self.assertEqual((line.big_bets, line.big_covers), (3, 2))
        self.assertEqual((line.games, line.winners), (5, 5))         # the winner is a different question
        self.assertAlmostEqual(line.miss, (4 + 6 + 2.2 + 2 + 2) / 5)
        self.assertAlmostEqual(line.cover_rate, 2 / 3)
        market = nflverse._score("Market", 2025, spread, margin, spread)
        self.assertEqual((market.bets, market.cover_rate), (0, None))    # the line never bets against itself

    def test_seasons_add_up(self):
        from acql.feeds import nflverse
        lines = [nflverse.Line("Hybrid", 2024, 100, 66, 10.0, 60, 31, 8, 5),
                 nflverse.Line("Hybrid", 2025, 300, 190, 9.0, 150, 75, 12, 5),
                 nflverse.Line("Market", 2025, 300, 200, 9.5)]
        total = nflverse._total("Hybrid", lines)
        self.assertEqual((total.season, total.games, total.winners, total.bets, total.covers), (0, 400, 256, 210, 106))
        self.assertAlmostEqual(total.miss, 9.25)                      # weighted by games
        found = nflverse.Families(lines + [total], {})
        self.assertEqual(found.seasons, [2024, 2025])
        self.assertIs(found.of("Hybrid"), total)
        self.assertIsNone(found.of("Football"))

    def test_form_needs_a_few_games_and_is_a_rate(self):
        from collections import deque
        from acql.feeds import nflverse
        # (plays, points added, yards, turnovers) for the team, then the same for its opponent
        game = (60.0, 6.0, 360.0, 1.0, 50.0, -5.0, 250.0, 2.0)
        self.assertEqual(nflverse._form(deque([game, game])), (0.0, 0.0, 0.0, 0.0, 0.0, 2))
        attack, allowed, yards, yards_allowed, turnovers, n = nflverse._form(deque([game] * 4))
        self.assertEqual((attack, allowed, yards, yards_allowed, turnovers, n), (0.1, -0.1, 6.0, 5.0, -1.0, 4))
        self.assertEqual(nflverse._team("OAK"), "LV")
        self.assertEqual(nflverse._team("KC"), "KC")


class TestEveryonesPicks(_ScratchStore):
    """The whole pool's cards for a week, side by side, and how it split on each game."""

    def _season(self):
        from acql.pool import picks
        from acql.models import Game, Week
        season = Season()
        for key, name in (("phelpy", "Phelpy"), ("ann", "Ann"), ("bob", "Bob")):
            season.players[key] = Player(key=key, display=name)
        season.weeks[5] = Week(5, games=[
            Game(1, home="Dallas", away="Tampa Bay", winner="Dallas", line=9.0, line_favourite="Dallas"),
            Game(2, home="Green Bay", away="Chicago"),
        ])
        sheet = picks.WeekPicks(5, games=[
            picks.GamePicks("Tampa Bay", "Dallas", "dal by 9", {"Phelpy": "Dallas", "Ann": "Tampa Bay", "Bob": "Dallas"}),
            picks.GamePicks("Chicago", "Green Bay", "", {"Phelpy": "Chicago", "Ann": "Green Bay", "Bob": "Green Bay"}),
        ], losers={"Phelpy": ["Miami", "Arizona", "Cleveland"]}, suicide={"Ann": "Detroit"})
        return season, sheet

    def test_every_card_is_laid_out_and_marked(self):
        from unittest import mock
        from acql.pool import everyone
        season, sheet = self._season()
        with mock.patch.object(everyone.picks, "load", return_value=sheet):
            board = everyone.build(season, 5)
        self.assertTrue(board.from_sheet)
        self.assertEqual([c.name for c in board.coaches], ["Ann", "Bob", "Phelpy"])
        first, second = board.games
        self.assertEqual((first.away, first.home, first.winner, first.line), ("TB", "DAL", "DAL", "Dallas by 9"))
        self.assertEqual((first.on_away, first.on_home, first.favourite), (1, 2, "DAL"))
        self.assertEqual((second.winner, second.favourite, second.on_home), ("", "GB", 2))
        mine = board.coach("phelpy")
        self.assertEqual(mine.picks, {1: "DAL", 2: "CHI"})
        self.assertEqual([mine.state(g) for g in board.games], ["right", "open"])
        self.assertEqual(board.coach("ann").state(first), "wrong")
        self.assertEqual((mine.right(board.games), mine.with_crowd(board.games)), (1, (1, 2)))
        self.assertEqual((mine.losers, board.coach("ann").suicide), (("MIA", "ARI", "CLE"), "DET"))
        self.assertEqual((board.most_popular.index, len(board.decided)), (1, 1))
        self.assertAlmostEqual(first.share, 2 / 3)
        self.assertIsNone(board.coach(""))

    def test_a_week_with_nothing_is_an_empty_board(self):
        from acql.pool import everyone
        board = everyone.build(Season(), 9)
        self.assertEqual((board.games, board.coaches, board.most_split), ([], [], None))
        self.assertEqual(everyone.short("Los Angeles Chargers"), "LAC")
        self.assertEqual(everyone.short(""), "")


class TestWhyItWasPicked(_ScratchStore):
    """How a pick's number was made, and the quarterback line under it."""

    def _case(self, **over):
        from types import SimpleNamespace
        from acql.model import evidence, pricing
        from acql.models import Game
        game = Game(1, home="Buffalo", away="Miami", **over)
        choice = SimpleNamespace(pick="Buffalo", other="Miami", chance=0.71)
        price = pricing.Price(game, 7.4, 8.7, "market + models", 6.5, 3.0, False, None)
        return game, choice, price, evidence

    def test_the_steps_run_from_the_line_to_the_chance(self):
        game, choice, price, evidence = self._case()
        steps = evidence._steps(game, choice, True, price, None)
        self.assertEqual([s.name for s in steps], ["Betting line", "Computer models", "Deadeye", "Straight up", "Chance"])
        self.assertEqual((steps[0].points, steps[0].words), (6.5, "Buffalo \u22126.5"))
        self.assertAlmostEqual(steps[1].points, 0.9)                       # what the models added
        self.assertEqual((steps[2].words, steps[-1].words), ("Buffalo by 7.4", "71%"))
        case = evidence.Case(1, "Miami @ Buffalo", "Buffalo", "Miami", 0.71, "", False, (), steps)
        self.assertIn("Betting line: Buffalo \u22126.5 (+6.5)", case.made())
        self.assertTrue(case.made().endswith("Chance: 71%"))

    def test_the_road_side_and_a_pool_line_read_from_the_picks_side(self):
        from types import SimpleNamespace
        game, _choice, price, evidence = self._case(line=7.0, line_favourite="Buffalo")
        choice = SimpleNamespace(pick="Miami", other="Buffalo", chance=0.52)
        steps = evidence._steps(game, choice, False, price, None)
        self.assertEqual(steps[0].points, -6.5)                            # Miami is getting the points
        self.assertEqual(steps[2].words, "Buffalo by 7.4")                 # and is still expected to lose
        self.assertEqual((steps[3].name, steps[3].words), ("Pool's line", "Buffalo must win by 7"))
        unpriced = evidence.pricing.Price(game, None, 8.7, "Vegas", None, None, False, None)
        self.assertEqual(evidence._steps(game, choice, False, unpriced, None), ())

    def test_a_quarterback_out_counts_for_the_other_side(self):
        from acql.model import evidence
        from acql import store
        self.assertIsNone(evidence._quarterback("Buffalo", "Miami", "buf", "mia"))      # no report read yet
        store.put("injuries", "latest", {"at": 1_790_000_000.0, "teams": {
            "mia": [{"name": "Tua T", "position": "QB", "status": "out", "what": "Hip", "at": 1.0}],
            "buf": [{"name": "A Tackle", "position": "OT", "status": "out", "what": "Knee", "at": 1.0}]}})
        ours = evidence._quarterback("Buffalo", "Miami", "buf", "mia")
        self.assertEqual((ours.verdict, ours.headline), (evidence.FOR, "Buffalo: none listed \u00b7 Miami: Tua T out"))
        theirs = evidence._quarterback("Miami", "Buffalo", "mia", "buf")
        self.assertEqual(theirs.verdict, evidence.AGAINST)
        store.put("injuries", "latest", {"at": 1_790_000_000.0, "teams": {
            "mia": [{"name": "Tua T", "position": "QB", "status": "questionable", "what": "Hip", "at": 1.0}]}})
        self.assertEqual(evidence._quarterback("Buffalo", "Miami", "buf", "mia").verdict, evidence.INFO)

    def test_a_team_is_judged_on_what_it_adds_less_what_it_allows(self):
        from acql.feeds import nflverse
        good = nflverse.TeamForm("buf", 8, 0.12, -0.05, 6.1, 4.9, -0.5, 2, 4)
        poor = nflverse.TeamForm("mia", 8, -0.04, 0.09, 4.8, 6.0, 0.6, 27, 30)
        self.assertAlmostEqual(good.net, 0.17)
        self.assertLess(poor.net, 0)


class TestReadiness(unittest.TestCase):
    """Whether the dashboard's inputs are up to date, judged against the clock."""

    def test_time_is_said_plainly(self):
        from acql import readiness
        self.assertEqual([readiness.ago(s) for s in (None, 30, 720, 5 * 3600, 3 * 86400)],
                         ["never", "just now", "12 min ago", "5 h ago", "3 days ago"])
        self.assertEqual([readiness.until(s) for s in (40 * 60, 5 * 3600, 3 * 86400, -7200)],
                         ["in 40 min", "in 5 h", "in 3 days", "2 h ago"])

    def test_how_fresh_is_fresh_enough_depends_on_the_deadline(self):
        from acql import headless, readiness
        self.assertEqual(readiness.line_cadence(4 * 86400), headless.LINES_EVERY)
        self.assertEqual(readiness.line_cadence(10 * 3600), headless.LINES_EVERY_LOCK_DAY)
        self.assertEqual(readiness.line_cadence(3600), headless.LINES_EVERY_LAST_HOURS)
        self.assertEqual(readiness.freshness(600, 3600), readiness.OK)
        self.assertEqual(readiness.freshness(7200, 3600), readiness.WARN)
        self.assertEqual(readiness.freshness(7200, 3600, matters=False), readiness.NOTE)
        self.assertEqual((readiness.freshness(None, 3600), readiness.freshness(None, 3600, matters=False)),
                         (readiness.BAD, readiness.NOTE))

    def test_the_verdict_is_the_worst_thing_found(self):
        from acql.readiness import AUTOMATION, BAD, CARD, NOTE, OK, SOURCES, WARN, Check, Readiness
        fine = Readiness([Check(SOURCES, "Scores", OK, "read just now"), Check(CARD, "Your own card", NOTE, "not saved")])
        self.assertTrue(fine.ready)
        self.assertEqual((fine.verdict, fine.worst), ("Ready", NOTE))
        poor = Readiness(fine.checks + [Check(SOURCES, "Betting lines", WARN, "old"),
                                        Check(AUTOMATION, "The nightly run", BAD, "not set up"),
                                        Check(SOURCES, "Weather", WARN, "old")])
        self.assertFalse(poor.ready)
        self.assertEqual((poor.verdict, poor.worst), ("1 thing broken, 2 to look at", BAD))
        groups = poor.grouped()
        self.assertEqual([name for name, _rows in groups], [SOURCES, CARD, AUTOMATION])
        self.assertEqual([c.state for c in groups[0][1]], [WARN, WARN, OK])        # the worst first

    def test_scheduled_jobs_are_read_for_what_they_last_did(self):
        from acql import readiness
        now = 1_790_000_000.0
        states = {
            "ACQL Daily": {"found": True, "last": now - 3600, "next": now + 79200, "result": 0},
            "ACQL Game-day poll": {"found": True, "last": now - 600, "next": now + 60, "result": 1},
            "ACQL Deadeye email": {"found": False},
        }
        daily, poll, mail, training, brief = readiness._tasks(states, now)
        self.assertEqual((training.state, brief.state), (readiness.BAD, readiness.BAD))   # neither registered here
        self.assertEqual((daily.state, daily.headline), (readiness.OK, "Last ran 60 min ago, next in 22 h"))
        self.assertEqual(poll.state, readiness.WARN)                         # its last run ended in an error
        self.assertEqual(mail.state, readiness.BAD)                          # never set up on this computer
        late = dict(states, **{"ACQL Daily": {"found": True, "last": now - 3 * 86400, "next": None, "result": 0}})
        self.assertEqual(readiness._tasks(late, now)[0].state, readiness.WARN)     # three days without a nightly run
        unknown = readiness._tasks(None, now)
        self.assertEqual((len(unknown), unknown[0].state), (1, readiness.NOTE))


class TestWebsitePages(unittest.TestCase):
    """The pool website's template: its pages are wired up, and it stays the league's own."""

    def setUp(self):
        from acql.ui import poolpage
        self.page = poolpage.template()

    def test_every_page_in_the_menu_can_be_drawn(self):
        import re
        views = re.search(r"var VIEWS = \{(.*?)\n  \};", self.page, re.S).group(1)
        names = re.findall(r"^\s+(\w+): \{ label:", views, re.M)
        self.assertIn("picks", names)
        drawn = set(re.findall(r'case "(\w+)": return \w+View\(\);', self.page)) | {"home"}
        self.assertEqual(set(names) - drawn, set())
        menus = re.findall(r'return \[("home".*?)\]\.filter\(available\);', self.page)
        self.assertTrue(menus)
        for listed in menus:
            self.assertEqual({v.strip('" ') for v in listed.split(",")} - set(names), set())

    def test_everyones_picks_follows_the_scores(self):
        self.assertIn("function picksView()", self.page)
        self.assertIn('["home", "live", "picks", "pools", "me"].indexOf(state.view)', self.page)
        self.assertIn('if (view === "live" || view === "picks") return hasLive();', self.page)

    def test_the_model_and_the_dashboard_stay_off_the_site(self):
        for word in ("Deadeye", "Operation Sunday"):
            self.assertNotIn(word, self.page)


class TestRecipe(unittest.TestCase):
    """The recipe taken apart: the share simulator, the ingredients, and a line's readings."""

    def _nudge(self, index, picks, credited="", models=30):
        from acql.model import recipe
        return recipe.Nudge(index, f"Away {index} @ Home {index}", 3.0, models, 2.0, 1.5,
                            {share: pick for share, pick in picks.items()}, credited)

    def test_the_shares_run_from_none_to_all_and_include_the_one_in_use(self):
        from acql.model import recipe
        steps = recipe.shares(0.33)
        self.assertEqual((steps[0], steps[-1], len(steps)), (0.0, 1.0, 22))
        self.assertIn(0.33, steps)
        self.assertEqual(len(recipe.shares(0.45)), 21)                    # already one of the twentieths

    def test_the_starting_weight_is_the_approved_quarter(self):
        # Sam's choice on 2026-10-12 (0.33 before; see predictions.CONSENSUS_WEIGHT).
        from acql.model import predictions, pricing, recipe
        self.assertEqual(predictions.CONSENSUS_WEIGHT, 0.25)
        self.assertEqual(pricing.Context(1).weight, 0.25)
        self.assertIn(0.25, predictions.fit_steps(predictions.CONSENSUS_WEIGHT))
        self.assertIn(0.25, recipe.shares(predictions.CONSENSUS_WEIGHT))
        self.assertEqual(pricing.model_weight(None), (0.25, None))      # no season yet: the start

    def test_a_flip_is_a_different_side_from_the_one_in_use(self):
        from acql.model import recipe
        steady = self._nudge(1, {0.0: ("Home 1", 0.60), 0.33: ("Home 1", 0.62), 1.0: ("Home 1", 0.66)}, "Home 1")
        turns = self._nudge(2, {0.0: ("Away 2", 0.52), 0.33: ("Home 2", 0.51), 1.0: ("Home 2", 0.56)}, "Away 2")
        thin = self._nudge(3, {0.0: ("Home 3", 0.7), 0.33: ("Home 3", 0.7), 1.0: ("Home 3", 0.7)}, models=4)
        sim = recipe.Simulation(5, 0.33, (0.0, 0.33, 1.0), [steady, turns, thin])
        self.assertEqual([g.index for g in sim.flipped(0.0)], [2])
        self.assertEqual(sim.flipped(1.0), [])
        self.assertEqual(sim.first_flip(), 0.0)
        self.assertEqual((sim.record(0.33), sim.record(0.0)), ((1, 2), (2, 2)))   # the flip would have been right
        self.assertEqual((turns.right(0.0), turns.right(0.33), thin.right(0.33)), (True, False, None))
        self.assertTrue(thin.thin)
        self.assertFalse(steady.thin)
        self.assertIn("at 0%, 1 (Home 2 \u2192 Away 2)", recipe.summary(sim))
        calm = recipe.Simulation(5, 0.33, (0.0, 0.33, 1.0), [steady])
        self.assertIn("Nothing flips", recipe.summary(calm))
        none = recipe.Simulation(6, 0.33, (0.0, 0.33, 1.0), [self._nudge(1, {0.33: ("Home 1", 0.6)}, models=0)])
        self.assertIn("No published models", recipe.summary(none))

    def test_an_ingredient_is_counted_by_the_picks_it_turned(self):
        from acql.model import recipe
        part = recipe.Ingredient(recipe.MODELS, games=45, changed=3, helped=2, examples=("a", "b", "c"))
        self.assertEqual((part.hurt, part.net), (1, 1))
        idle = recipe.Ingredient(recipe.MONEYLINE, games=0)
        self.assertEqual((idle.changed, idle.hurt, idle.net), (0, 0, 0))

    def test_the_side_a_number_takes_follows_the_pools_rule(self):
        from acql.model import recipe
        from acql.models import Game
        straight = Game(1, home="Buffalo", away="Miami")
        lined = Game(2, home="Buffalo", away="Miami", line=7.0, line_favourite="Buffalo")
        self.assertEqual(recipe._side(straight, 3.0), "Buffalo")
        self.assertEqual(recipe._side(straight, -3.0), "Miami")
        self.assertEqual(recipe._side(lined, 3.0), "Miami")               # favoured, but not by the pool's seven
        self.assertEqual(recipe._side(lined, 12.0), "Buffalo")
        self.assertEqual(recipe._side(straight, None), "")

    def test_a_line_is_read_from_the_picks_side(self):
        from acql.model import recipe
        flight = recipe.Flight(1, "Miami @ Buffalo", "Buffalo", [(100.0, 3.0), (200.0, 3.5), (300.0, 4.5)], models=6.0)
        self.assertEqual((flight.opened, flight.latest, flight.moved), (3.0, 4.5, 1.5))
        self.assertTrue(flight.toward_models)                             # toward the pick, and toward the models
        away = recipe.Flight(2, "Miami @ Buffalo", "Buffalo", [(100.0, 3.0), (200.0, 1.5)], models=6.0)
        self.assertEqual((away.moved, away.toward_models), (-1.5, False))
        still = recipe.Flight(3, "Miami @ Buffalo", "Buffalo", [(100.0, 3.0)], models=6.0)
        self.assertEqual((still.moved, still.toward_models), (0.0, None))
        self.assertEqual((recipe.Flight(4, "x", "y", []).opened, recipe.Flight(4, "x", "y", []).moved), (None, 0.0))


class TestGhostRecord(unittest.TestCase):
    """The ghost's won-lost record is out of the same chances as a coach's."""

    @staticmethod
    def _week(number, regular, losers):
        from acql.model import ghost
        return ghost.GhostWeek(week=number, regular=regular, losers=losers, losers_known=True, games=16, place=1,
                               tied=0, entrants=35, best=regular + losers, payout=0.0, money=0.0, rebuilt=False)

    def _weeks(self):
        return [self._week(1, 12, 1), self._week(2, 13, 1), self._week(3, 7, 0), self._week(4, 11, 0)]

    def test_big_loser_picks_are_chances_too(self):
        from acql.model import ghost
        mine = ghost.GhostSeason(self._weeks(), 4, field_chances=75)
        self.assertEqual((mine.wins, mine.losses, mine.chances), (45, 30, 75))
        self.assertAlmostEqual(mine.pct, 0.6)

    def test_without_the_pools_count_it_uses_its_own_weeks(self):
        from acql.model import ghost
        alone = ghost.GhostSeason(self._weeks(), 4)
        self.assertEqual((alone.chances, alone.losses), (76, 31))
        short = ghost.GhostSeason(self._weeks()[:3], 4, field_chances=75)      # a week with no card
        self.assertEqual((short.chances, short.wins, short.losses), (57, 34, 23))

    def test_the_pools_count_is_what_every_coach_is_out_of(self):
        from acql.model import ghost
        season = Season()
        for name, wins, losses in (("a", 44, 31), ("b", 39, 36), ("c", 30, 45), ("late", 20, 36), ("new", 0, 0)):
            season.players[name] = Player(key=name, display=name, wins=wins, losses=losses)
        self.assertEqual(ghost.field_chances(season), 75)
        self.assertEqual(ghost.field_chances(Season()), 0)


class TestWording(unittest.TestCase):
    """Counts and places, said one way for the whole app."""

    def test_counts_and_places(self):
        from acql import wording
        self.assertEqual([wording.plural(n, "game") for n in (0, 1, 2)], ["0 games", "1 game", "2 games"])
        self.assertEqual((wording.plural(3, "coach", "coaches"), wording.plural(1, "coach", "coaches")),
                         ("3 coaches", "1 coach"))
        self.assertEqual([wording.ordinal(n) for n in (1, 2, 3, 4, 11, 12, 13, 21, 22, 103, 111)],
                         ["1st", "2nd", "3rd", "4th", "11th", "12th", "13th", "21st", "22nd", "103rd", "111th"])
        self.assertEqual((wording.suffix(35), wording.suffix(2)), ("th", "nd"))


class TestWorkedOutOnce(_ScratchStore):
    """Each shortcut gives exactly the answer the long way gave."""

    def test_a_win_chance_is_the_same_sum_looked_up(self):
        from acql import analytics
        for sd in (analytics.SPREAD_SD, 15.3):
            for spread in (-9.5, 0.0, 3.0, 7.0, 13.5):
                chances = analytics.margin_distribution(round(spread, 1), sd)
                for need in (-70, -3, 0, 1, 7, 8, 15, 60, 61, 99):
                    self.assertEqual(analytics.cover_chance(spread, need, sd),
                                     sum(p for margin, p in chances if margin >= need))
        self.assertEqual(analytics.cover_chance(6.5, 7.5), analytics.cover_chance(6.5, 8))

    def test_stored_forecasts_are_built_once_until_the_store_changes(self):
        import statistics
        from acql.model import predictions
        opinions = (4.0, 5.0, 6.5, 7.0, 9.0, 3.0)
        first = predictions.Forecast("Tampa Bay", "Dallas", market=6.0, opinions=opinions)
        predictions.save(5, [first], "test")
        once, source = predictions.load(5)
        again, _ = predictions.load(5)
        self.assertEqual((source, len(once)), ("test", 1))
        self.assertIs(once[0], again[0])                  # the same forecast, with what it has worked out
        self.assertIsNot(once, again)                     # in a list of the caller's own
        self.assertEqual(once[0].disagreement, statistics.pstdev(predictions.trimmed(list(opinions))))
        self.assertEqual(once[0], first)                  # what it worked out is no part of what it is
        predictions.save(5, [first, predictions.Forecast("Chicago", "Green Bay", market=-2.5)], "newer")
        later, source = predictions.load(5)
        self.assertEqual((source, len(later)), ("newer", 2))

    def test_the_machine_is_grown_once_and_read_back(self):
        from unittest import mock
        import numpy as np
        from acql.model import contenders
        rng = np.random.default_rng(3)
        X = rng.normal(size=(900, len(contenders.INPUTS)))
        y = X[:, 0] * 3 + rng.normal(size=900)
        line, start, trees = contenders._fitted(X, y)
        with mock.patch.object(contenders, "boost", side_effect=AssertionError("grown again")):
            line2, start2, trees2 = contenders._fitted(X, y)
        self.assertEqual((start, trees, line[3]), (start2, trees2, line2[3]))
        for a, b in zip(line[:3], line2[:3]):
            self.assertTrue(np.array_equal(a, b))
        with mock.patch.object(contenders, "boost", return_value=(0.0, [])) as grown:
            contenders._fitted(X, y + 1)                  # other games: fitted afresh
        self.assertEqual(grown.call_count, 1)

    def test_the_money_tables_skip_only_what_cannot_matter(self):
        import numpy as np
        from acql.pool import moneyplan
        # one coin everybody lands: all 35 cards score 1 and share the week
        money, first = moneyplan.field_tables(np.ones((5, 1)), 34)
        self.assertTrue(np.allclose(first[:, 1], 1 / 35) and np.allclose(money[:, 1], 0) and np.allclose(money[:, 0], -1))
        # a coin that never lands changes nothing, wherever it sits
        rng = np.random.default_rng(11)
        coins = np.concatenate([rng.random((300, 6)), rng.random((300, 9)) * (rng.random((300, 9)) < 0.15)], axis=1)
        money, first = moneyplan.field_tables(coins, 34)
        padded = np.insert(coins, 4, 0.0, axis=1)
        money2, first2 = moneyplan.field_tables(padded, 34)
        self.assertTrue(np.array_equal(first, first2[:, :-1]))
        self.assertTrue(np.allclose(money, money2[:, :-1], atol=1e-9))   # one more column of rounding, no more


class TestChartsDrawOnce(unittest.TestCase):
    """A chart asked to draw what it already shows doesn't."""

    @classmethod
    def setUpClass(cls):
        import os

        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        try:
            from PySide6.QtWidgets import QApplication
        except ImportError as exc:  # pragma: no cover - Qt absent
            raise unittest.SkipTest(f"PySide6 unavailable: {exc}")
        cls.app = QApplication.instance() or QApplication([])

    def test_the_same_grid_is_drawn_once_and_a_changed_one_again(self):
        from acql.ui.charts import HeatmapChart
        from acql.ui.theme import DARK, LIGHT
        drawn = []

        class Counting(HeatmapChart):
            def _draw(self, *args, **kwargs):
                drawn.append(1)
                super()._draw(*args, **kwargs)

        chart = Counting(DARK)
        rows, cols, values = ["Ann", "Bob"], ["1", "2"], [[1.0, 2.0], [3.0, None]]
        chart.plot(rows, cols, values, tooltip_fn=lambda r, c, v: "old")
        chart.plot(list(rows), list(cols), [list(v) for v in values], tooltip_fn=lambda r, c, v: "new")
        self.assertEqual(len(drawn), 1)
        self.assertEqual(chart._heatmap[3]("Ann", "1", 1.0), "new")      # the new words, on the same drawing
        grid = chart.figure.axes[0]
        self.assertEqual((len(grid.patches), len(grid.collections)), (0, 1))   # one shape, however many cells
        values[0][0] = 5.0                                               # the same list, changed in place
        chart.plot(rows, cols, values, tooltip_fn=lambda r, c, v: "new")
        self.assertEqual(len(drawn), 2)
        chart.empty()
        chart.plot(rows, cols, values, tooltip_fn=lambda r, c, v: "new")
        self.assertEqual(len(drawn), 3)
        chart.set_palette(LIGHT)                                         # other colours: drawn again
        chart.plot(rows, cols, values, tooltip_fn=lambda r, c, v: "new")
        self.assertEqual(len(drawn), 4)

    def test_a_chart_given_something_it_cannot_read_is_always_drawn(self):
        from acql.ui import charts
        self.assertIsNone(charts._mark(([object()], {})))
        self.assertIsNone(charts._mark(((lambda: 1,), {})))
        self.assertIsNotNone(charts._mark(((lambda: 1,), {}), True))
        self.assertNotEqual(charts._mark(([1],)), charts._mark(([1.0],)))
        self.assertIsNone(charts._mark((list(range(charts.MARK_BUDGET + 5)),)))


class TestAmericanOdds(unittest.TestCase):
    """American prices to chances, and the books' margin taken out."""

    def test_valid_prices(self):
        from acql.feeds.odds import _implied
        for price, chance in [(-200, 2 / 3), (-150, 0.6), (-110, 110 / 210), (-100, 0.5), (100, 0.5),
                              (110, 100 / 210), (130, 100 / 230), (200, 1 / 3), (-10000, 10000 / 10100)]:
            self.assertAlmostEqual(_implied(price), chance, places=12, msg=price)

    def test_impossible_prices_are_not_read(self):
        from acql.feeds.odds import _implied
        from acql.training.data import _implied as training_implied
        for bad in (-99, 99, 50, -50, 0, 0.5, "", "-110", "abc", None, True, float("nan"),
                    float("inf"), float("-inf")):
            self.assertIsNone(_implied(bad), bad)
            self.assertIsNone(training_implied(bad), bad)

    def _event(self, *books):
        return [{"home_team": "Kansas City Chiefs", "away_team": "Denver Broncos",
                 "commence_time": "2026-10-18T17:00:00Z",
                 "bookmakers": [{"markets": [
                     {"key": "h2h", "outcomes": [{"name": "Kansas City Chiefs", "price": h},
                                                 {"name": "Denver Broncos", "price": a}]},
                     {"key": "spreads", "outcomes": [{"name": "Kansas City Chiefs", "point": -4.5}]}]}
                     for h, a in books]}]

    def test_no_vig_sums_to_one_and_books_are_combined_by_median(self):
        from acql.feeds.odds import _implied, parse
        h, a = _implied(-150), _implied(130)
        self.assertGreater(h + a, 1)
        self.assertAlmostEqual(h / (h + a) + a / (h + a), 1.0, places=12)
        one = parse(self._event((-150, 130)))[0]
        self.assertAlmostEqual(one.home_win, round(0.6 / (0.6 + 100 / 230), 4))
        # two identical books (a duplicate) change nothing; a conflicting third moves the median
        self.assertEqual(parse(self._event((-150, 130), (-150, 130)))[0].home_win, one.home_win)
        three = parse(self._event((-150, 130), (-150, 130), (-110, -110)))[0]
        self.assertEqual(three.home_win, one.home_win)

    def test_a_book_with_an_impossible_or_missing_price_is_left_out(self):
        from acql.feeds.odds import parse
        good = parse(self._event((-150, 130)))[0].home_win
        self.assertEqual(parse(self._event((-150, 130), (50, -50)))[0].home_win, good)
        self.assertEqual(parse(self._event((-150, 130), (-150, None)))[0].home_win, good)
        self.assertIsNone(parse(self._event((50, -50)))[0].home_win)     # the spread still stands


class TestMarketColumnsAreNotModels(unittest.TestCase):
    """A betting line or a summary column must never be counted as a model."""

    HEAD = "road,home,line,lineopen,lineca,linesag,linemass,linepig"
    ROW = "Tampa Bay,Dallas,3.5,3,3.5,4.5,2,6"

    def _read(self, head, row):
        from acql.model import predictions
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "nflpredictions week 3.csv"
            path.write_text(f"{head}\n{row}\n", encoding="utf-8")
            return predictions.read_csv(path)[0]

    def test_extra_market_and_summary_columns_leave_the_consensus_alone(self):
        plain = self._read(self.HEAD, self.ROW)
        self.assertEqual(sorted(plain.opinions), [2.0, 4.5, 6.0])
        more = self._read(self.HEAD + ",linemidweek,linemed,linemin,linemax,lineavg,linemedian,linestd",
                          self.ROW + ",-20,9,-30,30,4.2,4.5,1.7")
        self.assertEqual(sorted(more.opinions), sorted(plain.opinions))
        self.assertEqual(more.consensus, plain.consensus)
        self.assertEqual(more.market, 3.5)                       # the market itself is still read
        self.assertEqual(more.opening, 3.0)

    def test_the_market_still_enters_through_the_blend(self):
        f = self._read(self.HEAD, self.ROW)
        # edge = consensus - market; the blend moves the market a share of that
        self.assertAlmostEqual(f.edge, f.consensus - 3.5)
        self.assertAlmostEqual(f.expected(0.25), 3.5 + f.model_weight(0.25) * f.live_edge())
        self.assertEqual(f.expected(0.0), 3.5)


class TestPredictionsGate(_ScratchStore):
    """A season archive must never be taken for one week's predictions file.

    On 2026-10-18 nfl24.csv (285 games of 2024, in Downloads) was read as week
    7's file because its division games covered most of the week's pairings.
    """

    TEAMS = ["Chicago", "New England", "Green Bay", "Detroit", "Dallas", "Philadelphia",
             "Buffalo", "Miami", "Denver", "Seattle", "Atlanta", "Tampa Bay", "Houston",
             "Tennessee", "Baltimore", "Pittsburgh", "Arizona", "Los Angeles Rams",
             "Minnesota", "Carolina", "Cleveland", "Cincinnati", "Kansas City", "Las Vegas"]

    def setUp(self):
        super().setUp()
        from acql.models import Game
        t = self.TEAMS
        self.slate = [Game(i + 1, t[2 * i], t[2 * i + 1]) for i in range(12)]
        self.dir = self.folder / "files"
        self.dir.mkdir()

    def _write(self, name, rows, extra_head="", extra=""):
        path = self.dir / name
        lines = ["road,home,line,linesag,linemass" + extra_head]
        lines += [f"{road},{home},3.5,4,2{extra}" for road, home in rows]
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        return path

    def _week_rows(self):
        return [(g.away, g.home) for g in self.slate]

    def test_a_weeks_file_passes(self):
        from acql.model import predictions
        self.assertIsNone(predictions.refusal(self._write("nflpredictions.csv", self._week_rows()), self.slate))

    def test_archive_name_scores_size_and_duplicates_are_refused(self):
        from acql.model import predictions
        week = self._week_rows()
        season = week * 3 + [(a, b) for a in self.TEAMS[:12] for b in self.TEAMS[12:]]
        self.assertIn("archive", predictions.refusal(self._write("nfl24.csv", week)))
        self.assertIn("final scores", predictions.refusal(
            self._write("season.csv", week, ",hscore,rscore,week", ",24,17,7")))
        self.assertIn("games", predictions.refusal(self._write("season2.csv", season), self.slate))
        self.assertIn("twice", predictions.refusal(self._write("dupes.csv", week + week[:1])))
        wild = self._write("wild.csv", week)
        wild.write_text(wild.read_text().replace(",3.5,", ",45,", 1), encoding="utf-8")
        self.assertIn("betting line", predictions.refusal(wild))

    def test_find_file_skips_the_archive_and_records_why(self):
        from acql.model import predictions
        week = self._week_rows()
        season = week * 2 + [(a, b) for a in self.TEAMS[:12] for b in self.TEAMS[12:]]
        self._write("nflpredictions.csv", week)
        time.sleep(0.05)
        newest = self._write("nfl lines.csv", season)        # newer, covers the slate, but a season
        original = predictions.search_folders
        predictions.search_folders = lambda workbook=None: [self.dir]
        try:
            found = predictions.find_file(7, None, self.slate)
        finally:
            predictions.search_folders = original
        self.assertEqual(found.name, "nflpredictions.csv")
        refused = {Path(r["path"]).name: r["reason"] for r in predictions.turned_away()}
        self.assertIn(newest.name, refused)


class TestSubmittedCardCheck(_ScratchStore):
    """The sheet's row for you against the card you saved: a slip on the form is caught once."""

    def _pair(self, sheet_side="Philadelphia", losers=("dal", "nyg", "car")):
        from acql.pool import cards, picks
        saved = cards.Snapshot(at=1.0, take={1: "Chicago", 2: "Philadelphia"}, chance={1: .6, 2: .55},
                               label={1: "New England @ Chicago", 2: "Dallas @ Philadelphia"},
                               losers=["dal", "nyg", "car"])
        sheet = picks.WeekPicks(7, games=[
            picks.GamePicks("New England", "Chicago", "", {"Phelpy": "Chicago", "Jules": "New England"}),
            picks.GamePicks("Dallas", "Philadelphia", "", {"Phelpy": sheet_side, "Jules": "Dallas"}),
        ], losers={"Phelpy": list(losers)})
        return saved, sheet

    def test_a_matching_sheet_passes(self):
        from acql.pool import cardcheck
        found = cardcheck.compare(7, *self._pair(), "phelpy ")
        self.assertTrue(found.matches)
        self.assertEqual(found.games, 2)

    def test_a_different_pick_and_big_loser_are_reported_once(self):
        from acql.pool import cardcheck
        found = cardcheck.compare(7, *self._pair("Dallas", ("dal", "nyg", "mia")), "Phelpy")
        self.assertEqual([d.what for d in found.diffs], ["Game 2", "Big Losers"])
        self.assertIn("sheet has Dallas", found.lines()[0])
        self.assertTrue(cardcheck.keep(found))
        self.assertFalse(cardcheck.keep(found))          # the same difference is not news twice

    def test_nothing_to_compare_says_why(self):
        from acql.pool import cardcheck
        saved, sheet = self._pair()
        self.assertIn("no saved card", cardcheck.compare(7, None, sheet, "Phelpy").reason)
        self.assertIn("isn't on the pick sheet", cardcheck.compare(7, saved, sheet, "Nobody").reason)


class TestConsistencySentinel(_ScratchStore):
    """Sources that describe the same game must agree (consistency.py).

    The 2026-10-18 case: the week-7 record had Chicago -6.5 beside moneylines
    that made Chicago 54% - the line had come from a 2024 archive file."""

    def _season(self, pool_line=None):
        from acql.models import Game, Season, Week
        game = Game(1, "Chicago", "New England", line=pool_line, line_favourite="Chicago" if pool_line else "")
        return Season(title="ACQL 2026", weeks={7: Week(7, [game])})

    def _forecast(self, market, win, total=50.0):
        from acql.model import predictions
        predictions.save(7, [predictions.Forecast("New England", "Chicago", market=market, opinions=(market,),
                                                  market_win=win, total=total)], "test.csv")

    def test_agreeing_sources_raise_nothing(self):
        from acql import consistency
        self._forecast(3.0, 0.58)
        self.assertEqual(consistency.week(self._season(3.5), 7), [])

    def test_line_against_moneyline_total_and_pool_line(self):
        from acql import consistency
        self._forecast(6.5, 0.54, total=12.0)
        kinds = {p.kind for p in consistency.week(self._season(12.0), 7)}
        self.assertEqual(kinds, {"moneyline", "total", "pool line"})

    def test_open_record_against_the_line_now(self):
        from acql import consistency, store
        self._forecast(3.0, 0.58)
        store.put("known", "2026-7", {"at": 1.0, "frozen": False, "games": {
            "1": {"line": "Chicago −6.5", "favourite": "Chicago", "pool_line": ""}}})
        found = [p for p in consistency.week(self._season(), 7) if p.kind == "record"]
        self.assertEqual(len(found), 1)
        self.assertIn("6.5", found[0].text)


class TestFlipPoints(unittest.TestCase):
    """The margin at which a pick turns: just short of it the pick holds, just past it it changes."""

    def test_flip_points_on_straight_and_lined_games(self):
        from acql import analytics
        from acql.model import flips
        from acql.models import Game
        cases = [(Game(1, "Chicago", "New England"), 3.0),
                 (Game(2, "Houston", "New York Giants", line=10.0, line_favourite="Houston"), 10.2),
                 (Game(3, "Minnesota", "Indianapolis", line=7.0, line_favourite="Minnesota"), 4.0),
                 (Game(4, "Dallas", "Philadelphia"), -6.5)]
        for game, margin in cases:
            at = flips.flip_margin(game, margin, 13.0)
            self.assertIsNotNone(at, game)
            before = analytics.predict_margin(game, margin, sd=13.0).pick
            step = 0.05 if at > margin else -0.05
            self.assertEqual(analytics.predict_margin(game, at - step, sd=13.0).pick, before)
            past = analytics.predict_margin(game, at + 10 * step, sd=13.0)
            self.assertTrue(past.pick != before or past.reason == "pick 'em", (game, margin, at))

    def test_key_numbers_and_fragility(self):
        from acql.model import flips
        f = flips.Flip(1, "a @ b", "b", 0.55, 3.5, 2.5, 0.8)
        self.assertTrue(f.fragile and f.crosses_key)
        self.assertFalse(flips.Flip(1, "a @ b", "b", 0.6, 5.5, 4.0, 1.5).crosses_key)


class TestStandingNotes(_ScratchStore):
    """A note the load repeats every run is logged once a day, saying how long it has stood."""

    def test_repeats_are_held_back_and_resurface_daily(self):
        from acql import headless
        t = 1_790_000_000.0
        self.assertEqual(headless._new_notes(["A", "B"], t), ["A", "B"])
        self.assertEqual(headless._new_notes(["A", "B", "C"], t + 600), ["C"])
        again = headless._new_notes(["A"], t + 2 * 86400)
        self.assertEqual(again, ["A (standing for 2 days)"])
        self.assertEqual(headless._new_notes(["B"], t + 2 * 86400 + 60), ["B"])   # gone and back: new again


class TestBackupMirror(_ScratchStore):
    """Each day's copy is also kept outside the synced folder, checked sound, and pruned."""

    def test_mirror_copies_checks_and_prunes(self):
        from acql import store
        store.put("meta", "x", 1)
        daily = self.folder / "daily"
        mirror = self.folder / "mirror"
        for day in range(1, 26):                          # 25 days of October 2026
            target = daily / f"acql-2026-10-{day:02d}.db"
            store.backup_now(target)
            status = store.mirror_backup(target, mirror)
        self.assertTrue(status["ok"])
        self.assertGreaterEqual(status["documents"], 1)
        kept = sorted(p.name for p in mirror.glob("acql-*.db"))
        self.assertEqual(kept[-14:], [f"acql-2026-10-{d:02d}.db" for d in range(12, 26)])
        self.assertIn("acql-2026-10-04.db", kept)          # a Sunday, kept as a weekly copy
        self.assertNotIn("acql-2026-10-05.db", kept)
        broken = mirror / "acql-2026-10-26.db"
        broken.write_bytes(b"not a database")
        self.assertFalse(store.verify(broken)[0])


class TestStandingIssues(_ScratchStore):
    """Notes the nightly run keeps repeating show once on Readiness, with when they began."""

    def test_standing_after_a_day_and_fix_clears(self):
        from acql import headless, readiness
        old = "Stored pick sheets from another season were ignored (week 1). Clear them."
        day = 86400
        headless._new_notes([old, "fresh note"], now=1_000_000)
        self.assertIsNone(readiness._standing(1_000_000 + 3600))      # not a day yet
        headless._new_notes([old], now=1_000_000 + 2 * day)
        headless._new_notes([old, "<odd> note"], now=1_000_000 + 2 * day + 60)
        found = readiness._standing(1_000_000 + 2 * day + 120)
        self.assertEqual(found.name, "Standing issues")
        self.assertEqual(found.fix, "old_sheets")
        self.assertIn("1 note", found.headline)
        self.assertNotIn("fresh note", found.detail)                  # gone from the last run
        self.assertNotIn("<odd>", found.detail)
        readiness.forget_standing(readiness.OLD_SHEETS)
        self.assertIsNone(readiness._standing(1_000_000 + 2 * day + 120))


class TestCatchUp(_ScratchStore):
    """Missed jobs that are safe to run late start when the dashboard opens; timed ones don't."""

    def test_only_safe_overdue_jobs_and_not_twice(self):
        from acql import heal
        now = 2_000_000_000.0
        states = {
            "ACQL Daily": {"found": True, "last": now - 3 * 86400, "next": now + 3600, "result": 0},
            "ACQL Training": {"found": True, "last": None, "next": None, "result": None},
            "ACQL Deadeye email": {"found": True, "last": now - 30 * 86400, "next": None, "result": 0},
            "ACQL Sunday brief": {"found": True, "last": now - 30 * 86400, "next": None, "result": 0},
        }
        ran = []
        self.assertEqual(heal.catch_up(states, now, run=ran.append), ["ACQL Daily", "ACQL Training"])
        self.assertEqual(heal.catch_up(states, now + 3600, run=ran.append), [])   # not again so soon
        self.assertEqual(len(ran), 2)
        fine = dict(states, **{"ACQL Daily": {"found": True, "last": now - 3600, "next": None, "result": 0},
                               "ACQL Training": {"found": True, "last": None, "next": now + 86400}})
        self.assertEqual(heal.overdue(fine, now), [])                # recent, or about to run
        self.assertEqual(heal.overdue(None, now), [])

    def test_a_failed_start_is_tried_again(self):
        from acql import heal
        now = 2_000_000_000.0
        states = {"ACQL Daily": {"found": True, "last": now - 3 * 86400, "next": None}}

        def refuse(_name):
            raise OSError("access denied")
        self.assertEqual(heal.catch_up(states, now, run=refuse), [])
        self.assertEqual(heal.catch_up(states, now + 60, run=lambda n: None), ["ACQL Daily"])


class TestAutopsyBrief(_ScratchStore):
    """The week's look back: written from kept facts, sent once, and never once it is stale."""

    def test_compose_says_what_happened(self):
        from acql.site import autopsybrief as a
        from acql.version import MODEL_NAME
        f = a.Facts(6, you="Phelpy", hits=["Kansas City (A @ B)"], misses=["Seattle (Seattle @ Denver)"],
                    rank=2, model_right=7, model_games=14, luck=(8, 7.8),
                    splits=[("LAC @ KC", "Kansas City", "LA Chargers", "you")],
                    swing=("Houston @ Jacksonville", "Houston", 2, 35),
                    big_loser=[("Arizona", False)], suicide=("Buffalo", True), spots=["Division games: 11 of 27"])
        b = a.compose(None, 6, found=f)
        self.assertIn("Week 6 autopsy: you 1-1", b.subject)
        for bit in ("1 of 2, week rank 2", f"{MODEL_NAME}: 7 of 14", "deserved 7.8 wins and got 8",
                    "you were right", "only 2 of 35", "Arizona missed", "Buffalo - still alive", "Division games"):
            self.assertIn(bit, b.text)
        self.assertIn("Only you get this email", b.html)

    def test_sent_once_and_skipped_when_stale(self):
        from unittest import mock
        from acql import headless, store
        from acql.feeds import oddslock
        from acql.site import autopsybrief as a
        brief = a.Brief(6, "s", "t", "h")
        sent = []
        with mock.patch.object(a, "finished", return_value=6), mock.patch.object(a, "compose", return_value=brief),                 mock.patch.object(headless.results, "season_year", return_value=2026):
            with mock.patch.object(oddslock, "lock_time", return_value=2e9):
                self.assertEqual(headless._autopsy_brief(None, now=1e9, send=lambda b: sent.append(b) or ["me"]), 6)
                self.assertIsNone(headless._autopsy_brief(None, now=1e9, send=sent.append))   # once only
            store.put("autopsybrief", "sent", 5)
            with mock.patch.object(oddslock, "lock_time", return_value=1e9 - 1):          # week 7 already locked
                self.assertIsNone(headless._autopsy_brief(None, now=1e9, send=sent.append))
            self.assertEqual(store.get("autopsybrief", "sent", 0), 6)
        self.assertEqual(len(sent), 1)


class TestRiverChart(unittest.TestCase):
    """The day's chance as a line; the big swings are the looks that moved it ten points."""

    def test_swings(self):
        from PySide6.QtWidgets import QApplication
        QApplication.instance() or QApplication([])
        from acql.ui.pages.gameday import RiverChart
        from acql.ui.theme import DARK
        chart = RiverChart(DARK)
        chart.set_points([(3, 0.40), (1, 0.58), (2, 0.55), (4, 0.62)])
        self.assertEqual(chart.swings(), [(3, 0.55, 0.40), (4, 0.40, 0.62)])
        chart.resize(400, 120)
        chart.grab()            # paints without error


class TestRawArchive(unittest.TestCase):
    """Every payload kept as it came, once per content, findable by source and time."""

    def setUp(self):
        from acql import rawarchive
        self.was = rawarchive._OVERRIDE
        self.root = Path(tempfile.mkdtemp(prefix="acql-raw-"))
        rawarchive.use(self.root)

    def tearDown(self):
        from acql import rawarchive
        rawarchive.use(self.was)

    def test_keep_dedupe_latest_read(self):
        from acql import rawarchive
        a = rawarchive.keep("odds", '{"x": 1}', now=1000)
        b = rawarchive.keep("odds", '{"x": 1}', now=2000)
        rawarchive.keep("odds", '{"x": 2}', now=3000)
        self.assertEqual(a, b)
        self.assertEqual(len(list(self.root.rglob("*.gz"))), 2)          # same content, one file
        self.assertEqual(len(rawarchive.index("odds")), 3)              # but every fetch is on record
        self.assertEqual(rawarchive.read(rawarchive.latest("odds", before=2500)), b'{"x": 1}')
        self.assertEqual(rawarchive.read(rawarchive.latest("odds")), b'{"x": 2}')
        self.assertIsNone(rawarchive.latest("injuries"))
        self.assertTrue(rawarchive.open_copy(rawarchive.latest("odds")).read_text().startswith("{"))

    def test_a_fetch_keeps_its_payload(self):
        import io, json as _json
        from acql import rawarchive
        from acql.feeds import results

        class Reply(io.BytesIO):
            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False
        results.fetch(2026, 3, opener=lambda r: Reply(_json.dumps({"events": []}).encode()))
        row = rawarchive.latest("espn-scores")
        self.assertEqual(row["label"], "2026 week 3")


class TestGoldenWeek(_ScratchStore):
    """A whole made-up week priced end to end, held to a fingerprint.

    Any change to how a game is priced - the share of the models, the home
    lean, the curve, the pool's line rule, the Big Loser chances - changes the
    picks, and that needs Sam's approval. This fails on every such change,
    approved or not; when one is approved, run this test, read the new values
    it prints, and put the new fingerprint below (say so in the commit).
    Made-up numbers only: the repository is public.
    """

    FINGERPRINT = "15892b819a2530ac"
    TEAMS = [("New England", "Chicago", 3.0, [1.4, -2.8, 2.4, 1.1, 2.6, 1.0, 3.5], 0.59, 50.5),
             ("Pittsburgh", "New Orleans", -3.5, [-1.2, -4.0, -2.5, -3.3, -0.4, -5.1, -2.0], 0.38, 44.5),
             ("San Francisco", "Atlanta", -3.0, [-4.4, -1.5, -3.8, -2.2, -6.0, -2.9, -3.1], 0.39, 49.5),
             ("Cincinnati", "Baltimore", 2.5, [3.9, 1.0, 2.2, 4.4, 0.5, 2.8, 3.0], 0.57, 51.0),
             ("New York Giants", "Houston", 10.0, [9.0, 11.5, 8.2, 10.6, 12.0, 9.4, 7.7], 0.83, 41.0),
             ("Indianapolis", "Minnesota", 6.5, [5.1, 7.9, 6.0, 4.4, 8.8, 6.3, 5.5], 0.73, 45.5),
             ("Miami", "New York Jets", 5.0, [2.2, 6.4, 4.8, 3.9, 7.1, 5.5, 4.0], 0.68, 42.0),
             ("Dallas", "Philadelphia", -4.5, [-3.0, -6.2, -4.1, -5.5, -2.4, -4.9, -3.8], 0.33, 47.5)]

    def _outputs(self):
        from acql import analytics, consistency
        from acql.model import flips, predictions, pricing
        from acql.models import Game, Season, Week
        forecasts = [predictions.Forecast(road, home, market=m, opening=m - 0.5, opinions=tuple(o),
                                          market_win=w, total=t) for road, home, m, o, w, t in self.TEAMS]
        predictions.save(7, forecasts, "golden.csv")
        games = [Game(i + 1, home, road) for i, (road, home, *_r) in enumerate(self.TEAMS)]
        games[4].line, games[4].line_favourite = 10.0, "Houston"       # pool lines on two games
        games[5].line, games[5].line_favourite = 7.0, "Minnesota"
        season = Season(title="ACQL 2026", weeks={7: Week(7, games)})
        prices = pricing.price_week(season, 7, games)
        out = {"prices": [[p.game.index, round(p.margin, 4), round(p.sd, 4), p.prediction().pick,
                           round(p.prediction().chance, 5)] for p in prices],
               "flips": [[f.index, None if f.distance is None else round(f.distance, 3)]
                         for f in flips.week_flips(season, 7)],
               "big_losers": [[b.team, round(b.chance, 5)] for b in pricing.big_loser_candidates(prices)[:5]],
               "sources_agree": [p.kind for p in consistency.week(season, 7)],
               "curve": [round(analytics.cover_chance(s, n), 5) for s, n in ((3, 3), (6.5, 7), (10, 10), (1, 1))]}
        return out

    def test_the_week_comes_out_as_it_did(self):
        import hashlib
        import json
        out = self._outputs()
        text = json.dumps(out, sort_keys=True)
        found = hashlib.sha256(text.encode()).hexdigest()[:16]
        self.assertEqual(found, self.FINGERPRINT,
                         "Pricing changed - picks may differ. If the change was approved, set FINGERPRINT to "
                         f"{found!r}. The new values:\n{json.dumps(out, indent=1)}")


class TestWaterfall(_ScratchStore):
    """A game's steps add up to exactly the margin pricing gives it (#23)."""

    def test_steps_add_up_on_the_golden_week(self):
        from acql.model import pricing, waterfall
        self.TEAMS = TestGoldenWeek.TEAMS
        TestGoldenWeek._outputs(self)
        from acql.models import Game
        games = [Game(i + 1, home, road) for i, (road, home, *_r) in enumerate(TestGoldenWeek.TEAMS)]
        ctx = pricing.context(None, 7, games)
        moved = 0
        for g in games:
            price = pricing.price(g, ctx)
            found = waterfall.steps(g, ctx, price)
            self.assertEqual(found[0].kind, "start")
            self.assertEqual(found[-1].kind, "end")
            self.assertAlmostEqual(waterfall.total(found), price.margin, places=9)
            moved += sum(1 for s in found if s.kind == "move")
        self.assertGreater(moved, len(games))       # the models (and more) show on most games


class TestExplorer(unittest.TestCase):
    """Any slice of the graded games, with its size and a plain interval (#31)."""

    def _games(self):
        from acql.model.autopsy import Graded
        out = []
        for i in range(40):
            out.append(Graded(week=1 + i % 6, index=i, matchup="A @ B", pick="B", chance=0.6, right=i % 3 != 0,
                              margin=3.0, line=2.5 if i % 2 else 7.0, line_pick="B", line_right=i % 4 != 0,
                              pool_line=7.0 if i % 5 == 0 else None, home_pick=i % 2 == 0,
                              divisional=i % 4 == 0, night=i % 8 == 0))
        return out

    def test_slices(self):
        from acql.model import explorer
        games = self._games()
        everything = explorer.ask(games, explorer.Slice())
        self.assertEqual(everything.games, 40)
        self.assertAlmostEqual(everything.expected, 24.0)
        close = explorer.ask(games, explorer.Slice(band="Pick'em to 3"))
        self.assertEqual(close.games, 20)
        both = explorer.ask(games, explorer.Slice(band="7 to 9.5", side="Home pick", division="Yes"))
        self.assertTrue(all(g % 4 == 0 for g in [x.index for x in games if explorer.matches(x, both.slice)]))
        self.assertTrue(both.few)
        self.assertIn("Too few games", both.words())
        self.assertLess(both.low, both.rate)
        self.assertGreater(both.high, both.rate)
        self.assertEqual(explorer.ask([], explorer.Slice()).words(), "No graded games match.")


class TestLiveCheck(unittest.TestCase):
    """The live rule replayed from scoring plays: scores rebuilt at each checkpoint, bands formed (#96)."""

    def test_score_at_and_study(self):
        from acql.feeds.timeline import Moment
        from acql.model import livecheck as L
        moments = [Moment(None, 0, 7, 1, "10:00", "TD"), Moment(None, 3, 7, 2, "0:30", "FG"),
                   Moment(None, 3, 14, 4, "5:00", "TD")]
        self.assertEqual(L.score_at(moments, 1.0), (0, 0))
        self.assertEqual(L.score_at(moments, 0.75), (7, 0))
        self.assertEqual(L.score_at(moments, 0.5), (7, 3))
        self.assertEqual(L.score_at(moments, 0.1), (7, 3))      # the late TD came with 5:00 left, after 6:00
        self.assertGreater(L.chance_at(14, 3, 0.05, 0.0, 13.5), 0.97)
        self.assertAlmostEqual(L.chance_at(0, 0, 1.0, 0.0, 13.5), 0.5, places=2)
        calls = [L.Call(1, "x", 0.5, 0.9, True), L.Call(1, "x", 0.5, 0.9, False), L.Call(1, "y", 0.5, 0.1, False)]
        st = L.study(None, (2, calls))
        self.assertAlmostEqual(st.brier[0.5][0], (0.01 + 0.81 + 0.01) / 3)
        top = [b for b in st.bands if b.low == 0.8][0]
        self.assertEqual((top.calls, top.came), (2, 0.5))


class TestDeadlineView(_ScratchStore):
    """A past week from what was written before its deadline, with gaps said plainly (#7)."""

    def test_a_week_as_it_stood(self):
        from acql import store
        from acql.model import deadline, known
        from acql.models import Game, Season, Week
        from acql.pool import cards
        g1 = Game(1, "Chicago", "Minnesota", winner="Chicago", home_points=24, away_points=17)
        g2 = Game(2, "Dallas", "Philadelphia")
        season = Season(title="ACQL 2026", weeks={6: Week(6, [g1, g2])})
        store.put(known.KIND, known._key(2026, 6), {"at": 1000.0, "frozen": True, "games": {
            "1": {"matchup": "Minnesota @ Chicago", "line": "Chicago -3", "pool_line": ""},
            "2": {"matchup": "Philadelphia @ Dallas", "line": "Philadelphia -4.5", "pool_line": "Philadelphia by 7"}}})
        cards._put(6, "model", cards.Snapshot(at=2000.0, take={1: "Chicago", 2: "Dallas"}, chance={1: 0.6, 2: 0.52},
                                              label={1: "Minnesota @ Chicago", 2: "Philadelphia @ Dallas"}))
        from unittest import mock
        from acql.feeds import results
        with mock.patch.object(results, "season_year", return_value=2026):
            self.assertEqual(deadline.weeks(season), [6])
            v = deadline.view(season, 6)
        self.assertTrue(v.frozen)
        self.assertEqual([(r.index, r.line, r.deadeye, r.result) for r in v.rows],
                         [(1, "Chicago -3", "Chicago", "Chicago"), (2, "Philadelphia -4.5", "Dallas", "")])
        self.assertEqual(v.rows[1].pool_line, "Philadelphia by 7")
        self.assertEqual(v.caveats(), ["No card of yours is on record for this week."])


class TestFinalRecheck(_ScratchStore):
    """Finals are asked about again at 15 minutes and 12 hours; a changed score replaces the kept one (#17)."""

    def test_corrections_are_kept_once_each(self):
        from acql.feeds import recheck, results
        S = results.Score
        kick = 1_000_000.0
        results.save(2026, 6, [S("min", "chi", 17, 24, "post", True, kick), S("phi", "dal", 20, 20, "post", True, kick)])
        now = kick + 4 * 3600                      # first seen final here
        self.assertEqual(recheck.run(2026, 6, now=now, fetch=lambda y, w: []), [])   # tracked, nothing owed yet
        self.assertFalse(recheck.due(2026, [6], now=now + 60))
        self.assertTrue(recheck.due(2026, [6], now=now + 16 * 60))
        corrected = [S("min", "chi", 17, 27, "post", True, kick, detail="Final"), S("phi", "dal", 20, 20, "post", True, kick)]
        made = recheck.run(2026, 6, now=now + 16 * 60, fetch=lambda y, w: corrected)
        self.assertEqual([(c.game, c.before, c.after) for c in made], [("MIN @ CHI", (17, 24), (17, 27))])
        kept = {s.home: (s.away_points, s.home_points) for s in results.load(2026, 6)}
        self.assertEqual(kept["chi"], (17, 27))
        self.assertEqual(recheck.run(2026, 6, now=now + 20 * 60, fetch=lambda y, w: corrected), [])   # not owed again
        self.assertTrue(recheck.due(2026, [6], now=now + 12 * 3600 + 1))
        self.assertEqual(recheck.run(2026, 6, now=now + 12 * 3600 + 1, fetch=lambda y, w: corrected), [])
        self.assertFalse(recheck.due(2026, [6], now=now + 99 * 3600))          # both looks done
        self.assertEqual(len(recheck.corrections(2026, [6])), 1)


class TestLoserWatch(unittest.TestCase):
    """A Big Loser pick against the 15-point line while it is played (#15)."""

    def test_words(self):
        from types import SimpleNamespace as NS
        from acql.pool import gameday as GD
        live = NS(state="in", margin=-11, blowout_home=0.31, blowout_away=0.0)       # home down 11
        self.assertEqual(GD.loser_watch(live, GD.HOME), f"needs 4 more · 31% to finish {GD.BLOWOUT}+")
        past = NS(state="in", margin=17, blowout_home=0.0, blowout_away=0.82)        # away down 17
        self.assertEqual(GD.loser_watch(past, GD.AWAY), f"{GD.BLOWOUT}+ now, must hold · 82% to finish {GD.BLOWOUT}+")
        self.assertEqual(GD.loser_watch(NS(state="pre", margin=None, blowout_home=0.12, blowout_away=0.0), GD.HOME),
                         f"12% to lose by {GD.BLOWOUT}+")
        self.assertEqual(GD.loser_watch(NS(state="final", margin=-20, blowout_home=1.0, blowout_away=0.0), GD.HOME), "")


class TestBudgetForecast(unittest.TestCase):
    """The month's pace projected to its end (#42)."""

    def test_pace(self):
        import datetime as dt
        from acql.feeds import odds
        day = 86400
        start = dt.datetime(2026, 10, 1, 12).timestamp()
        log = [[start, 500], [start + 10 * day, 440], [start + 20 * day, 380]]
        now = start + 20 * day
        projected, end, per_day = odds.budget_forecast(now, log)
        self.assertEqual(per_day, 6)
        self.assertEqual(dt.datetime.fromtimestamp(end).day, 31)
        self.assertTrue(310 <= projected <= 320)
        self.assertIsNone(odds.budget_forecast(now, [[now - 60, 400], [now, 397]]))       # under a day of readings
        self.assertIsNone(odds.budget_forecast(now, [[start - 40 * day, 500], [now, 380]]))   # last month's reading


class TestRulings(_ScratchStore):
    """One ruling per week that doesn't add up, written down once (#57)."""

    def test_put_get_clear(self):
        from acql.pool import rulings
        self.assertIsNone(rulings.get(2026, 3))
        self.assertEqual(rulings.words(None), "Not ruled yet")
        rulings.put(2026, 3, "official", "commissioner's hand count", now=1_800_000_000)
        found = rulings.get(2026, 3)
        self.assertEqual(found["ruling"], "official")
        self.assertIn("Official totals stand (ruled 2027-01-15) - commissioner's hand count", rulings.words(found))
        rulings.put(2026, 3, "")
        self.assertIsNone(rulings.get(2026, 3))
        with self.assertRaises(ValueError):
            rulings.put(2026, 3, "maybe")


class TestFavouriteGrid(unittest.TestCase):
    """How the pool takes favourites, by line size and week (#49)."""

    def test_grid(self):
        from acql.pool import crowd
        def ob(week, points, share, weight=1.0):
            sit = crowd.Situation("A", "B", points, True, None)
            return crowd.Observation(sit, share, 30, week, weight)
        bands, weeks, values, counts = crowd.favourite_grid(
            [ob(1, 2.5, 0.6), ob(1, 1.0, 0.8), ob(1, 7.0, 0.9), ob(2, 10.5, 0.75), ob(2, 3.0, 0.5, weight=0.4)])
        self.assertEqual(weeks, [1, 2])
        self.assertEqual(values[0][0], 70.0)               # under 3, week 1: (60 + 80) / 2
        self.assertEqual(counts[0][0], 2)
        self.assertEqual(values[3][0], 90.0)               # 7 to 9.5
        self.assertEqual(values[4][1], 75.0)               # 10 or more, week 2
        self.assertIsNone(values[1][1])                    # an older season's sheet isn't this pool's habit


class TestDriftOutlook(unittest.TestCase):
    """Only games still at their opening line; the expected move and whether it turns the pick (#21)."""

    def test_at_opening(self):
        from types import SimpleNamespace as NS
        from acql.model import drift
        self.assertTrue(drift.at_opening(NS(opening=-3.0, market=-3.0)))
        self.assertTrue(drift.at_opening(NS(opening=None, market=-3.0)))
        self.assertFalse(drift.at_opening(NS(opening=-3.0, market=-3.5)))

    def test_read_needs_the_file(self):
        from acql.model import drift
        folder = Path(tempfile.mkdtemp(prefix="acql-drift-"))
        self.assertIsNone(drift.read(folder / "missing.json"))
        (folder / "d.json").write_text('{"from": {"opening": {"move": 0.4, "games": 2000}}}', encoding="utf-8")
        self.assertEqual(drift.read(folder / "d.json")["from"]["opening"]["move"], 0.4)


class TestClvShadow(unittest.TestCase):
    """Each model weighted by how often the line later moved its way, on earlier weeks only (#27)."""

    def _r(self, week, named, opening, close, actual):
        from acql.model.computers import Reading
        return Reading(week, 1, "A @ B", close, None, None, tuple(sorted(named.values())), named=named,
                       actual=actual, opening=opening, close=close)

    def test_hits_and_weights(self):
        from acql.model import computers as C
        r = self._r(1, {"good": 5.0, "bad": 1.0}, 3.0, 4.0, 7.0)
        self.assertTrue(C.clv_hit(5.0, r))
        self.assertFalse(C.clv_hit(1.0, r))
        self.assertIsNone(C.clv_hit(5.0, self._r(1, {}, 3.0, 3.0, 7.0)))        # the line didn't move
        w = C.clv_weights([r, r, r])
        self.assertGreater(w["good"], w["bad"])
        self.assertEqual(w["bad"], C.CLV_FLOOR)


class TestSeasonGoal(unittest.TestCase):
    """The goal suggests a card with its reason, and never one until it's set (#90)."""

    def test_suggestions(self):
        from acql.pool import briefing, goal
        self.assertIsNone(goal.suggest("", 3.0, 10))
        self.assertEqual(goal.suggest("money", 3.0, 10).style, briefing.STYLE_MONEY)
        self.assertEqual(goal.suggest("top3", 3.0, 10).style, briefing.STYLE_STRAIGHT)
        self.assertEqual(goal.suggest("title", 7.0, 12).style, briefing.STYLE_STRAIGHT)    # time to catch up
        self.assertEqual(goal.suggest("title", 4.0, 5).style, briefing.STYLE_WIN)          # far back, late
        self.assertEqual(goal.suggest("title", 1.0, 3).style, briefing.STYLE_STRAIGHT)     # close, late


class TestPrePublishCheck(unittest.TestCase):
    """No page with a merge, a placeholder, a private word or a broken script is uploaded."""

    def test_problems(self):
        from acql.site import webpublish
        ok = lambda page: webpublish.page_problems(page, parse=lambda _p: [])
        self.assertEqual(ok("<html><p>fine</p></html>"), [])
        marked = "\n".join(["<p>a</p>", "<" * 7 + " HEAD", "x", "=" * 7, "y", ">" * 7 + " main", ""])
        self.assertIn("merge-conflict", ok(marked)[0])
        self.assertIn("__POOL_DATA__", ok("<script>__POOL_DATA__</script>")[0])
        self.assertIn("Deadeye", ok("<p>Deadeye says</p>")[0])
        self.assertEqual(webpublish.page_problems("<p>x</p>", parse=lambda _p: ["broken"]), ["broken"])

    def test_a_bad_page_is_never_uploaded(self):
        from acql.site import webpublish
        settings = webpublish.WebSettings(user="u", repo="r", sealed_token="x")
        calls = []
        with self.assertRaises(webpublish.WebPublishError) as caught:
            webpublish.publish(settings, "<p>Deadeye</p>", opener=lambda *a, **k: calls.append(a),
                               check=lambda page: webpublish.page_problems(page, parse=lambda _p: []))
        self.assertIn("still shows the last good page", str(caught.exception))
        self.assertEqual(calls, [])


class TestRivalsAndLevers(unittest.TestCase):
    """Rival lens: only the games two cards differ on (#16); levers: a touchdown's worth (#14)."""

    def _board(self):
        from types import SimpleNamespace as NS
        from acql.pool import gameday as GD
        def st(i, state, leader, p_home):
            return NS(index=i, state=state, leader=leader, p_home=p_home, p_away=1 - p_home)
        states = [st(1, "final", GD.HOME, 1.0), st(2, "in", GD.AWAY, 0.4), st(3, "pre", GD.NOBODY, 0.5),
                  st(4, "in", GD.HOME, 0.7)]
        you = NS(card=NS(name="Phelpy", you=True, picks={1: GD.HOME, 2: GD.HOME, 3: GD.AWAY, 4: GD.HOME}),
                 total_now=5)
        jules = NS(card=NS(name="Jules", you=False, picks={1: GD.AWAY, 2: GD.AWAY, 3: GD.AWAY, 4: GD.HOME}),
                   total_now=6)
        return NS(you=you, standings=[jules, you], states=states, roots=[])

    def test_rivals(self):
        from acql.pool import gameday as GD
        found = GD.rivals(self._board(), ["Jules", "Nobody"])
        self.assertEqual(len(found), 1)
        r = found[0]
        self.assertEqual([s.index for s, _m, _t in r.differ], [1, 2])
        self.assertEqual(r.ahead_now, 0)          # you won game 1, they lead game 2
        self.assertEqual(r.total_gap, -1)

    def test_no_levers_without_live_stakes(self):
        from acql.pool import gameday as GD
        self.assertEqual(GD.levers(self._board()), [])


class TestSurvivorPath(unittest.TestCase):
    """One team a week, each once, surest pairing first (#56)."""

    def test_choose(self):
        from acql.pool import survivorpath as sp
        cells = {(7, "buf"): (0.80, "Miami", True), (7, "kc"): (0.78, "Denver", True),
                 (8, "buf"): (0.85, "Jets", True), (8, "kc"): (0.60, "Raiders", False),
                 (8, "phi"): (0.70, "Dallas", True)}
        steps = sp.choose(cells, [7, 8])
        self.assertEqual([s.week for s in steps], [7, 8])
        picked = {s.week: s.team for s in steps}
        from acql.model import predictions
        self.assertEqual(picked[8], predictions.display_team("buf"))     # the surest cell goes first
        self.assertEqual(picked[7], predictions.display_team("kc"))      # so week 7 takes the next team
        self.assertEqual(steps[1].alternate, predictions.display_team("phi"))
        self.assertAlmostEqual(sp.Path("x", steps).survive_all, 0.85 * 0.78)


class TestVisualDiff(unittest.TestCase):
    """Changed regions boxed, scores per page, baselines kept per version (#4)."""

    def test_compare_and_baselines(self):
        from PIL import Image, ImageDraw
        from acql.ui import visualdiff
        root = Path(tempfile.mkdtemp(prefix="acql-vd-"))
        old_run, new_run = root / "r1", root / "r2"
        old_run.mkdir()
        new_run.mkdir()
        base = Image.new("RGB", (200, 120), (20, 20, 20))
        base.save(old_run / "01-Results-Standings.png")
        changed = base.copy()
        ImageDraw.Draw(changed).rectangle((50, 40, 90, 70), fill=(240, 200, 60))
        changed.save(new_run / "03-Results-Standings.png")
        self.assertEqual(visualdiff.run(old_run, "2.18.2")[0], None)       # first run: nothing to compare
        against, found = visualdiff.run(new_run, "2.19.0")
        self.assertEqual(against, "2.18.2")
        self.assertEqual(found[0].key, "Results-Standings.png")
        self.assertGreater(found[0].score, 0)
        self.assertEqual(len(found[0].boxes), 1)
        self.assertTrue((new_run / "diff-03-Results-Standings.png").exists())
        self.assertIn("1 of 1 pages changed", "\n".join(visualdiff.lines(against, found)))


class TestSheetChanges(_ScratchStore):
    """A sheet sent again: each changed pick, card and side pool written down (#43)."""

    def test_diff_and_note(self):
        from acql.pool import picks, sheetchanges
        g = lambda picks_: picks.GamePicks("Minnesota", "Chicago", "", picks_)
        old = picks.WeekPicks(5, [g({"Ann": "Chicago", "Bo": "Minnesota"})], suicide={"Ann": "Buffalo"})
        new = picks.WeekPicks(5, [g({"Ann": "Minnesota", "Bo": "Minnesota", "Cy": "Chicago"})],
                              suicide={"Ann": "Detroit"})
        found = sheetchanges.note(5, old, new, "week5.xlsx", now=1000.0)
        self.assertIn("Cy: card added", found)
        self.assertIn("Ann: Minnesota @ Chicago changed from Chicago to Minnesota", found)
        self.assertIn("Ann: suicide pick changed from Buffalo to Detroit", found)
        self.assertEqual(len(found), 3)
        self.assertEqual(sheetchanges.note(5, new, new), [])
        self.assertEqual(sheetchanges.recent([5], now=2000.0)[0][3], found)


class TestChanceTrail(_ScratchStore):
    """Ghost marks: a chance from about a day ago, for the same side (#9)."""

    def test_ghosts(self):
        from acql.pool import chancetrail as ct
        day = 86400
        self.assertTrue(ct.record(7, {1: ("Chicago", 0.55), 2: ("Dallas", 0.60)}, now=1_000_000))
        self.assertFalse(ct.record(7, {1: ("Chicago", 0.56)}, now=1_000_000 + 60))     # within the hour
        now = 1_000_000 + day
        found = ct.ghosts(7, {1: ("Chicago", 0.61), 2: ("Philadelphia", 0.41)}, now=now)
        self.assertEqual(found, {1: 0.55})        # game 2 changed side: 1 - 0.60 = 0.40, under 3 points
        self.assertEqual(ct.ghosts(7, {1: ("Chicago", 0.61)}, now=1_000_000 + 3600), {})   # nothing a day old


class TestTwins(unittest.TestCase):
    """The most alike past games and how they ended, beside what the line says (#37)."""

    def test_find(self):
        from acql.model import twins
        past = tuple(twins.Past(2015, 5, f"g{i}", 3.0 + (i % 3) * 0.5, 44.0, 0.0, i % 2 == 0, m)
                     for i, m in enumerate([7, -3, 3, 10, 1, -7, 17, 4, 2, 21] * 4))
        t = twins.find(3.0, 44.0, 5, 0.0, True, past=past, count=10)
        self.assertEqual(t.count, 10)
        self.assertTrue(all(g.fav_spread == 3.0 for g in t.games[:5]))
        self.assertEqual(t.fav_won, sum(1 for g in t.games if g.fav_margin > 0))
        self.assertTrue(0.5 < t.line_win < 0.7)


class TestShortGrids(unittest.TestCase):
    """A workbook week's grid short of a coach's official count is filled from the sheet, never past it."""

    def test_fill(self):
        from unittest import mock
        from pathlib import Path as _P
        from acql.models import Game, PlayerWeek, Season, Week
        from acql.pool import picks
        g1 = Game(1, "Buffalo", "Los Angeles Chargers", winner="Buffalo", home_points=24, away_points=16,
                  line=8.0, line_favourite="Buffalo")
        g2 = Game(2, "Chicago", "Philadelphia", winner="Chicago", home_points=27, away_points=7)
        g3 = Game(3, "Denver", "Los Angeles Rams", winner="Denver", home_points=30, away_points=26)
        week = Week(3, [g1, g2, g3])
        week.source = _P("workbook.xlsx")
        kit = PlayerWeek(player="Kit", week=3, correct_picks={3: "Denver"}, regular_wins=2)   # official: 2 right
        ann = PlayerWeek(player="Ann", week=3, correct_picks={3: "Denver"}, regular_wins=1)
        week.lines = {"kit": kit, "ann": ann}
        season = Season(title="ACQL 2026", weeks={3: week})
        gp = lambda away, home, picks_: picks.GamePicks(away, home, "", picks_)
        sheet = picks.WeekPicks(3, [gp("Los Angeles Chargers", "Buffalo", {"Kit": "Buffalo", "Ann": "Buffalo"}),
                                    gp("Philadelphia", "Chicago", {"Kit": "Chicago", "Ann": "Chicago"}),
                                    gp("Los Angeles Rams", "Denver", {"Kit": "Denver", "Ann": "Denver"})])
        with mock.patch.object(picks, "load", return_value=sheet):
            self.assertEqual(picks.fill_short_grids(season), 1)
        self.assertEqual(sorted(kit.correct_picks), [1, 3])     # one filled: the official count is 2
        self.assertEqual(sorted(ann.correct_picks), [3])        # already at the official count


class TestRoundFour(_ScratchStore):
    """Seal, source health, story checker, publish times, chance words, ritual look-back, alerts."""

    def test_seal_keeps_first_fingerprint(self):
        from unittest import mock
        from acql.pool import seal
        with mock.patch.object(seal, "fingerprint", return_value="abcd"):
            self.assertIsNone(seal.seal(2026, 7, lock=100.0, now=50.0))          # before the deadline
            first = seal.seal(2026, 7, lock=100.0, now=200.0)
        self.assertEqual((first.code, first.matches, first.label), ("abcd", True, "W7 · abcd"))
        with mock.patch.object(seal, "fingerprint", return_value="ffff"):
            later = seal.seal(2026, 7, lock=100.0, now=900.0)
        self.assertEqual((later.code, later.matches), ("abcd", False))           # rewritten records show

    def test_source_health(self):
        from acql import sourcehealth
        calls = sourcehealth.watched("odds", lambda x: x * 2)
        self.assertEqual(calls(2), 4)
        def bad():
            raise OSError("down")
        with self.assertRaises(OSError):
            sourcehealth.watched("odds", bad)()
        h = sourcehealth.health("odds")
        self.assertEqual((h.calls, h.good), (2, 1))
        self.assertIn("down", h.last_note)
        self.assertEqual(sourcehealth.watched("odds", lambda opener=None: 1)(opener=object()), 1)
        self.assertEqual(sourcehealth.health("odds").calls, 2)                   # test stand-ins not recorded

    def test_story_checker(self):
        from acql.site import storyteller
        sheet = {"week": 6, "winners": ["Ann"], "top_score": 13}
        self.assertEqual(storyteller.problems("Ann won week 6 with 13.", sheet), [])
        self.assertTrue(storyteller.problems("Ann won with 14.", sheet))
        self.assertTrue(storyteller.problems("Deadeye liked it.", sheet))

    def test_publish_window(self):
        import time as _t
        from acql import publishtimes
        base = _t.mktime((2026, 10, 6, 10, 0, 0, 0, 0, -1))                     # a Tuesday 10 am
        for w, extra in ((1, 0), (2, 7 * 86400 + 3600), (3, 14 * 86400 + 7200)):
            self.assertTrue(publishtimes.note("predictions", w, 2026, base + extra))
        self.assertFalse(publishtimes.note("predictions", 1, 2026, base))
        found = publishtimes.window("predictions", 2026)
        self.assertEqual((found.weekday, found.weeks), ("Tue", 3))
        self.assertIn("usually Tue 10 am-12 pm", found.words())

    def test_alert_preferences(self):
        from acql.site import cardmail
        self.assertEqual(cardmail.alert_channel("watch"), "email")
        cardmail.set_alert_channel("watch", "none")
        self.assertEqual(cardmail.send_note("s", "t", event="watch"), [])        # muted: nothing sent

    def test_words_and_look_back(self):
        from acql.pool import ritual
        from acql.ui import extras
        self.assertEqual(extras.chance_words(0.5), "a coin flip")
        self.assertEqual(extras.chance_words(0.6), "likely, about 3 in 5")
        self.assertEqual(extras.chance_words(0.25), "unlikely, about 1 in 4")
        self.assertIn("too few", ritual.LookBack(17, 5, 12).words())
        self.assertIn("6 more points", ritual.LookBack(40, 10, 16).words())


class TestClinchedAndOut(unittest.TestCase):
    """Clinched and out are exact: from what is banked and what each card can still reach."""

    def test_marks(self):
        from acql.pool import gameday as GD
        from types import SimpleNamespace as NS

        def standing(key, banked, most, in_field=True):
            card = GD.Card(key, key, in_field=in_field)
            return GD.Standing(card, banked, 0, banked, 0, most, 1, None, None, None)
        field = [standing("a", 12, 13), standing("b", 9, 12), standing("c", 8, 11), standing("d", 10, 12)]
        bd = NS(standings=field, you=None, ghost=standing("g", 7, 9, in_field=False),
                states=[NS(started=True)])
        marks = GD.certainty(bd)
        self.assertEqual(marks.get("c"), GD.OUT)            # 11 can't reach a's banked 12
        self.assertEqual(marks.get("g"), GD.OUT)
        self.assertEqual(marks.get("a"), GD.SHARE)          # b and d can at best tie its 12
        self.assertNotIn("b", marks)
        field[1] = standing("b", 9, 11)
        field[3] = standing("d", 10, 11)
        self.assertEqual(GD.certainty(bd).get("a"), GD.CLINCHED)   # 12 banked beats every 11
        field[1] = standing("b", 9, 13)
        self.assertNotIn("a", GD.certainty(bd))                    # b can still pass it
        bd.states = [NS(started=False)]
        self.assertEqual(GD.certainty(bd), {})                     # nothing before kickoff


class TestWebsitePieces(unittest.TestCase):
    """The pool page is kept as pieces that join into one page, and its helpers are tested in a browser (#83)."""

    def test_pieces_join_in_order_into_one_page(self):
        from acql.ui import poolpage
        names = [p.name for p in poolpage.parts()]
        self.assertEqual(names[0], "000-page-top.html")
        self.assertEqual(names[-1], "999-page-end.html")
        page = poolpage.template()
        self.assertEqual(page.count('<script id="pool-app">'), 1)
        self.assertTrue(page.rstrip().endswith("})();\n</script>"))
        self.assertIn("__POOL_DATA__", page)

    def test_site_fingerprint_works_and_follows_the_data(self):
        # Every automatic website update starts here: when it raised (Oct 23, a
        # local name hiding template()), the site stopped publishing.
        from acql.ui import poolpage
        empty = poolpage.site_fingerprint({})
        self.assertIsInstance(empty, str)
        self.assertTrue(empty)
        self.assertEqual(empty, poolpage.site_fingerprint({}))
        self.assertNotEqual(empty, poolpage.site_fingerprint({"announcement": "Week 8 is up"}))

    def test_the_page_helpers(self):
        sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "packaging"))
        import jsrun
        from acql.ui import poolpage
        helpers = (poolpage.PARTS / "020-small-helpers.js").read_text(encoding="utf-8")
        try:
            got = jsrun.run(helpers, "[money(1234.5), money(-3), dollars(-12.4), dollars(0), plural(1, 'win'), "
                                     "plural(2, 'win'), count(3, 'game'), count(12, 'game'), cap('abc'), "
                                     "listNames(['A']), listNames(['A', 'B', 'C']), listNames(['A', 'B', 'C', 'D']), "
                                     "[1, 2, 3, 4, 11, 12, 13, 21, 22, 101, 111].map(ordinal).join(' ')]")
        except jsrun.NoBrowser:
            self.skipTest("no Chrome or Edge to run the page's script in")
        self.assertEqual(got, ["$1,234.50", "−$3.00", "−$12", "$0", "1 win", "2 wins", "three games",
                               "12 games", "Abc", "A", "A, B & C", "A, B & 2 more",
                               "1st 2nd 3rd 4th 11th 12th 13th 21st 22nd 101st 111th"])


class TestWebsiteCurve(unittest.TestCase):
    """The website prices games on the dashboard's own curve: same chances, both surfaces."""

    def test_page_data_carries_the_curve(self):
        from acql import analytics
        from acql.ui import poolpage
        found = poolpage.curve()
        self.assertEqual(len(found["keys"]), len(analytics.MARGIN_RANGE))
        self.assertEqual((found["sd"], found["low"]), (analytics.SPREAD_SD, analytics.MARGIN_RANGE[0]))

    def test_the_page_and_the_dashboard_agree(self):
        import html
        import json
        import re
        import subprocess
        sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "packaging"))
        import guard
        browser = guard.chrome()
        if browser is None:
            self.skipTest("no Chrome or Edge to run the page's script in")
        from acql import analytics
        from acql.pool import gameday
        from acql.ui import poolpage
        page = poolpage.template()
        engine = page[page.index("  var SQRT2"):page.index("  function oldOutlook")]
        phi = page[page.index("  function phi(z) {"):page.index("  // A game level at the end of regulation")]
        ks = [1, 0, 3, 4, 7, 8, 15, -14]
        pre = [-10.5, -3.0, 0.0, 2.5, 6.5, 13.5]
        live = [(-3, 0.5, 3.0), (4, 0.25, 3.0), (0, 0.05, -6.5), (14, 0.75, 0.0)]
        probe = (f"<!doctype html><title>x</title><script>var DATA = {{curve: {json.dumps(poolpage.curve())}}};"
                 f"var OT_TIE = {gameday.OT_TIE}, OT_MARGIN = {gameday.OT_MARGIN};\n{phi}\n{engine}\n"
                 "function at(d, k) { var i = Math.max(0, k - d.lo); return i >= d.c.length ? 0 : 1 - (i ? d.c[i - 1] : 0); }"
                 f"var ks = {json.dumps(ks)};"
                 f"document.title = JSON.stringify([{json.dumps(pre)}.map(function (s) {{ var d = preDist(s); return ks.map(function (k) {{ return at(d, k); }}); }}),"
                 f"{json.dumps(live)}.map(function (c) {{ var d = liveDist(c[0], c[1], c[2]); return ks.map(function (k) {{ return at(d, k); }}); }})]);"
                 "</script>")
        folder = Path(tempfile.mkdtemp(prefix="acql-test-"))
        try:
            (folder / "p.html").write_text(probe, encoding="utf-8")
            done = subprocess.run([browser, "--headless=new", "--disable-gpu", f"--user-data-dir={folder / 'u'}",
                                   "--dump-dom", (folder / "p.html").as_uri()], capture_output=True, timeout=90)
        finally:
            shutil.rmtree(folder, ignore_errors=True)
        js_pre, js_live = json.loads(html.unescape(re.search(rb"<title>(.*?)</title>", done.stdout).group(1).decode()))
        for s, row in zip(pre, js_pre):
            want = [sum(p for d, p in analytics.margin_distribution(s) if d >= k) for k in ks]
            for a, b in zip(row, want):
                self.assertAlmostEqual(a, b, places=4, msg=f"before kickoff, spread {s}")
        for (now, left, exp), row in zip(live, js_live):
            want = [gameday.outlook(now, left, exp, analytics.SPREAD_SD, started=True).at_least(k) for k in ks]
            for a, b in zip(row, want):
                self.assertAlmostEqual(a, b, places=4, msg=f"in play {now}, {left}, {exp}")


class TestDecisionJournal(_ScratchStore):
    """Games taken against Deadeye, with a reason each, scored by reason once graded."""

    def _season(self):
        from acql.models import Game, Season, Week
        from acql.pool import cards, picks
        games = [Game(1, "Chicago", "New England", winner="Chicago"),
                 Game(2, "Philadelphia", "Dallas", winner="Dallas"),
                 Game(3, "Houston", "Miami", winner="Houston"),
                 Game(4, "Detroit", "Green Bay")]                       # still to play
        cards._put(7, "model", cards.Snapshot(at=1.0, take={1: "Chicago", 2: "Philadelphia", 3: "Houston", 4: "Detroit"},
                                              chance={}, label={}))
        sheet = picks.WeekPicks(7, games=[
            picks.GamePicks("New England", "Chicago", "", {"Phelpy": "NEW ENGLAND"}),
            picks.GamePicks("Dallas", "Philadelphia", "", {"Phelpy": "dallas"}),
            picks.GamePicks("Miami", "Houston", "", {"Phelpy": "Houston"}),
            picks.GamePicks("Green Bay", "Detroit", "", {"Phelpy": "Green Bay"})])
        original = picks.load
        picks.load = lambda week, **_: sheet if week == 7 else None
        self.addCleanup(lambda: setattr(picks, "load", original))
        return Season(title="ACQL 2026", weeks={7: Week(7, games)})

    def test_overrides_reasons_and_tallies(self):
        from acql.pool import journal
        from acql.version import MODEL_NAME
        season = self._season()
        found = journal.overrides(season, "Phelpy")
        self.assertEqual([(o.index, o.yours, o.outcome) for o in found],
                         [(1, "New England", MODEL_NAME), (2, "Dallas", "you"), (4, "Green Bay", "to play")])
        journal.give_reason(7, 2, "injury", "  their QB  was out ")
        journal.give_reason(7, 1, "gut")
        found = journal.overrides(season, "Phelpy")
        self.assertEqual((found[1].tag, found[1].note), ("injury", "their QB was out"))
        by = {t.tag: t for t in journal.tallies(found)}
        self.assertEqual((by["injury"].you, by["gut"].deadeye, by["all"].games, by["all"].decided), (1, 1, 3, 2))
        journal.give_reason(7, 1, "", "")                           # nothing given: forgotten
        self.assertEqual(journal.overrides(season, "Phelpy")[0].tag, "")


class TestFlightRecorder(_ScratchStore):
    """A live look adds a point only when a score, a game's state or your chance moved."""

    def test_points_only_on_change(self):
        from types import SimpleNamespace as NS
        from acql.models import Season
        from acql.pool import recorder
        season = Season(title="ACQL 2026")

        def board(home, chance, live=True):
            st = NS(index=1, state="in" if live else "final", home_points=home, away_points=7, left=0.5)
            you = NS(chance=chance, place=3, total_now=8)
            return NS(states=[st], live=[st] if live else [], you=you)
        self.assertTrue(recorder.record(season, 7, board=board(10, 0.20), now=1.0))
        self.assertFalse(recorder.record(season, 7, board=board(10, 0.202), now=2.0))   # noise
        self.assertTrue(recorder.record(season, 7, board=board(17, 0.31), now=3.0))     # a touchdown
        self.assertFalse(recorder.record(season, 7, board=board(17, 0.5, live=False), now=4.0))  # no game on
        kept = recorder.trace(2026, 7)
        self.assertEqual([p["chance"] for p in kept], [0.2, 0.31])


class TestSinceLastHere(unittest.TestCase):
    """The opening digest reads the log back to the last visit: problems and quiet days first."""

    def test_digest_from_a_log(self):
        from acql import digest
        folder = Path(tempfile.mkdtemp(prefix="acql-test-"))
        try:
            log = folder / "acql-log.txt"
            log.write_text("\n".join([
                "[2026-10-16 03:00:00] headless: loaded 35 players, week 5; ESPN results for week(s) [5]",
                "[2026-10-18 03:00:00] headless: loaded 35 players, week 5; ESPN results for week(s) [6]",
                "[2026-10-18 03:00:10] headless: website updated: https://example.org/",
                "[2026-10-18 09:00:00] headless: sunday brief not sent: Couldn't reach Gmail",
                "[2026-10-18 09:10:00] headless: note: Week 6 is not covered",
                "[2026-10-18 10:00:00] slow: Game Day took 900 ms to refresh",
            ]), encoding="utf-8")
            since = time.mktime(time.strptime("2026-10-15 20:00", "%Y-%m-%d %H:%M"))
            now = time.mktime(time.strptime("2026-10-20 18:00", "%Y-%m-%d %H:%M"))
            texts = [i.text for i in digest.gather(since, now, log)]
            self.assertTrue(texts[0].startswith("1 problem in the scheduled runs"))
            self.assertIn("No scheduled run on Sat Oct 17, Mon Oct 19", texts[1])
            self.assertIn("Results read for week 6", texts)
            self.assertTrue(any(t.startswith("The website was updated 1 time") for t in texts))
            self.assertEqual(digest.gather(now - 600, now, log), [])          # back within the hour: nothing
        finally:
            shutil.rmtree(folder, ignore_errors=True)


class TestTrainingCatchUp(_ScratchStore):
    """An overdue training answer starts the weekly run once - as its own process, never an import."""

    def test_overdue_answer_starts_one_run(self):
        from acql import headless
        from acql.model import challengers
        started = []
        original = challengers.read
        try:
            challengers.read = lambda path=None: challengers.Answer(made=1_790_000_000.0)
            now = 1_790_000_000.0 + 9 * 86400
            self.assertTrue(headless._heal_training(now=now, start=started.append))
            self.assertFalse(headless._heal_training(now=now + 3600, start=started.append))  # not again so soon
            self.assertEqual(started, [[sys.executable, "-m", "acql.training", "weekly"]])
            challengers.read = lambda path=None: challengers.Answer(made=now)
            self.assertFalse(headless._heal_training(now=now + 2 * 86400, start=started.append))  # fresh: nothing
        finally:
            challengers.read = original


class TestMeaningfulAlerts(_ScratchStore):
    """Only a pick changing side, or a deadline close with no card saved, is worth an email."""

    def test_only_changed_sides_are_reported(self):
        from acql.pool import alerts
        before = {"1": ["Chicago", 3.0], "2": ["Dallas", -4.5]}
        now = {"1": ["New England", 1.5], "2": ["Dallas", -5.0]}
        turned = alerts.compare_picks(before, now)
        self.assertEqual([(t.index, t.was, t.now) for t in turned], [(1, "Chicago", "New England")])
        self.assertIn("from +3 to +1.5", turned[0].text())

    def test_the_guardian_speaks_once_and_only_for_card_savers(self):
        from acql.feeds import oddslock
        from acql.pool import alerts, cards
        lock = oddslock.lock_time(2026, 8)
        soon = lock - 3600
        self.assertIsNone(alerts.deadline_due(2026, 8, soon))           # never saved a card: stays quiet
        cards._put(7, "mine", cards.Snapshot(at=1.0, take={1: "Chicago"}, chance={1: .6}, label={1: "a @ b"}))
        self.assertIn("due in about 1 hour", alerts.deadline_due(2026, 8, soon))
        self.assertIsNone(alerts.deadline_due(2026, 8, soon + 60))      # once a week
        self.assertIsNone(alerts.deadline_due(2026, 9, oddslock.lock_time(2026, 9) - 5 * 3600))  # too early


class TestComputerModels(_ScratchStore):
    """Each computer model kept by name, and the crowd of them measured against the betting line."""

    FILE = ("lineopen,line,road,home,linesag,linemass,linemidweek,lineavg,linestd,neutral,phwin\n"
            "3,3.5,Tampa Bay,Dallas,4.5,2,3.5,3.3,1.0,0,0.6\n"
            "-2,-2.5,Chicago,Green Bay,-1,,-2.5,-1.7,0.8,0,0.4\n")

    @staticmethod
    def _reading(i, line, actual, named, opinions=()):
        from acql.model import computers
        return computers.Reading(week=3 + i // 16, index=i % 16 + 1, matchup=f"A{i} @ H{i}", line=line,
                                 middle=named.get("good"), apart=1.0 + (i % 2), opinions=tuple(opinions),
                                 named=dict(named), deadeye=line, actual=actual)

    def _season(self, games=40):
        """Made-up games: one model near the result, its copy, one that runs backwards, one that is the line."""
        out = []
        for i in range(games):
            line = float((i * 7) % 15 - 7)
            actual = line + ((i * 5) % 13 - 6)
            view = ((i * 3) % 7 - 3) / 2
            named = {"good": actual + view / 4, "twin": actual + view / 4 + 0.5, "wild": line + 9 + (i % 5),
                     "back": -line + (i % 3), "midweek": line, "odd": line + view}
            out.append(self._reading(i, line, actual, named))
        return out

    def test_the_file_is_read_with_its_names(self):
        from acql.model import computers
        csv_file = self.folder / "nflpredictions week 3.csv"
        csv_file.write_text(self.FILE, encoding="utf-8")
        named = computers.read_named(csv_file)
        # linemidweek is a betting line, not a computer model
        self.assertEqual(named[("tb", "dal")], {"sag": 4.5, "mass": 2.0})
        self.assertEqual(named[("chi", "gb")], {"sag": -1.0})                       # a blank is no opinion
        self.assertEqual(computers.label("sag"), "Sagarin (sag)")
        self.assertEqual(computers.label("zzz"), "ZZZ")

    def test_names_are_kept_only_where_the_file_still_matches(self):
        from acql.model import computers
        from acql import store
        csv_file = self.folder / "nflpredictions week 3.csv"
        csv_file.write_text(self.FILE, encoding="utf-8")
        store.put("forecasts", "3", {"source": str(csv_file), "games": [
            {"road": "Tampa Bay", "home": "Dallas", "opinions": [2.0, 3.5, 4.5]},
            {"road": "Chicago", "home": "Green Bay", "opinions": [-9.0, -2.5]},      # priced from another copy
        ]})
        store.put("forecasts", "4", {"source": str(self.folder / "gone.csv"), "games": []})
        self.assertEqual(computers.keep(), {3: 1})
        kept = store.get(computers.KIND, "3")
        self.assertEqual([(g["road"], g["home"]) for g in kept["games"]], [("tb", "dal")])
        self.assertEqual(kept["games"][0]["named"]["sag"], 4.5)
        stamp = kept["kept"]
        self.assertEqual(computers.keep(), {3: 1})                                   # nothing new: not rewritten
        self.assertEqual(store.get(computers.KIND, "3")["kept"], stamp)

    def test_each_model_is_measured_against_the_line_on_its_own_games(self):
        from acql.model import computers
        found = self._season()
        board = {s.code: s for s in computers.leaderboard(found)}
        self.assertEqual(computers.leaderboard(found)[0].code, "good")
        good, wild, line = board["good"], board["wild"], board["midweek"]
        self.assertEqual((good.games, good.of, good.ranked), (40, 40, True))
        self.assertGreater(good.edge, 0)
        self.assertLess(wild.edge, 0)
        self.assertAlmostEqual(line.edge, 0.0)
        self.assertEqual((line.closer, line.further, line.cover_games), (0, 0, 0))
        self.assertGreater(good.closer, good.further)
        self.assertLessEqual(good.closer + good.further, good.games)
        few = computers.leaderboard(found[:10])
        self.assertFalse(any(s.ranked for s in few))

    def test_the_residual_says_how_much_of_the_models_view_showed_up(self):
        from acql.model import computers
        rows = [self._reading(i, 0.0, 0.5 * x, {"good": x}) for i, x in enumerate([-4, -3, -2, -1, 0, 1, 2, 3, 4, 5])]
        half = computers.residual(rows, share=0.33)
        self.assertAlmostEqual(half.slope, 0.5)
        self.assertEqual(half.verdict, "more")
        flat = computers.residual([self._reading(i, 0.0, 0.0, {"good": x}) for i, x in enumerate(range(-5, 5))], 0.33)
        self.assertAlmostEqual(flat.slope, 0.0)
        self.assertEqual(flat.verdict, "nothing")
        short = computers.residual(rows[:4])
        self.assertIsNone(short.slope)
        self.assertIn("Too few", short.reading)
        noisy = computers.residual(self._season())
        self.assertLess(noisy.low, noisy.slope)
        self.assertIn(str(noisy.games), noisy.reading)

    def test_ways_to_average_are_scored_on_the_same_games(self):
        from acql.model import computers
        opinions = (2.0, 2.5, 3.0, 3.0, 3.5, 4.0, 30.0)          # one wild model
        found = [self._reading(i, 1.0, 3.0, {"good": 3.0}, opinions) for i in range(12)]
        ways = computers.averages(found)
        self.assertEqual(ways[-1].name, computers.LINE_ROW)
        named = {w.name: w for w in ways}
        self.assertLess(named["Median"].miss, named["Plain average"].miss)
        self.assertLess(named[computers.IN_USE].miss, named["Plain average"].miss)
        self.assertTrue(named[computers.IN_USE].in_use)
        self.assertEqual({w.games for w in ways}, {12})
        self.assertEqual(computers.averages([]), [])

    def test_copies_and_bad_columns_are_found(self):
        from acql.model import computers
        found = self._season()
        alike = computers.look_alikes(found)
        self.assertIn(("good", "twin"), [(a, b) for a, b, _, _ in alike.pairs])
        self.assertEqual(alike.crowd, 6)
        self.assertLess(alike.effective, 5)                       # six columns, fewer opinions
        flags = {f.code: f.what for f in computers.checks(found)}
        self.assertEqual(flags["back"], "backwards")
        self.assertEqual(flags["midweek"], "a betting line")
        self.assertEqual(flags["wild"], "far from the rest")
        self.assertNotIn("good", flags)

    def test_a_season_with_no_models_is_an_empty_report(self):
        from acql.model import computers
        report = computers.build(Season())
        self.assertEqual((report.found, report.board, report.slate, report.blend), ([], [], [], None))
        self.assertIsNone(report.residual.slope)


class TestLongRecord(unittest.TestCase):
    """Each computer model's season totals since 2000, and what past form is worth."""

    PAGE = ("<h3>2012 Season Totals</h3>Through 2013-02-04<table border='0' CLASS='results_table'>"
            "<tr><th>Rank</th><th>System</th><th>Pct. Correct</th><th>Against Spread</th><th>Absolute Error</th>"
            "<th>Bias</th><th>Mean Square Error</th><th>games</th><th>suw</th><th>sul</th><th>atsw</th><th>atsl</th></tr>"
            "<tr><td align='center'><font size=2>1</font></td><td><font size=2>Sonny  Moore</font></td>"
            "<td>0.66292</td><td>0.51200</td><td>10.2500</td><td>0.50000</td><td>170.100</td><td>267</td>"
            "<td>177</td><td>90</td><td>128</td><td>122</td></tr>"
            "<tr><td>2</td><td>Line (updated)</td><td>0.65</td><td>0.5</td><td>10.0</td><td>0.1</td><td>160.0</td>"
            "<td>267</td><td>170</td><td>92</td><td>0</td><td>0</td></tr>"
            "<tr><td>not a row</td></tr></table>")

    @staticmethod
    def _row(season, name, miss, games=267, cover=0.5, ats=(130, 130)):
        from acql.model import longrecord
        return longrecord.Row(season, name, 0.64, cover, miss, 0.2, games, ats[0], ats[1])

    def _found(self):
        """2010-2013: ten models that keep their order every season, the line ahead of them all."""
        found = {}
        for year in range(2010, 2014):
            rows = [self._row(year, "Line (updated)", 10.0), self._row(year, "System Median", 10.5),
                    self._row(year, "Computer Adjusted Line", 9.95, ats=(60, 40))]
            for i in range(10):
                name = "Sagarin Rating" if i == 0 else f"Model {i}"
                beat_spread = (i + year) % 2 == 0           # a good year against the spread, then a bad one
                rows.append(self._row(year, name, 10.2 + 0.1 * i, cover=0.56 if beat_spread else 0.46,
                                      ats=(150, 118) if beat_spread else (123, 145)))
            found[year] = rows
        found[2014] = [self._row(2014, "Line (updated)", 9.0, games=60), self._row(2014, "Model 1", 8.0, games=60)]
        return found

    def test_a_season_page_is_read(self):
        from acql.model import longrecord
        rows = longrecord.parse(self.PAGE, 2012)
        self.assertEqual([r.name for r in rows], ["Sonny Moore", "Line (updated)"])
        moore = rows[0]
        self.assertEqual((moore.season, moore.games, moore.ats_won, moore.ats_lost), (2012, 267, 128, 122))
        self.assertAlmostEqual(moore.miss, 10.25)
        self.assertEqual((moore.key, rows[1].key), ("sonnymoore", longrecord.LINE))
        self.assertEqual(longrecord.key_of("Turnover Adj. regressio"), longrecord.key_of("Turnover Adj. regression"))
        folder = Path(tempfile.mkdtemp(prefix="acql-test-"))
        try:
            (folder / "nflresults-2012.html").write_text(self.PAGE, encoding="utf-8")
            self.assertEqual(sorted(longrecord.seasons(folder)), [2012])
            self.assertEqual(longrecord.build(folder).careers, [])        # one season is not a record
            self.assertEqual(longrecord.seasons(folder / "nowhere"), {})
        finally:
            shutil.rmtree(folder, ignore_errors=True)

    def test_each_model_is_set_beside_the_line_over_its_own_seasons(self):
        from acql.model import longrecord
        found = self._found()
        careers = longrecord.careers(found)
        self.assertEqual(len(careers), 10)                   # no lines, no averages, no half season
        best = careers[0]
        self.assertEqual((best.name, best.seasons, best.first, best.last, best.games), ("Sagarin Rating", 4, 2010, 2013, 1068))
        self.assertAlmostEqual(best.edge, -0.2)
        self.assertEqual((best.ahead, best.of, best.current), (0, 4, False))
        self.assertTrue(next(c for c in careers if c.name == "Model 1").current)
        self.assertAlmostEqual(best.cover, (150 + 123 + 150 + 123) / (4 * 268))

    def test_accuracy_carries_over_and_the_spread_record_does_not(self):
        from acql.model import longrecord
        carry = longrecord.carry_over(self._found())
        self.assertEqual((carry.pairs, carry.hot_seasons), (30, 3))
        self.assertAlmostEqual(carry.miss_r, 1.0)
        self.assertEqual(carry.verdict, "strong")
        self.assertLess(carry.cover_r, 0)
        self.assertAlmostEqual(carry.top_was, carry.top_next)
        self.assertGreater(carry.cover_was, carry.cover_next)
        self.assertEqual((carry.hot_beat_typical, carry.hot_beat_line), (3, 0))
        self.assertEqual(longrecord.carry_over({}).verdict, "none")

    def test_the_line_the_crowd_and_form(self):
        from acql.model import longrecord
        found = self._found()
        lines = longrecord.by_season(found)
        self.assertEqual([s.season for s in lines], [2010, 2011, 2012, 2013])
        self.assertEqual((lines[0].models, lines[0].ahead, lines[0].best), (10, 0, "Sagarin Rating"))
        record = longrecord.Record(longrecord.careers(found), longrecord.carry_over(found), lines, 2010, 2013,
                                   longrecord.form(found, 2014))
        self.assertEqual(record.crowd[:2], (0, 4))
        self.assertAlmostEqual(record.crowd[2], -0.5)
        self.assertEqual((record.nudged[0], record.nudged[1], record.nudged[3], record.nudged[4]), (4, 4, 240, 160))
        self.assertEqual(record.nudged_lately, (240, 160, 4))
        self.assertEqual(record.model_seasons, (0, 40))
        self.assertAlmostEqual(record.standing("sag"), 0.45)          # the weekly file's "sag" is Sagarin Rating
        self.assertIsNone(record.standing("zzz"))
        self.assertEqual(record.career("sag").name, "Sagarin Rating")
        self.assertEqual(longrecord.form(found, 2011), {})            # one season is not form
        self.assertNotIn("model1", {k for k, v in longrecord.form(found, 2014).items() if v > 0.4})

    def test_crowds_are_chosen_on_form_from_before_the_season(self):
        from acql.model import computers
        codes = list("abcdefghij")
        form = {code: 1.0 - i / 5 for i, code in enumerate(codes)}       # a is the best, j the worst
        crowds = computers.chosen(codes + ["new"], form)
        self.assertEqual(crowds[computers.PROVEN], set("abcde"))
        self.assertEqual(crowds[computers.WITHOUT_WEAKEST], set("abcdefgh") | {"new"})
        self.assertEqual(computers.chosen(codes[:6], form), {})          # too few with a record to choose from
        found = []
        for i in range(12):
            named = {code: 3.0 + (4.0 * n if n >= 5 else 0.1 * n) for n, code in enumerate(codes)}
            found.append(computers.Reading(week=3, index=i + 1, matchup=f"A{i} @ H{i}", line=1.0, middle=3.0, apart=1.0,
                                           opinions=tuple(sorted(named.values())), named=named, deadeye=1.0, actual=3.0))
        ways = {w.name: w for w in computers.averages(found, form)}
        self.assertLess(ways[computers.PROVEN].miss, ways[computers.IN_USE].miss)      # the good half were the close ones
        self.assertIn(computers.WITHOUT_WEAKEST, ways)
        self.assertNotIn(computers.PROVEN, {w.name for w in computers.averages(found)})
        board = computers.leaderboard(found * 3)
        self.assertIsNone(computers.form_check(board[:4], form))


class TestModelLabForm(unittest.TestCase):
    """Straight up and the pool's line are two records; a record carries its size; a miss is one game."""

    def _lab(self):
        from acql.model import modellab
        # three recipes over six weeks of two games; game 1 of each week has a pool line
        sides = {"A": lambda w, i: True, "B": lambda w, i: w <= 2, "C": lambda w, i: i == 2}
        entries, fixtures = [], {}
        for name, takes_home in sides.items():
            entry = modellab.Entry(name, "", incumbent=name == "A")
            for week in range(1, 7):
                for index in (1, 2):
                    home_credited = not (week == 6 and index == 1)      # one upset, in the last week
                    home = takes_home(week, index)
                    hit = home == home_credited
                    chance = 0.70 if home else 0.30
                    entry.log[(week, index)] = modellab.Play(week, index, chance, hit, 0.5, hit, index == 1,
                                                             6.0 if home else -4.0)
                    tally = entry.by_week.setdefault(week, [0, 0])
                    tally[0] += hit
                    tally[1] += 1
                    entry.games += 1
                    entry.right += hit
                    fixtures[(week, index)] = modellab.Fixture(
                        week, index, f"Road{index} @ Home{index}", f"Home{index}", f"Road{index}", home_credited,
                        7.0 if home_credited else -21.0, index == 1)
            entries.append(entry)
        return modellab.Lab(entries, 12, tuple(range(1, 7)), decided=12, fixtures=fixtures)

    def test_a_record_carries_how_far_luck_moves_it(self):
        from acql.model import modellab
        small, large = modellab.Tally(8, 10), modellab.Tally(120, 200)
        self.assertGreater(small.rate, large.rate)
        self.assertGreater(small.reach, 3 * large.reach)        # 8 of 10 is the less sure of the two
        self.assertIsNone(modellab.Tally().reach)
        self.assertGreater(modellab.Tally(8, 8).reach, 0)       # a perfect record is not a certainty
        self.assertEqual(small + large, modellab.Tally(128, 210))

    def test_winners_and_pool_lines_are_counted_apart(self):
        lab = self._lab()
        a = lab.entry("A")
        self.assertEqual((a.tally("winners").games, a.tally("lined").games), (12, 6))
        self.assertEqual((a.tally("winners").right, a.tally("lined").right), (11, 5))
        leaders, record = lab.best("lined")
        self.assertEqual(([e.name for e in leaders], record.right, record.games), (["A"], 5, 6))
        self.assertEqual([e.name for e in lab.best("winners")[0]], ["A"])

    def test_form_sets_the_last_weeks_beside_the_ones_before(self):
        from acql.model import modellab
        lab = self._lab()
        form = {f.name: f for f in lab.form(weeks=4)}
        self.assertEqual((form["B"].early.right, form["B"].early.games), (4, 4))
        self.assertEqual((form["B"].late.right, form["B"].late.games), (1, 8))
        self.assertEqual(form["B"].reading, modellab.EARLY)     # four games before is too few to say
        hot = modellab.Form("X", *[modellab.Tally()] * 4, early=modellab.Tally(20, 40), late=modellab.Tally(36, 40))
        level = modellab.Form("Y", *[modellab.Tally()] * 4, early=modellab.Tally(24, 40), late=modellab.Tally(27, 40))
        self.assertEqual((hot.reading, level.reading), (modellab.BETTER, modellab.STEADY))
        # the verdict is on the winners the table shows, and two spreads of luck is not enough
        self.assertEqual(form["B"].late, form["B"].winners_late)
        nearly = modellab.Form("Z", *[modellab.Tally()] * 4, early=modellab.Tally(20, 40), late=modellab.Tally(29, 40))
        self.assertTrue(2 < nearly.luck < modellab.FORM_LUCK)
        self.assertEqual(nearly.reading, modellab.STEADY)
        rolling = lab.rolling("A", weeks=4)
        self.assertEqual([week for week, _l, _s in rolling], [1, 2, 3, 4, 5, 6])
        self.assertEqual((rolling[-1][1].games, rolling[-1][2].games), (8, 12))
        self.assertEqual((rolling[-1][1].right, rolling[-1][2].right), (7, 11))

    def test_the_room_and_who_stood_against_it(self):
        from acql.model import modellab
        lab = self._lab()
        rooms = {(r.fixture.week, r.fixture.index): r for r in lab.rooms()}
        upset = rooms[(6, 1)]           # A took the home side; B and C the road side, which was credited
        self.assertEqual((upset.home, upset.took, upset.priced, upset.missed), (False, 2, 3, False))
        self.assertEqual(set(upset.right), {"B", "C"})
        agreed, games, came = lab.agreement()
        self.assertEqual((agreed, games, came), (2, 12, 2))     # only weeks 1 and 2, game 2: all three at home
        against = lab.against_the_room()
        self.assertEqual(against["A"], (3, 3))                  # weeks 3-5, game 1: nobody else had it
        self.assertEqual(against["B"][1], 0)
        kinds = [(s.week, s.kind) for s in lab.shocks() if s.week == 3]
        self.assertIn((3, modellab.SUREST), kinds)
        self.assertIn((3, modellab.SPLIT), kinds)
        worst = lab.upset()
        self.assertTrue(worst is not None and worst.missed and worst.off > 0)


class TestLedger(_ScratchStore):
    """A locked week's numbers are written once, marked by whether each game had kicked off."""

    NOW = 1_791_000_000.0       # early October 2026

    def _season(self):
        from acql.models import Game, Season, Week
        season = Season(title="ACQL 2026")
        season.weeks[4] = Week(4, games=[
            Game(1, home="Denver", away="Seattle", winner="Seattle", home_points=17, away_points=24),
            Game(2, home="Dallas", away="Miami"),
            Game(3, home="Buffalo", away="Chicago"),
        ])
        season.weeks[12] = Week(12, games=[Game(1, home="Detroit", away="Green Bay")])
        return season

    @staticmethod
    def _numbers(value):
        return lambda season, week: {1: {"Deadeye": value, "Betting line only": value + 0.5},
                                     2: {"Deadeye": -1.0, "Betting line only": None}, 3: {}}

    def test_a_locked_week_is_written_once(self):
        from acql.model import ledger
        season = self._season()
        self.assertEqual(ledger.keep(season, self.NOW, self._numbers(3.0)), [4])    # week 12 is still open
        self.assertIsNone(ledger.kept(2026, 12))
        rows = {row["index"]: row for row in ledger.kept(2026, 4)["games"]}
        self.assertEqual(sorted(rows), [1, 2])                  # a game no recipe priced is not written
        # game 1 was played; game 2 has no kickoff time on file, so "before" cannot be claimed
        self.assertEqual((rows[1]["before"], rows[2]["before"]), (False, None))
        self.assertEqual(rows[2]["margins"], {"Deadeye": -1.0})
        # asked again with different numbers: what was written stands
        self.assertEqual(ledger.keep(season, self.NOW + 86400, self._numbers(9.0)), [])
        self.assertEqual({r["index"]: r for r in ledger.kept(2026, 4)["games"]}[1]["margins"]["Deadeye"], 3.0)

    def test_the_replay_is_set_beside_what_was_written(self):
        from acql.model import ledger, modellab
        ledger.keep(self._season(), self.NOW, self._numbers(3.0))

        def lab(deadeye):
            entries = []
            for name, margin in (("Deadeye", deadeye), ("Betting line only", 3.5)):
                entry = modellab.Entry(name, "")
                entry.log[(4, 1)] = modellab.Play(4, 1, 0.6, False, 0.9, False, False, margin)
                entries.append(entry)
            fixture = modellab.Fixture(4, 1, "Seattle @ Denver", "Denver", "Seattle", False, -7.0)
            return modellab.Lab(entries, 1, (4,), fixtures={(4, 1): fixture})

        same = ledger.check(lab(3.0), 2026)
        self.assertEqual((same.graded, same.at_the_time, same.changed), (1, 0, 0))
        self.assertEqual((same.lines[0].numbers, same.lines[0].standing),
                         (2, "0 of 2 before kickoff \u00b7 1 not known"))
        self.assertEqual(same.lines[0].unknown, 1)
        moved = ledger.check(lab(4.5), 2026)
        self.assertEqual((moved.changed, moved.lines[0].worst), (1, "Deadeye"))

    def test_before_kickoff_is_only_claimed_with_a_kickoff(self):
        from acql.model import ledger
        self.assertIs(ledger.written_before({"before": True, "kickoff": 1_791_100_000.0}), True)
        self.assertIs(ledger.written_before({"before": False, "kickoff": None}), False)
        self.assertIsNone(ledger.written_before({"before": None, "kickoff": None}))
        # a row written before this was fixed: marked "before" with no kickoff to show for it
        self.assertIsNone(ledger.written_before({"before": True, "kickoff": None}))
        every = ledger.Line(4, 0.0, 0.0, "", 3, 0, unknown=3)
        self.assertEqual(every.standing, "kickoff times not known")
        self.assertEqual(ledger.Line(4, 0.0, 0.0, "", 3, 3).standing, "written before kickoff")
        self.assertEqual(ledger.Line(4, 0.0, 0.0, "", 3, 0).standing, "written after the games")

    def test_a_draft_dates_the_numbers_that_did_not_move(self):
        from types import SimpleNamespace
        from unittest import mock
        from acql.model import ledger
        from acql.feeds import oddslock
        season = self._season()
        del season.weeks[4]                 # only the week still open
        lock = oddslock.lock_time(2026, 12)
        kickoff = lock + 7200
        numbers = lambda value: (lambda season, week: {1: {"Deadeye": value, "Betting line only": 1.0}})  # noqa: E731
        found = mock.patch.object(ledger.results, "find",
                                  lambda scores, away, home: (SimpleNamespace(kickoff=kickoff), False))
        with found:
            # three days out: too early for a draft; an hour out: a draft, and nothing written for good
            self.assertEqual(ledger.keep(season, lock - 3 * 86400, numbers(2.0)), [])
            self.assertEqual(ledger.drafted(2026, 12), {})
            self.assertEqual(ledger.keep(season, lock - 3600, numbers(2.0)), [])
            self.assertIsNone(ledger.kept(2026, 12))
            self.assertEqual(ledger.drafted(2026, 12)[1]["at"], lock - 3600)
            # the computer slept through the deadline and the game: the numbers are the draft's, so its date stands
            self.assertEqual(ledger.keep(season, kickoff + 86400, numbers(2.0)), [12])
        row = ledger.kept(2026, 12)["games"][0]
        self.assertEqual((row["at"], row["before"]), (lock - 3600, True))
        self.assertIs(ledger.written_before(row), True)

    def test_a_draft_that_moved_dates_nothing(self):
        from types import SimpleNamespace
        from unittest import mock
        from acql.model import ledger
        from acql.feeds import oddslock
        season = self._season()
        del season.weeks[4]                 # only the week still open
        lock = oddslock.lock_time(2026, 12)
        kickoff = lock + 7200
        numbers = lambda value: (lambda season, week: {1: {"Deadeye": value}})  # noqa: E731
        with mock.patch.object(ledger.results, "find",
                               lambda scores, away, home: (SimpleNamespace(kickoff=kickoff), False)):
            ledger.keep(season, lock - 3600, numbers(2.0))
            ledger.keep(season, kickoff + 86400, numbers(2.5))      # the number moved after the draft
        row = ledger.kept(2026, 12)["games"][0]
        self.assertEqual((row["at"], row["before"], row["margins"]), (kickoff + 86400, False, {"Deadeye": 2.5}))

    def test_the_cheap_look_waits_between_tries(self):
        from acql.model import ledger
        self.assertFalse(ledger.due(2026, 1_780_000_000.0))     # June: no week has locked
        ledger.tried(self.NOW)
        self.assertFalse(ledger.due(2026, self.NOW + 60))


class TestSharedDesign(unittest.TestCase):
    """The window and the website are cut from the same tokens (acql/ui/tokens.py)."""

    def test_the_published_page_carries_the_tokens_and_shared_icons(self):
        import json
        import re
        from acql.ui import iconset, poolpage, tokens
        page = poolpage._with_design(poolpage.template())
        self.assertIn(tokens.css(), page)
        self.assertNotIn("__POOL_ICONS__", page)
        shared = json.loads(json.loads(re.search(r"var SHARED_ICONS = (\".*?\");\n", page).group(1)))
        self.assertEqual(shared, iconset.web_icons())

    def test_web_corners_come_off_the_scale(self):
        import re
        from acql.ui import poolpage
        page = poolpage.template()
        loose = [v for v in re.findall(r"border-radius: ([^;{}]+);", page)
                 for part in v.split() if re.fullmatch(r"([2-9]|\d\d+)px", part)]
        self.assertEqual(loose, [])

    def test_the_stylesheet_uses_only_scale_corners(self):
        import re
        try:
            from acql.ui import theme
        except ImportError as exc:  # no Qt-free theme module? nothing to check
            self.skipTest(str(exc))
        allowed = set(theme.RADIUS.values()) | {0, 1, 2, 5, 9}   # 5 and 9: half a scrollbar or slider handle
        for palette in theme.PALETTES.values():
            sheet = theme.stylesheet(palette)
            corners = {int(v) for v in re.findall(r"border-radius: (\d+)px", sheet)}
            self.assertLessEqual(corners, allowed, palette.name)

class TestUpdaterInstall(unittest.TestCase):
    """Installing, failing, rolling back and undoing - on a scratch project, never this folder."""

    def setUp(self):
        import tempfile
        self.tmp = Path(tempfile.mkdtemp(prefix="acql-upd-"))
        self.root = self.tmp / "proj"
        (self.root / "acql").mkdir(parents=True)
        (self.root / "data").mkdir()
        (self.root / "run.py").write_text("old run\n")
        (self.root / "README.md").write_text("old readme\n")
        (self.root / "schedule.ps1").write_text("old schedule\n")
        (self.root / "acql" / "__init__.py").write_text("")
        (self.root / "acql" / "version.py").write_text('VERSION = "1.0.0"\n')
        (self.root / "acql" / "gone.py").write_text("old module\n")
        (self.root / "acql-picks.json").write_text('{"mine": 1}')
        (self.root / "data" / "acql.db").write_text("database")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _zip(self, files, remove=(), version="1.1.0", damage=None):
        import hashlib
        import json
        import zipfile
        from acql import updater
        path = self.tmp / f"ACQL-Dashboard-{version}-update.zip"
        top = f"ACQL-Dashboard-{version}-update"
        manifest = {"version": version, "files": {}, "remove": list(remove), "requires": []}
        with zipfile.ZipFile(path, "w") as zf:
            for rel, text in files.items():
                data = text.encode()
                manifest["files"][rel] = hashlib.sha256(data).hexdigest()
                zf.writestr(f"{top}/{updater.PAYLOAD}/{rel}", b"x" + data if rel == damage else data)
            zf.writestr(f"{top}/{updater.MANIFEST}", json.dumps(manifest))
        return updater.open_update(path)

    def _good(self):
        return {"run.py": "new run\n", "README.md": "new readme\n", "schedule.ps1": "new schedule\n",
                "acql/version.py": 'VERSION = "1.1.0"\n', "acql/added.py": "new module\n"}

    def test_a_good_update_installs_and_leaves_the_data(self):
        from acql import updater
        done = updater.install(self._zip(self._good(), remove=["acql/gone.py"]), self.root,
                               check=lambda r: "", say=lambda *_: None)
        self.assertEqual(done.version, "1.1.0")
        self.assertEqual((self.root / "README.md").read_text(), "new readme\n")
        self.assertFalse((self.root / "acql" / "gone.py").exists())
        self.assertEqual((self.root / "acql-picks.json").read_text(), '{"mine": 1}')
        self.assertEqual((self.root / "data" / "acql.db").read_text(), "database")

    def _git(self, *args, cwd=None):
        import subprocess
        done = subprocess.run(["git", "-c", "user.name=t", "-c", "user.email=t@t", "-c", "core.autocrlf=false",
                               *args], cwd=str(cwd or self.root), capture_output=True, text=True)
        self.assertEqual(done.returncode, 0, done.stderr)

    def _checkout(self):
        """The scratch project as a git checkout with everything pushed to a scratch remote."""
        remote = self.tmp / "remote.git"
        self._git("init", "--bare", str(remote), cwd=self.tmp)
        self._git("init")
        self._git("add", "-A")
        self._git("commit", "-m", "start")
        self._git("remote", "add", "origin", str(remote))
        self._git("push", "origin", "HEAD:refs/heads/main")
        self._git("fetch", "origin")

    def test_a_checkout_with_work_not_on_github_is_never_overwritten(self):
        import shutil as _shutil
        from acql import updater
        if _shutil.which("git") is None:
            self.skipTest("git isn't installed")
        self._checkout()
        self.assertEqual(updater.work_at_risk(self.root), "")
        module = self.root / "acql" / "gone.py"
        data = module.read_bytes()                                             # line endings only: not work
        module.write_bytes(data.replace(b"\r\n", b"\n") if b"\r\n" in data else data.replace(b"\n", b"\r\n"))
        self.assertEqual(updater.work_at_risk(self.root), "")
        module.write_text("my edit\n")                                         # real work, not committed
        with self.assertRaises(updater.UpdateError) as caught:
            updater.install(self._zip(self._good()), self.root, check=lambda r: "", say=lambda *_: None)
        self.assertIn("1 changed file not committed (acql/gone.py)", str(caught.exception))
        self.assertEqual((self.root / "run.py").read_text(), "old run\n")      # nothing was touched
        self.assertFalse((self.root / "backups").exists())
        self._git("add", "-A")
        self._git("commit", "-m", "mine")                                       # committed, not pushed
        self.assertIn("1 commit not pushed to GitHub", updater.work_at_risk(self.root))
        (self.root / "acql" / "brand_new.py").write_text("x = 1\n")
        self.assertIn("1 new file not committed (acql/brand_new.py)", updater.work_at_risk(self.root))
        (self.root / "acql" / "brand_new.py").unlink()
        self._git("push", "origin", "HEAD:refs/heads/main")
        self._git("fetch", "origin")
        self.assertEqual(updater.work_at_risk(self.root), "")                  # all on GitHub: installs
        done = updater.install(self._zip(self._good()), self.root, check=lambda r: "", say=lambda *_: None)
        self.assertEqual(done.version, "1.1.0")

    def test_a_failed_start_puts_every_overwritten_file_back(self):
        from acql import updater
        with self.assertRaises(updater.UpdateError):
            updater.install(self._zip(self._good(), remove=["acql/gone.py"]), self.root,
                            check=lambda r: "ImportError: boom", say=lambda *_: None)
        self.assertEqual((self.root / "run.py").read_text(), "old run\n")
        # README.md and schedule.ps1 aren't .py files: the backup used to leave them out
        self.assertEqual((self.root / "README.md").read_text(), "old readme\n")
        self.assertEqual((self.root / "schedule.ps1").read_text(), "old schedule\n")
        self.assertEqual((self.root / "acql" / "gone.py").read_text(), "old module\n")
        self.assertFalse((self.root / "acql" / "added.py").exists())
        self.assertEqual(updater.installed_version(self.root), "1.0.0")

    def test_undo_brings_back_the_version_before(self):
        from acql import updater
        updater.install(self._zip(self._good(), remove=["acql/gone.py"]), self.root,
                        check=lambda r: "", say=lambda *_: None)
        self.assertEqual(updater.undo_last(self.root), "1.0.0")
        self.assertEqual((self.root / "schedule.ps1").read_text(), "old schedule\n")
        self.assertEqual((self.root / "acql" / "gone.py").read_text(), "old module\n")
        self.assertFalse((self.root / "acql" / "added.py").exists())

    def test_a_damaged_file_is_refused_before_anything_changes(self):
        from acql import updater
        with self.assertRaises(updater.UpdateError):
            updater.install(self._zip(self._good(), damage="run.py"), self.root,
                            check=lambda r: "", say=lambda *_: None)
        self.assertEqual((self.root / "run.py").read_text(), "old run\n")
        self.assertFalse((self.root / "backups").exists())

    def test_paths_outside_the_code_are_refused_before_anything_changes(self):
        from acql import updater
        for bad in ("../escaped.py", "/abs.py", "C:/x.py", "acql/../../escaped.py", "data/acql.db",
                    "backups/x.zip", ".venv/x.py", "acql-picks.json", "a//b.py"):
            with self.subTest(files=bad), self.assertRaises(updater.UpdateError):
                updater.install(self._zip({**self._good(), bad: "evil"}), self.root,
                                check=lambda r: "", say=lambda *_: None)
            with self.subTest(remove=bad), self.assertRaises(updater.UpdateError):
                updater.install(self._zip(self._good(), remove=[bad]), self.root,
                                check=lambda r: "", say=lambda *_: None)
        self.assertFalse((self.tmp / "escaped.py").exists())
        self.assertEqual((self.root / "data" / "acql.db").read_text(), "database")
        self.assertEqual((self.root / "acql-picks.json").read_text(), '{"mine": 1}')
        self.assertEqual((self.root / "run.py").read_text(), "old run\n")
        self.assertFalse((self.root / "backups").exists())

    def test_the_real_release_passes_the_path_check(self):
        from acql import updater
        zips = sorted((ROOT / "updates").glob("ACQL-Dashboard-*-update.zip"))
        if not zips:
            self.skipTest("no release zip in updates/")
        update = updater.open_update(zips[-1])
        for rel in [*update.files, *update.remove]:
            self.assertEqual(updater.safe_path(rel), rel)


class TestAtomicWrites(unittest.TestCase):
    """aliases.json and acql-settings.json are replaced whole or left as they were."""

    def test_a_failed_write_leaves_the_old_aliases(self):
        import os
        from unittest import mock
        with tempfile.TemporaryDirectory() as scratch:
            path = Path(scratch) / "aliases.json"
            table = AliasTable()
            table.add("Zachary N", "Zach N")
            table.save(path)
            before = path.read_bytes()
            other = AliasTable()
            other.add("Debbie", "Debe")
            with mock.patch.object(os, "replace", side_effect=OSError("disk full")):
                other.save(path)                     # save() swallows OSError, as before
            self.assertEqual(path.read_bytes(), before)
            self.assertEqual(sorted(p.name for p in Path(scratch).iterdir()), ["aliases.json"])
            self.assertEqual(AliasTable.load(path).key("zachary n"), normalize("Zach N"))
            other.save(path)
            self.assertEqual(AliasTable.load(path).key("Debbie"), normalize("Debe"))

    def test_settings_round_trip_through_the_atomic_write(self):
        from acql import config
        with tempfile.TemporaryDirectory() as scratch:
            path = Path(scratch) / "acql-settings.json"
            real = config.SETTINGS_FILE
            config.SETTINGS_FILE = path
            try:
                Settings(buy_in=55.0, last_week_viewed=4).save()
                again = Settings.load()
            finally:
                config.SETTINGS_FILE = real
            self.assertEqual((again.buy_in, again.last_week_viewed), (55.0, 4))
            self.assertEqual([p.name for p in Path(scratch).iterdir()], ["acql-settings.json"])


class TestPieSplits(unittest.TestCase):
    """stats.xls's Pie Splits column (suicide + BL + top 10): side-pool money,
    read as a season total and, between weekly exports, by the week it was paid."""

    class _Sheet:
        def __init__(self, week, rows):
            grid = {(2, 5): f"Week {week}", (3, 11): "(suicide + BL + top 10)", (4, 1): "Coach",
                    (4, 11): "Pie Splits"}
            for r, (coach, won, pie) in enumerate(rows, start=5):
                grid.update({(r, 0): r - 4, (r, 1): coach, (r, 7): won, (r, 11): pie})
            self.grid, self.nrows, self.ncols = grid, 5 + len(rows), 40

        def cell_value(self, r, c):
            return self.grid.get((r, c), "")

    def _read(self, snapshots):
        from unittest import mock
        from acql.models import SourceFile
        from acql.names import AliasTable
        from acql.sources import stats_xls
        people = {}
        for week, rows in snapshots:
            sheet = self._Sheet(week, rows)
            book = mock.Mock(sheet_names=lambda: ["stats"], sheet_by_name=lambda name, s=sheet: s)
            with mock.patch("xlrd.open_workbook", lambda path, b=book: b):
                stats_xls.parse_stats(SourceFile(Path(f"stats wk{week}.xls"), "stats", week), AliasTable(), people)
        return people

    def test_payout_lands_in_its_week(self):
        people = self._read([(4, [("Gr8granpa", 0, ""), ("Line-Pony", 75.5, "")]),
                             (5, [("Gr8granpa", 45, 45), ("Line-Pony", 103.5, "")])])
        gr8 = next(p for p in people.values() if p.display == "Gr8granpa")
        self.assertEqual((gr8.pie_splits, gr8.weekly_pie), (45.0, {5: 45.0}))
        pony = next(p for p in people.values() if p.display == "Line-Pony")
        self.assertEqual((pony.pie_splits, pony.weekly_pie), (0.0, {}))

    def test_running_total_is_split_by_week(self):
        people = self._read([(5, [("Gr8granpa", 45, 45)]), (8, [("Gr8granpa", 75, 75)])])
        gr8 = next(iter(people.values()))
        self.assertEqual(gr8.weekly_pie, {5: 45.0, 8: 30.0})

    def test_winnings_page_counts_it(self):
        try:
            from acql.ui.pages.winnings import WinningsPage
        except Exception as exc:  # noqa: BLE001
            raise unittest.SkipTest(f"PySide6 unavailable: {exc}")
        from acql.models import Season
        people = self._read([(5, [("Gr8granpa", 45, 45)])])
        season = Season()
        season.players = people
        standings, source = WinningsPage._standings(season)
        self.assertEqual((source, standings[0].won, standings[0].weekly), ("stats", 45.0, {5: 45.0}))

    def test_net_is_stats_net_pts(self):
        """stats.xls is the pool's own book: its NET PTS is the net everywhere, whatever the weeks add up to."""
        from acql.models import Player, Season
        gr8 = Player(key="gr8", display="Gr8granpa", points_won=45.0, points_paid=-62.0, net_points=-17.0,
                     weekly_pie={5: 45.0})
        zach = Player(key="zach", display="Zach N", points_won=176.0, points_paid=-59.0, net_points=117.0,
                      weekly_winnings={3: 167.0})
        self.assertEqual((gr8.net(40.0), gr8.total_winnings), (-17.0, 45.0))
        self.assertEqual((zach.net(40.0), zach.total_winnings), (117.0, 176.0))
        # with no standings export, the weeks against the buy-in, as before
        self.assertEqual(Player(key="x", display="X", weekly_winnings={1: 50.0}).net(40.0), 10.0)
        try:
            from acql.ui.pages.winnings import WinningsPage
        except Exception as exc:  # noqa: BLE001
            raise unittest.SkipTest(f"PySide6 unavailable: {exc}")
        season = Season()
        season.players = {"gr8": gr8, "zach": zach}
        standings, source = WinningsPage._standings(season)
        self.assertEqual(source, "stats")
        self.assertEqual({s.name: (s.won, s.paid, s.net) for s in standings},
                         {"Gr8granpa": (45.0, 62.0, -17.0), "Zach N": (176.0, 59.0, 117.0)})


class SurvivorFieldTests(unittest.TestCase):
    """The suicide pool field: shares among unspent teams, and the fit."""

    def test_shares_skip_spent_and_favour_safer(self):
        from acql.pool import survivorfield as sf
        share = sf._shares({"a": 0.7, "b": 0.5, "c": 0.6}, {"c"}, 10.0)
        self.assertEqual(set(share), {"a", "b"})
        self.assertAlmostEqual(sum(share.values()), 1.0)
        self.assertGreater(share["a"], share["b"])

    def test_fit_finds_sharp_pickers(self):
        from acql.pool import survivorfield as sf
        chances = {"a": 0.75, "b": 0.5, "c": 0.45}
        made = [sf.Made("x", w, "a", chances, frozenset()) for w in range(1, 9)]
        self.assertGreater(sf.fit(made), sf.fit(made[:1] + [sf.Made("y", 9, "c", chances, frozenset())] * 3))


class TestScoreboardRibbon(unittest.TestCase):
    """The ribbon over every page: the week's games in ticker order, exact
    kickoff times, your pick in gold, and a chip that opens Game Day."""

    @staticmethod
    def _game(index, state, kickoff, away="Dallas", home="Green Bay", away_points=None, home_points=None,
              clock=""):
        from types import SimpleNamespace
        from acql.pool import gameday as GD
        margin = None if home_points is None or away_points is None else home_points - away_points
        return SimpleNamespace(
            index=index, state=state, kickoff=kickoff, clock=clock,
            game=SimpleNamespace(away=away, home=home, label=f"{away} at {home}"),
            home_points=home_points, away_points=away_points, margin=margin,
            started=state in ("in", "final"), decided=state == "final",
            leader=(GD.HOME if (margin or 0) > 0 else GD.AWAY if (margin or 0) < 0 else GD.NOBODY),
            possession="", line_text=lambda: "straight up")

    def test_on_now_then_by_kickoff_then_the_finals(self):
        from acql.ui.ribbon import in_order
        games = [self._game(1, "final", 100), self._game(2, "pre", 900), self._game(3, "in", 500),
                 self._game(4, "pre", 700), self._game(5, "final", 300), self._game(6, "off", None)]
        self.assertEqual([g.index for g in in_order(games)], [3, 4, 2, 6, 5, 1])

    def test_kickoff_is_a_day_and_a_time_never_a_countdown(self):
        import time as _time
        from acql.ui.ribbon import kickoff_words
        soon = _time.mktime((2026, 10, 18, 13, 0, 0, 0, 0, -1))      # a Sunday, 1 pm
        self.assertEqual(kickoff_words(soon), "Sun 1:00 PM")
        self.assertEqual(kickoff_words(_time.mktime((2026, 10, 19, 20, 15, 0, 0, 0, -1))), "Mon 8:15 PM")
        self.assertEqual(kickoff_words(None), "")

    def test_only_while_a_week_is_being_played(self):
        from types import SimpleNamespace
        from acql.feeds import results
        from acql.ui.ribbon import playing_week
        year = 2026
        mid = results.week_starts(year, 6) + 4 * 86400
        season = SimpleNamespace(is_empty=False, results_year=year, title="ACQL 2026",
                                 weeks={6: SimpleNamespace(games=[object()])})
        self.assertEqual(playing_week(season, mid), 6)
        self.assertEqual(playing_week(season, results.week_starts(year, 7) + 3600), 0)   # no games set up
        self.assertEqual(playing_week(season, results.week_starts(year, 1) - 86400), 0)  # before the season
        self.assertEqual(playing_week(SimpleNamespace(is_empty=True), mid), 0)

    def test_chips_are_one_row_in_order_with_your_pick_in_gold(self):
        import os
        from types import SimpleNamespace
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        try:
            from PySide6.QtWidgets import QApplication, QLabel
        except ImportError as exc:  # pragma: no cover - Qt absent
            self.skipTest(f"PySide6 unavailable: {exc}")
        type(self).app = QApplication.instance() or QApplication([])
        from acql.pool import gameday as GD
        from acql.ui.ribbon import ScoreboardRibbon, kickoff_words
        from acql.ui.theme import PALETTES
        ribbon = ScoreboardRibbon(PALETTES["dark"])
        games = [self._game(1, "final", 100, "Houston", "Jacksonville", 34, 31, "Final"),
                 self._game(2, "pre", 900_000, "Arizona", "Los Angeles Rams"),
                 self._game(3, "in", 500, "Chicago", "Atlanta", 21, 7, "Q2 1:02")]
        you = SimpleNamespace(picks={3: GD.AWAY, 2: GD.HOME})
        ribbon.show_games(6, games, you)
        self.assertFalse(ribbon.isHidden())
        self.assertTrue(ribbon.live)
        self.assertEqual([c.index for c in ribbon.chips], [3, 2, 1])
        self.assertTrue(all(ribbon.lane_row.indexOf(c) >= 0 for c in ribbon.chips))   # one row
        self.assertEqual(ribbon.chips[1].state_label.text(), kickoff_words(900_000))
        gold = {c.index: [lab.full_text() for lab in c.findChildren(QLabel) if lab.property("pick") == "true"]
                for c in ribbon.chips}
        self.assertEqual(gold, {3: ["Chicago"], 2: ["Los Angeles Rams"], 1: []})
        chosen = []
        ribbon.game_chosen.connect(lambda week, index: chosen.append((week, index)))
        ribbon.chips[1].pressed.emit(ribbon.chips[1].index)
        self.assertEqual(chosen, [(6, 2)])
        ribbon.show_games(0, [], None)                                           # between weeks
        self.assertTrue(ribbon.isHidden())


class PickDeadlineTests(unittest.TestCase):
    def test_next_deadline_skips_one_already_past_and_stays_out_of_the_fingerprint(self):
        from unittest import mock
        from acql.feeds import oddslock, results
        from acql.ui import poolpage
        locks = {6: 1000.0, 7: 5000.0}
        from acql.site import mainsite
        with mock.patch.object(mainsite, "due", return_value=None), \
             mock.patch.object(results, "season_year", return_value=2026), \
             mock.patch.object(results, "calendar_weeks", return_value=[6, 7]), \
             mock.patch.object(oddslock, "lock_time", side_effect=lambda y, w: locks[w]):
            self.assertEqual(poolpage._deadline(Season(), now=2000.0), {"week": 7, "at": 5_000_000, "form": mainsite.FORM})
            self.assertIsNone(poolpage._deadline(Season(), now=6000.0))
        shown = poolpage._as_shown({"deadline": {"week": 7, "at": 1}, "weeks": []})
        self.assertNotIn("deadline", shown)

    def test_the_main_sites_due_time_is_read_as_eastern(self):
        import io
        from acql.site import mainsite
        page = ("<p><span>Week 7&nbsp; <a href=x>Enter Picks</a>&nbsp;&nbsp; <b>(Due Thursday 6:30pm)</b></span></p>"
                "<p>2026 Week 6 picks</p>").encode("cp1252")
        self.assertEqual(mainsite.parse_due(mainsite.page_text(page)), (7, 3, 18, 30))
        self.assertEqual(mainsite.parse_due("Week 12 Enter Picks (Due Wed 1 PM)"), (12, 2, 13, 0))
        self.assertIsNone(mainsite.parse_due("Line-Pony takes the week with 10 wins"))
        found = mainsite.due(2026, opener=lambda url, timeout: io.BytesIO(page))
        # Thursday 22 October 2026, 6:30pm EDT is 22:30 UTC
        self.assertEqual(found, {"week": 7, "at": 1792708200, "form": mainsite.FORM})
        linked = page.replace(b"<a href=x>", b"<a href=\"http://example.com/form.cgi\"><span style='color:#92D050'>")
        self.assertEqual(mainsite.due(2026, opener=lambda url, timeout: io.BytesIO(linked))["form"], "http://example.com/form.cgi")
        # in November, after the clocks go back, it is five hours behind
        from datetime import date
        self.assertEqual(mainsite.eastern(date(2026, 11, 12), 18, 30) - mainsite.eastern(date(2026, 11, 12), 0, 0), 18.5 * 3600)
        self.assertEqual(mainsite.eastern(date(2026, 11, 12), 18, 30), 1794526200)
        self.assertIsNone(mainsite.due(2026, opener=lambda url, timeout: io.BytesIO(b"<p>nothing here</p>")))

class TestOddsPapiBackup(_ScratchStore):
    """The second odds source: used only when the main feed can't answer, and only after a check agreed."""

    KICKOFF = "2030-01-06T18:00:00Z"
    # (participant1, participant2, participant1's margin, its decimal moneyline, participant2's)
    GAMES = [("Green Bay Packers", "Chicago Bears", 6.5, 1.40, 3.10),       # GB favoured
             ("Kansas City Chiefs", "Denver Broncos", 3.5, 1.60, 2.45),
             ("Dallas Cowboys", "Philadelphia Eagles", -2.5, 2.30, 1.65),   # PHI (listed second) favoured
             ("Buffalo Bills", "Miami Dolphins", 9.5, 1.20, 5.00),
             ("Detroit Lions", "Minnesota Vikings", 1.5, 1.80, 2.05)]
    # the week's schedule, (away, home): Green Bay is away at Chicago, the others listed first are at home
    PAIRS = [("gb", "chi"), ("den", "kc"), ("phi", "dal"), ("mia", "buf"), ("min", "det")]
    AGREED = {("gb", "chi"): -6.5, ("den", "kc"): 3.5, ("phi", "dal"): -2.5, ("mia", "buf"): 9.0, ("min", "det"): 1.5}

    def setUp(self):
        super().setUp()
        import json
        import os
        from unittest import mock
        from acql.feeds import oddspapi
        self.json = json
        self.patches = [mock.patch.dict(os.environ, {"ACQL_ODDSPAPI_KEY": "made-up", "ACQL_ODDS_API_KEY": "made-up"}),
                        mock.patch.object(oddspapi, "_pause", lambda s: None)]
        for p in self.patches:
            p.start()
        self.calls = []

    def tearDown(self):
        for p in self.patches:
            p.stop()
        super().tearDown()

    def _papi(self, flip=False):
        """OddsPapi's replies. `flip` writes each handicap the other way round."""
        names = {}
        for g in self.GAMES:
            n = len(names)
            names[str(n + 1)], names[str(n + 2)] = g[0], g[1]
        ids = {v: k for k, v in names.items()}
        markets = [
            {"marketId": 501, "marketName": "Moneyline (incl. Overtime)", "marketType": "moneyline", "sportId": 9,
             "playerProp": False, "handicap": 0, "period": "fulltime", "outcomes": [{"outcomeId": 501}, {"outcomeId": 502}]},
            {"marketId": 511, "marketName": "1st Half Moneyline", "marketType": "moneyline", "sportId": 9,
             "playerProp": False, "handicap": 0, "period": "1st half", "outcomes": [{"outcomeId": 511}, {"outcomeId": 512}]},
        ]
        fixtures, mid = [], 600
        for p1, p2, margin, a, b in self.GAMES:
            def side(oid, price):
                return {str(oid): {"players": {"0": {"price": price}}}}
            book = {"501": {"outcomes": {**side(501, a), **side(502, b)}},
                    "511": {"outcomes": {**side(511, 9.0), **side(512, 1.05)}}}
            for m, prices in ((margin, (1.91, 1.91)), (margin + 1, (2.40, 1.60))):   # the main line and an alternative
                mid += 2
                written = m if flip else -m
                markets.append({"marketId": mid, "marketName": f"Handicap {written}", "marketType": "spreads",
                                "sportId": 9, "playerProp": False, "handicap": written, "period": "fulltime",
                                "outcomes": [{"outcomeId": mid * 10}, {"outcomeId": mid * 10 + 1}]})
                book[str(mid)] = {"outcomes": {**side(mid * 10, prices[0]), **side(mid * 10 + 1, prices[1])}}
            fixtures.append({"participant1Id": int(ids[p1]), "participant2Id": int(ids[p2]), "startTime": self.KICKOFF,
                             "bookmakerOdds": {"pinnacle": {"markets": book}}})
        return {"sports": [{"sportId": 10, "sportName": "Soccer"}, {"sportId": 9, "sportName": "American Football"}],
                "tournaments": [{"tournamentId": 77, "tournamentName": "NCAAF"},
                                {"tournamentId": 88, "tournamentName": "NFL"}],
                "participants": names, "markets": markets, "odds-by-tournaments": fixtures,
                "account": {"request_limit": 250, "request_count": 12}}

    def _opener(self, replies, fail=None):
        import io
        import urllib.error

        class Reply(io.BytesIO):
            headers = {"x-requests-remaining": "300"}

        def open_url(request):
            url = request.full_url
            self.calls.append(url)
            path = url.split("?")[0].rstrip("/").rsplit("/", 1)[-1]
            if fail:
                raise urllib.error.HTTPError(url, fail, "no", {}, io.BytesIO(b"down for a bit"))
            return Reply(self.json.dumps(replies[path]).encode())
        return open_url

    def _main_feed(self, margins):
        """The Odds API's reply for {(away, home): home margin}."""
        full = {"gb": "Green Bay Packers", "chi": "Chicago Bears", "kc": "Kansas City Chiefs", "den": "Denver Broncos",
                "dal": "Dallas Cowboys", "phi": "Philadelphia Eagles", "buf": "Buffalo Bills", "mia": "Miami Dolphins",
                "det": "Detroit Lions", "min": "Minnesota Vikings"}
        return {"odds": [{"home_team": full[h], "away_team": full[a], "commence_time": self.KICKOFF,
                          "bookmakers": [{"markets": [{"key": "spreads", "outcomes": [
                              {"name": full[h], "point": -m}, {"name": full[a], "point": m}]}]}]}
                         for (a, h), m in margins.items()]}

    def test_lines_read_right_whichever_way_the_handicap_is_written(self):
        from acql import store
        from acql.feeds import odds, oddspapi
        for flip in (False, True):
            store.delete(oddspapi.KIND)
            store.delete("odds")
            quotes, left = oddspapi.fetch("made-up", opener=self._opener(self._papi(flip=flip)))
            self.assertEqual(left, 238)
            done = odds.update(18, quotes, self.PAIRS, now=0)
            kept = odds.stored(18).latest
            self.assertEqual(done.matched, 5)
            self.assertEqual(kept["gb@chi"], -6.5)       # Green Bay by 6.5, away at Chicago
            self.assertEqual(kept["den@kc"], 3.5)
            self.assertEqual(kept["phi@dal"], -2.5)      # the team listed second the favourite
            self.assertGreater(odds.stored(18).win["den@kc"], 0.55)
            self.assertNotIn("chi@gb", kept)              # the wrong way round never lands

    def test_discovery_is_kept_and_half_time_markets_are_left_out(self):
        from acql.feeds import oddspapi
        oddspapi.fetch("made-up", opener=self._opener(self._papi()))
        first = len(self.calls)
        oddspapi.fetch("made-up", opener=self._opener(self._papi()))
        self.assertEqual(len(self.calls) - first, 2)     # the odds, and the free account check
        ids = oddspapi._read()["ids"]
        self.assertEqual((ids["sport"], ids["tournament"]), (9, 88))
        self.assertEqual(list(ids["markets"]["ml"]), ["501"])
        self.assertEqual(ids["teams"]["1"], "gb")

    def test_the_backup_is_off_until_a_check_agrees_then_steps_in(self):
        from acql.feeds import odds, oddspapi
        papi = self._papi()
        down = self._opener({}, fail=500)
        with self.assertRaises(odds.OddsError):           # no check yet: a failure is still a failure
            odds.fetch_any("made-up", opener=down, backup_opener=self._opener(papi))
        self.assertFalse(oddspapi.confirmed())
        # the main feed answers, and the backup is checked against it
        quotes, left = odds.fetch_any("made-up", opener=self._opener(self._main_feed(self.AGREED)),
                                      backup_opener=self._opener(papi))
        self.assertEqual(left, 300)
        self.assertTrue(oddspapi.ready(), oddspapi.last_check())
        # now the main feed is down: Pinnacle's lines step in, and the main feed's count is untouched
        quotes, left = odds.fetch_any("made-up", opener=down, backup_opener=self._opener(papi))
        self.assertIsNone(left)
        done = odds.update(18, quotes, self.PAIRS, now=0)
        self.assertEqual(done.matched, 5)
        self.assertIn("OddsPapi", done.summary())
        # the reserve sends the daily run straight to it, without asking the main feed
        before = len(self.calls)
        odds.fetch_any("made-up", opener=down, backup_opener=self._opener(papi), skip_main="reserve")
        self.assertFalse(any("the-odds-api" in c for c in self.calls[before:]))

    def test_a_check_that_disagrees_keeps_it_off(self):
        from acql.feeds import odds, oddspapi
        off = dict(self.AGREED)
        off[("den", "kc")] = -3.5                         # one game the wrong way round
        odds.fetch_any("made-up", opener=self._opener(self._main_feed(off)), backup_opener=self._opener(self._papi()))
        self.assertFalse(oddspapi.ready())
        self.assertEqual(len(oddspapi.last_check()["apart"]), 1)
        self.assertIn("stays off", oddspapi.describe(oddspapi.last_check()))
        with self.assertRaises(odds.OddsError):
            odds.fetch_any("made-up", opener=self._opener({}, fail=500), backup_opener=self._opener(self._papi()))


def main() -> int:
    if STATS_PATH is None and WORKBOOK_PATH is None:
        print("No spreadsheets found in data/ - only the pure-logic tests will run.\n")
    loader = unittest.TestLoader()
    suite = loader.loadTestsFromModule(sys.modules[__name__])
    result = unittest.TextTestRunner(verbosity=2).run(suite)
    return 0 if result.wasSuccessful() else 1


if __name__ == "__main__":
    raise SystemExit(main())
