"""The pool's raw pick sheet - everyone's picks for a week, before grading.

This is the file the commissioner publishes each week and that gets pasted
into the workbook: a column per game, a row per coach, and at the end each
coach's three Big Loser picks and their Suicide pick. The sheet it comes on
looks like this:

    row 1                         ACQL WEEK 2 ...
    row 2                         (phi by 8)          the pool's line
    row 3   cleveland             new orleans         the road team
    row 5   TAMPA BAY             BALTIMORE           the home team
    row 7+  Phelpy   TAMPA BAY    BALTIMORE  ...      one row a coach

Nothing here grades anything. What it is for is the one number that cannot
be worked out from anywhere else: how much of the pool is on each side. A
pick that is right and popular wins nothing; the whole of the gain on the
field is in being right where the others are wrong, and until this file is
loaded that share has to be estimated from a curve.

The week is read from the sheet's own title rather than the file name,
because the files are numbered by something other than the week.

The same link serves the sheet again once the week is scored, graded: each
coach's total and regular wins in the first two columns, the name moved to
the third, and every wrong pick deleted. Read as a raw sheet, its totals
became the coaches' names and every winner looked like everyone's pick. So a
graded sheet is recognised, and each coach's card is rebuilt from it - every
pick is one of two teams, so a blank under a graded game means the coach had
the other side. Its Big Losers are only half there (a correct one is written
over with "1"), which is why the sheet says it was graded.
"""

from __future__ import annotations

import re
import time
from collections import Counter
from dataclasses import asdict, dataclass, field
from pathlib import Path

from .config import ROOT
from .models import Game, Player, PlayerWeek, Season, Week
from .predictions import DISPLAY, team_code

#: Where older versions kept the pick sheets; read in once by store.py.
PICKS_FILE = ROOT / "acql-picks.json"
# Marks a player as known only from a raw pick sheet, so a later workbook
# load can be told apart from this one and left in charge of the standings.
SHEET_SOURCE = "pick sheet: "

FIRST_PLAYER_ROW = 7
AWAY_ROW, HOME_ROW, NOTE_ROW = 3, 5, 2
TITLE = re.compile(r"WEEK\s*(\d+)", re.IGNORECASE)
LOSER_HEADER = re.compile(r"LOSER\s*(\d)", re.IGNORECASE)
SUICIDE_HEADER = re.compile(r"SUICIDE", re.IGNORECASE)
# Rows at the bottom that summarise rather than pick.
SUMMARY_ROW = re.compile(r"^\s*(picked|total|average|avg)\b", re.IGNORECASE)
NUMBER = re.compile(r"^\s*-?\d+(\.\d+)?\s*$")


class Unreadable(Exception):
    """The file could not be opened at all - the message says what to do."""


def _grid(path: str | Path) -> list[list[str]]:
    """The sheet as rows of strings, from either .xlsx or an older .xls.

    openpyxl reads the modern format; the old one needs xlrd, which is not
    always installed. Saying so plainly beats a traceback about magic bytes.
    """
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError as problem:            # pragma: no cover
            raise Unreadable("openpyxl isn't installed, so no spreadsheet can be read") from problem
        book = openpyxl.load_workbook(path, data_only=True)
        sheet = book["picks"] if "picks" in book.sheetnames else book.worksheets[0]
        return [
            ["" if cell.value is None else str(cell.value).strip() for cell in row]
            for row in sheet.iter_rows()
        ]
    try:
        import xlrd
    except ImportError as problem:
        raise Unreadable(
            f"{path.name} is in the old .xls format, which needs the xlrd package. "
            f"Either run \"pip install xlrd\", or open the file and save it again "
            f"as .xlsx - both work."
        ) from problem
    book = xlrd.open_workbook(path)
    names = book.sheet_names()
    sheet = book.sheet_by_name("picks" if "picks" in names else names[0])
    return [
        [_text(sheet.cell_value(r, c)) for c in range(sheet.ncols)]
        for r in range(sheet.nrows)
    ]


def _text(value: object) -> str:
    """A cell as the sheet shows it. The old format hands every number back
    as a float, so a Big Loser marked 1 would read "1.0" and go uncounted."""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _cell(rows: list[list[str]], row: int, col: int) -> str:
    """1-based, and off the end of the sheet is simply empty."""
    if 1 <= row <= len(rows) and 1 <= col <= len(rows[row - 1]):
        return rows[row - 1][col - 1]
    return ""


@dataclass(frozen=True)
class GamePicks:
    """One game, and who took which side."""

    away: str
    home: str
    note: str                       # the pool's line, as written on the sheet
    picks: dict[str, str] = field(default_factory=dict)     # coach -> team picked
    #: On a graded sheet, the side the commissioner credited - the one whose
    #: picks were left standing. "" on a raw sheet, or a game not yet graded.
    credited: str = ""

    @property
    def counted(self) -> int:
        return sum(1 for team in self.picks.values() if team_code(team))

    def share_on(self, team: str) -> float | None:
        """The fraction of the pool that took this team, or None if nobody did."""
        wanted = team_code(team)
        if not wanted or not self.counted:
            return None
        return sum(
            1 for pick in self.picks.values() if team_code(pick) == wanted
        ) / self.counted


@dataclass(frozen=True)
class WeekPicks:
    """Everyone's card for one week."""

    week: int
    games: list[GamePicks] = field(default_factory=list)
    losers: dict[str, list[str]] = field(default_factory=dict)    # coach -> 3 teams
    suicide: dict[str, str] = field(default_factory=dict)         # coach -> team
    source: str = ""
    #: Read from the scored sheet: the sides are rebuilt, but a correct Big
    #: Loser is only a "1", so the Big Loser shares are not to be trusted.
    graded: bool = False
    #: On a graded sheet, each coach's [total, regular wins] as the
    #: commissioner scored them, from the columns left of the name.
    scores: dict[str, list[int]] = field(default_factory=dict)

    @property
    def coaches(self) -> list[str]:
        names = {name for game in self.games for name in game.picks}
        return sorted(names | set(self.losers) | set(self.suicide))

    def game_for(self, away: str, home: str) -> GamePicks | None:
        """The sheet's row for a matchup, whichever way round it is written."""
        want = {team_code(away), team_code(home)}
        if "" in want or len(want) != 2:
            return None
        for game in self.games:
            if {team_code(game.away), team_code(game.home)} == want:
                return game
        return None

    def loser_counts(self) -> list[tuple[str, int]]:
        """How many coaches took each Big Loser, most popular first."""
        tally = Counter(
            team_code(team) for picks in self.losers.values() for team in picks
            if team_code(team)
        )
        return tally.most_common()

    def loser_share(self, team: str) -> float:
        """The fraction of coaches with this team among their three."""
        wanted = team_code(team)
        if not wanted or not self.losers:
            return 0.0
        return sum(
            1 for picks in self.losers.values()
            if any(team_code(t) == wanted for t in picks)
        ) / len(self.losers)


def chalk(week: WeekPicks) -> dict[str, float]:
    """How much each coach went with the crowd, 0..1, over one week.

    A coach's figure is the average share of the pool that was on the same
    side as them, game by game. Around 0.8 is this pool's habit; well under
    it is a coach who takes unpopular sides, which is the only way anyone
    finishes a week clear of the field.
    """
    totals: dict[str, list[float]] = {}
    for game in week.games:
        # Each pick is resolved to a team once, and the sides counted once.
        # Asking the game for `share_on` per coach re-counted every card for
        # every coach - thirty-five times the work, times sixteen games,
        # times a team lookup on each - which is what made the pages that
        # read this crawl.
        codes = {coach: team_code(team) for coach, team in game.picks.items()}
        tally = Counter(code for code in codes.values() if code)
        counted = sum(tally.values())
        if not counted:
            continue
        for coach, code in codes.items():
            if code:
                totals.setdefault(coach, []).append(tally[code] / counted)
    return {
        coach: sum(shares) / len(shares)
        for coach, shares in totals.items() if shares
    }


def load_all(*, every_season: bool = False) -> dict[int, WeekPicks]:
    """Every stored week, reading the file once.

    Sheets found to belong to another season are left out unless
    `every_season` is set - only the Data page, which lists them so they can
    be cleared, has any use for them.
    """
    out: dict[int, WeekPicks] = {}
    for key in _read_store():
        try:
            number = int(key)
        except (TypeError, ValueError):
            continue
        sheet = load(number, every_season=every_season)
        if sheet is not None:
            out[number] = sheet
    return out


# ---- which season a stored sheet belongs to ------------------------------
# A pick sheet says "week 3" and nothing about the year. The store keeps
# every sheet it has been given, so last season's week 3 sits under the same
# key as this season's - and it was being read as this one: its crowd was
# quoted as this week's crowd, its coaches ranked among this pool's, and
# weeks 4, 8, 10, 11 and 13 were built in September out of last autumn.
#
# The sheet's own games are the tell. Where the workbook has that week's
# slate, a sheet from this season shares nearly all of its matchups and one
# from another season shares almost none. Where the workbook has no slate
# yet, only the week coming up can have picks already; a sheet further ahead
# than that is from a season already played.

#: Share of a sheet's matchups that must be on this season's slate that week.
SAME_SEASON = 0.6
_OTHER_SEASON: set[int] = set()


def _pairs(games) -> set[frozenset]:
    out = set()
    for game in games:
        pair = frozenset((team_code(game.home), team_code(game.away)))
        if len(pair) == 2 and "" not in pair:
            out.add(pair)
    return out


def _season_slate(season: Season, number: int) -> set[frozenset]:
    """This season's games in week `number`, from whatever knows them: the
    workbook's week, else the NFL's own fixtures as the scoreboard listed
    them (results.py), else the schedule the app carries for its season."""
    week = season.weeks.get(number)
    if week is not None and getattr(week, "source", None) is not None and week.games:
        return _pairs(week.games)
    from . import results, schedule
    year = results.season_year(season)
    listed = {frozenset((s.away, s.home)) for s in results.load(year, number)}
    if listed:
        return listed
    if getattr(schedule, "SCHEDULE_YEAR", None) == year:
        return {frozenset((team_code(a), team_code(h))) for a, h in schedule.pairings(number)
                if team_code(a) and team_code(h)}
    return set()


def from_other_season(season: Season, number: int, sheet: WeekPicks) -> bool:
    """Whether a stored sheet is another season's week `number`."""
    theirs = _pairs(sheet.games)
    if not theirs:
        return False
    ours = _season_slate(season, number)
    if ours:
        return len(theirs & ours) < SAME_SEASON * len(theirs)
    owned = {
        n: w for n, w in season.weeks.items()
        if getattr(w, "source", None) is not None and w.games
    }
    if not owned:
        return False            # nothing of this season's to compare with yet
    return number > max(owned) + 1


def other_season_weeks() -> set[int]:
    """Weeks whose stored sheet was last found to be from another season."""
    return set(_OTHER_SEASON)


def forget_weeks(numbers) -> None:
    """Remove stored sheets for good - used to clear out another season's."""
    from . import store as db
    store = dict(_read_store())
    changed = False
    for number in numbers:
        if store.pop(str(number), None) is not None:
            changed = True
        try:
            if db.get(GRADED_KIND, str(number), None) is not None:
                db.delete(GRADED_KIND, str(number))
        except db.StoreError:
            pass
    if changed:
        _write_store(store)
    _OTHER_SEASON.difference_update(numbers)


def read_file(path: str | Path) -> WeekPicks:
    """Parse a pick sheet. Raises Unreadable if it isn't one."""
    rows = _grid(path)
    if not rows:
        raise Unreadable(f"{Path(path).name} has nothing in it")

    found = TITLE.search(" ".join(rows[0]))
    if not found:
        raise Unreadable(
            f"{Path(path).name} doesn't say which week it is - the top row of a "
            f"pick sheet reads something like \"ACQL WEEK 2\""
        )
    week = int(found.group(1))

    width = max(len(row) for row in rows)
    losers: dict[int, int] = {}          # column -> which of the three
    suicide_col = 0
    games: list[tuple[int, str, str, str]] = []
    for col in range(1, width + 1):
        header = _cell(rows, HOME_ROW, col)
        if LOSER_HEADER.search(header):
            losers[col] = int(LOSER_HEADER.search(header).group(1))
            continue
        if SUICIDE_HEADER.search(header):
            suicide_col = col
            continue
        away = _cell(rows, AWAY_ROW, col)
        if away and header and team_code(away) and team_code(header):
            games.append((col, away, header, _cell(rows, NOTE_ROW, col)))
    if not games:
        raise Unreadable(
            f"{Path(path).name} has no games in it - a pick sheet has the road "
            f"team on row {AWAY_ROW} and the home team on row {HOME_ROW}"
        )

    picks: dict[int, dict[str, str]] = {col: {} for col, *_ in games}
    loser_picks: dict[str, list[str]] = {}
    suicide: dict[str, str] = {}
    scores: dict[str, list[int]] = {}
    name_col = _name_column(rows, min(col for col, *_ in games))
    graded = name_col > 1 and _scores_to_the_left(rows, name_col)
    coaches: list[str] = []
    for number in range(FIRST_PLAYER_ROW, len(rows) + 1):
        name = _cell(rows, number, name_col)
        if not name or SUMMARY_ROW.match(name) or NUMBER.match(name):
            continue
        # A coach is a row with picks on it. Separators, notes and whatever
        # else is typed down the left-hand column are not coaches, and
        # counting them makes every share in the pool come out too small.
        # On a graded sheet a coach who got everything wrong has no picks
        # left, but still has a score beside the name.
        chosen = {col: _cell(rows, number, col) for col, *_ in games}
        scored = graded and NUMBER.match(_cell(rows, number, 1) or "x")
        if not any(team_code(team) for team in chosen.values()) and not scored:
            continue
        coaches.append(name)
        if graded:
            # total first, then regular wins, when there is room for both
            got = [_whole(_cell(rows, number, col)) for col in range(1, min(name_col, 3))]
            if got and got[0] is not None:
                scores[name] = [got[0], got[1] if len(got) > 1 and got[1] is not None else got[0]]
        for col, team in chosen.items():
            if team:
                picks[col][name] = team
        losing = [_cell(rows, number, col) for col in sorted(losers, key=losers.get)]
        if any(losing):
            loser_picks[name] = losing
        if suicide_col:
            pick = _cell(rows, number, suicide_col)
            if pick:
                suicide[name] = pick

    credited: dict[int, str] = {}
    if graded:
        for col, away, home, _note in games:
            credited[col] = _ungrade(picks[col], away, home, coaches)
    return WeekPicks(
        week=week,
        games=[GamePicks(away, home, note, picks[col], credited.get(col, ""))
               for col, away, home, note in games],
        losers=loser_picks,
        suicide=suicide,
        source=str(path),
        graded=graded,
        scores=scores,
    )


def _whole(text: str) -> int | None:
    try:
        return int(round(float(str(text).strip())))
    except (TypeError, ValueError):
        return None


def _name_column(rows: list[list[str]], first_game: int) -> int:
    """Which column holds the coaches' names: the one, left of the games,
    with the most plain words in the coaches' rows."""
    best, best_count = 1, -1
    for col in range(1, max(2, first_game)):
        count = sum(
            1 for number in range(FIRST_PLAYER_ROW, len(rows) + 1)
            for value in (_cell(rows, number, col),)
            if value and not NUMBER.match(value) and not SUMMARY_ROW.match(value)
            and not team_code(value)
        )
        if count > best_count:
            best, best_count = col, count
    return best


def _scores_to_the_left(rows: list[list[str]], name_col: int) -> bool:
    """A graded sheet keeps each coach's score in the columns before the name."""
    numbers = total = 0
    for number in range(FIRST_PLAYER_ROW, len(rows) + 1):
        if not _cell(rows, number, name_col):
            continue
        total += 1
        numbers += bool(NUMBER.match(_cell(rows, number, 1) or "x"))
    return total > 0 and numbers >= 0.8 * total


def _ungrade(picks: dict[str, str], away: str, home: str, coaches: list[str]) -> str:
    """Put back the wrong picks a graded sheet deleted, for one game, and
    return the side the commissioner credited (as the sheet writes it).

    Only the side the pool credited is left standing, so every coach with
    no pick here had the other side. A game still showing both sides has not
    been graded (it hadn't been played) and is left as it is; so is one with
    no picks left at all, which can't say which side that was.
    """
    shown = Counter(team_code(team) for team in picks.values() if team_code(team))
    if len(shown) != 1:
        return ""
    credited = next(iter(shown))
    other = away if team_code(home) == credited else home if team_code(away) == credited else ""
    if not other:
        return ""
    for name in coaches:
        if not team_code(picks.get(name, "")):
            picks[name] = other
    return home if other == away else away


def summary(week: WeekPicks) -> str:
    """One line about a sheet, for the page to show."""
    alive = sum(1 for pick in week.suicide.values() if team_code(pick))
    parts = [_plural(len(week.coaches), "coach", "coaches")]
    if week.losers:
        parts.append(f"{len(week.losers)} big-loser cards")
    if week.suicide:
        parts.append(f"{alive} alive in the suicide pool")
    return " \u00b7 ".join(parts)


def _plural(count: int, word: str, many: str = "") -> str:
    return f"{count} {word if count == 1 else (many or word + 's')}"


# ---- grading a week from the raw sheet --------------------------------------
@dataclass(frozen=True)
class Graded:
    """What attaching the pick sheets did to a season."""

    built: list[int] = field(default_factory=list)       # weeks created from a sheet
    strangers: list[int] = field(default_factory=list)   # sheets from another pool
    other_season: list[int] = field(default_factory=list)  # sheets from another year
    filled: list[int] = field(default_factory=list)      # weeks whose picks were filled in
    checked: list[int] = field(default_factory=list)     # weeks compared against the sheet
    disagreements: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)    # line notes that couldn't be read

    @property
    def touched(self) -> list[int]:
        return sorted(set(self.built) | set(self.filled) | set(self.checked))


def _decided(game: Game) -> bool:
    return bool(str(getattr(game, "winner", "") or "").strip())


def grade_week(week: Week, sheet: WeekPicks) -> dict[str, int]:
    """Each coach's correct picks for a week, from the raw sheet and the results.

    The workbook records the winner with the pool's line already applied, so
    grading against it needs no knowledge of the line - a pick is right when
    it names the team in the result row.
    """
    scores: dict[str, int] = {}
    for game in week.games:
        if not _decided(game):
            continue
        found = sheet.game_for(game.away, game.home)
        if found is None:
            continue
        winner = team_code(game.winner)
        for coach, team in found.picks.items():
            scores[coach] = scores.get(coach, 0) + (team_code(team) == winner)
    return scores


def _sync_players(season: Season, week: Week, source: str) -> None:
    """Give the season a coach for every name on a week it took from a sheet.

    Scoring a week fills `week.lines`, which is what the weekly pages read.
    Everything else in the dashboard - standings, form, head to head, the
    crowd chart - reads `season.players`, so without this a season loaded
    from pick sheets alone knows all 35 cards and none of the 35 coaches.

    Only weeks the sheet built or filled come through here, so nothing the
    workbook entered is ever written over.

    A week with no results yet is deliberately left out of everyone's record.
    Grading an unplayed week gives all 35 coaches nought, and nought is not a
    score - carried into the season it would drag every average down and put
    the whole pool level at the bottom of the standings.
    """
    scored = any(_decided(game) for game in week.games)
    tag = f"{SHEET_SOURCE}{source or 'pick sheet'}"
    for key, line in week.lines.items():
        person = season.players.get(key)
        if person is None:
            person = Player(key=key, display=line.player or key)
            season.players[key] = person
        person.sources.add(tag)
        if not scored:
            continue
        # stats.xls's own score for the week stays: it is the pool's
        # official number, and a difference is reported, not written over.
        if week.number in getattr(person, "official_weeks", ()):
            pass
        elif week.number in getattr(season, "provisional_weeks", ()):
            # a week the season holds out as in progress stays held out
            person.provisional_wins[week.number] = line.total_wins
        else:
            person.weekly_wins[week.number] = line.total_wins
            person.weekly_regular[week.number] = line.regular_wins
            if line.big_loser_wins:
                person.weekly_big_loser[week.number] = line.big_loser_wins
            else:
                person.weekly_big_loser.pop(week.number, None)
        if line.suicide_pick:
            person.suicide_picks[week.number] = line.suicide_pick


def sheet_only(season: Season) -> set[str]:
    """Coaches who exist because a pick sheet was loaded, and nowhere else.

    Ownership has to be read off the player rather than remembered, because
    sheets are attached again every time one is added or a result arrives.
    A coach the workbook has touched carries that file's name among their
    sources, and from then on their record is the workbook's to keep.
    """
    return {
        key for key, person in season.players.items()
        if person.sources and all(s.startswith(SHEET_SOURCE) for s in person.sources)
    }


def _rank_players(season: Season, ours: set[str]) -> None:
    """Standings for coaches the workbook has never seen.

    `wins`, `pos` and the rest are normally read straight out of stats.xls,
    which is authoritative and stays that way: a coach the workbook knows is
    left exactly as the workbook left them. A coach who exists only because
    a pick sheet was loaded has no such row, so their record is added up from
    the weeks that are scored - and positions are only handed out when the
    workbook has given out none, so the two numbering systems never mix.
    """
    if not ours:
        return
    games = {w: len(wk.games) for w, wk in season.weeks.items()}
    for key in ours:
        person = season.players.get(key)
        if person is None:
            continue
        played = sum(games.get(w, 0) for w in person.weekly_wins)
        person.wins = person.season_wins
        person.losses = max(0, played - person.wins)
        person.pct = person.wins / played if played else 0.0
    if any(p.pos for p in season.players.values() if p.key not in ours):
        return            # the workbook is running the standings; leave them be
    order = sorted(season.players.values(), key=lambda p: (-p.wins, p.display))
    for place, person in enumerate(order, start=1):
        person.pos = place


#: How much of a sheet has to be this pool before the sheet is treated as
#: this pool's. A week of the pool's own sheet matches nearly every name; a
#: different pool's matches almost none, so anything in between is safely
#: called ours - a few coaches sitting a week out must not disown the file.
POOL_SHARE = 0.4

#: A roster smaller than this cannot tell one pool from another, so it does
#: not get a vote. Disowning a real sheet is far more costly than taking a
#: stray one: a stranger shows up as a name to ignore, while a wrongly
#: rejected sheet takes a whole week of picks with it, silently.
MIN_ROSTER = 5


def belongs_here(season: Season, sheet: WeekPicks) -> bool:
    """Whether this sheet is about this pool's coaches.

    Two ways a sheet counts as ours: most of the names on it are coaches here,
    or it carries a good share of the roster. The second matters when the
    workbook has only been half read - a handful of known coaches against a
    full 35-name sheet would otherwise look like a stranger's file.

    With no roster yet, or barely one, there is nothing to compare against and
    the sheet is taken at its word - that is the first-run case, where the
    sheets are all the app has.
    """
    roster = {_plain(person.display) for person in season.players.values()}
    if len(roster) < MIN_ROSTER:
        return True
    names = [_plain(name) for name in sheet.coaches]
    if not names:
        return False
    known = sum(1 for name in names if name in roster)
    return (
        known >= max(2, POOL_SHARE * len(names))
        or known >= POOL_SHARE * len(roster)
    )


def _plain(name: str) -> str:
    """A coach's name reduced to what cannot vary between files.

    Case, spacing and punctuation are exactly what differs between a pick
    sheet and the workbook - "Jim T" against "JimT" - so none of them decide
    who somebody is.
    """
    return "".join(c for c in str(name).casefold() if c.isalnum())


def _coach_key(season: Season, key_for=None):
    """How to file a coach from a pick sheet so they land on their own row.

    The workbook keys coaches through the alias table, and the pool spells
    names differently between the two files - "MayeDay! MayeDay!" on the pick
    sheet against whatever the workbook has. Filing a sheet under a key of its
    own would give that coach two rows: half their weeks under each, and every
    pool-wide average taken over a pool twice its real size.

    So an existing coach is matched first, by name with case, spacing and
    punctuation set aside. Only a coach nobody has seen before gets a new key,
    and that key comes from the same alias table the workbook uses whenever
    one is to hand.

    Once the pool has a roster, a name that matches nobody on it is not one of
    the pool's coaches and comes back empty. Pick sheets are downloaded a file
    at a time and they are not all this pool's: a folder of them brought in
    two dozen strangers - Blitzburgh, Frosted Flaccos, Mac Daddy Jones - each
    filed as a coach on no wins, and thirty-five became fifty-nine. The
    workbook says who is in the pool; a sheet only says what they picked.
    """
    known = {_plain(person.display): key for key, person in season.players.items()}
    roster = bool(known)

    def resolve(coach: str) -> str:
        plain = _plain(coach)
        if plain in known:
            return known[plain]
        if roster:
            return ""                # not one of ours
        if key_for is not None:
            try:
                made = key_for(coach)
            except Exception:                    # noqa: BLE001 - never fail a load
                made = ""
            if made:
                known[plain] = made
                return made
        known[plain] = plain
        return plain

    return resolve


#: What the commissioner writes in the suicide column once a coach is out.
OUT_MARKERS = {"dead", "out", "x", "eliminated", "knocked out"}


def _same_side(a: str, b: str) -> bool:
    """Two results naming the same side ("Tie" only matches "Tie")."""
    ca, cb = team_code(a), team_code(b)
    if ca or cb:
        return ca == cb
    return " ".join(str(a).split()).casefold() == " ".join(str(b).split()).casefold()


def _stored_results(season: Season):
    """The scores kept for each week, by week number (nothing is fetched)."""
    from . import results
    year = results.season_year(season)
    return lambda number: results.load(year, number)


def _shared(a: WeekPicks, b: WeekPicks) -> float:
    """How much of sheet a's pool is on sheet b: the share of a's coaches b names."""
    ours = {_plain(name) for name in a.coaches}
    theirs = {_plain(name) for name in b.coaches}
    return len(ours & theirs) / len(ours) if ours else 0.0


def _graded_copy(season: Season, number: int, card: WeekPicks, given: WeekPicks | None) -> WeekPicks | None:
    """The commissioner's graded copy of this week of this season, if any -
    never another season's, and never another pool's grading of the same
    NFL week (which knows nothing of this pool's lines)."""
    graded = given if given is not None and given.graded else load_graded(number)
    if graded is None:
        return None
    ours, theirs = _pairs(card.games), _pairs(graded.games)
    if not theirs or len(ours & theirs) < SAME_SEASON * max(len(theirs), 1):
        return None             # another season's week
    if _shared(card, graded) < POOL_SHARE:
        return None             # another pool's
    if from_other_season(season, number, graded):
        return None
    return graded


def _unanimous(*sheet_games) -> bool:
    """Whether every coach took the same side of a game, on the first of
    these sheets that shows any picks for it. On a graded copy such a game
    looks the same graded or not - everyone's pick still standing - so the
    copy can't say whether it has been played."""
    for found in sheet_games:
        if found is None:
            continue
        sides = {team_code(team) for team in found.picks.values() if team_code(team)}
        if sides:
            return len(sides) == 1
    return False


def _attach_lines(week: Week, sheet: WeekPicks, out: "Graded") -> None:
    """The pool's line on every game the slate has none for, from the note
    the sheet carries above it. A note that can't be read is said, not guessed."""
    from . import lines
    for game in week.games:
        if game.line is not None:
            continue
        found = sheet.game_for(game.away, game.home)
        if found is None:
            continue
        line, problem = lines.line_from_note(found.note, game.away, game.home)
        if line is not None:
            game.line_favourite, game.line = line
        elif problem:
            out.problems.append(
                f"Week {week.number}, {game.label}: {problem} - the game is graded "
                f"without a line until the sheet says it plainly"
            )


def _apply_results(week: Week, card: WeekPicks, graded: WeekPicks | None, scores: list,
                   out: "Graded", *, owned: bool) -> None:
    """Results onto the week's games: the commissioner's grading first, then
    the final score with the pool's line. A week the workbook owns keeps its
    own result row; the other two only check it."""
    from . import results
    for game in week.games:
        official = ""
        if graded is not None:
            found = graded.game_for(game.away, game.home)
            if (found is not None and found.credited
                    and not _unanimous(card.game_for(game.away, game.home), found)):
                official = game.home if team_code(found.credited) == team_code(game.home) else game.away
        score, _ = results.find(scores, game.away, game.home)
        final = None
        if score is not None and score.final:
            final = results.points(game, score)
            if final is not None:
                game.home_points, game.away_points = final
        computed = results.credited(game, *final) if final is not None else ""
        label = f"week {week.number} game {game.index} ({game.label})"
        if owned:
            if not _decided(game):
                continue
            game.result_from = game.result_from or "workbook"
            if official and not _same_side(official, game.winner):
                out.disagreements.append(
                    f"{label}: the result row says {game.winner}, but the graded "
                    f"sheet credits {official} - one of the two is wrong")
            elif computed and not official and not _same_side(computed, game.winner):
                out.disagreements.append(
                    f"{label}: the result row says {game.winner}, but the final score "
                    f"({_score_text(game)}) with the pool's line gives {computed}")
            continue
        if official:
            game.winner, game.result_from = official, "graded sheet"
            if computed and not _same_side(computed, official):
                out.disagreements.append(
                    f"{label}: the graded sheet credits {official}, but the final score "
                    f"({_score_text(game)}) with the pool's line gives {computed} - "
                    f"the graded sheet was kept")
        elif computed:
            game.winner, game.result_from = computed, "ESPN"
        elif game.result_from in ("ESPN", "graded sheet"):
            game.winner, game.result_from = "", ""     # no longer known: say so


def _score_text(game: Game) -> str:
    from .predictions import display_team
    return (f"{display_team(game.away)} {game.away_points}, "
            f"{display_team(game.home)} {game.home_points}")


def _mark(cell: object) -> str:
    """A Big Loser cell: the team, or "1" once it came in (however a
    spreadsheet wrote the 1 - "1.0" from an old .xls stored before)."""
    text = str(cell or "").strip()
    return "1" if NUMBER.match(text) and float(text) == 1 else text


def _loser_slots(week: Week, card: list[str], marks: list[str] | None,
                 listed: set[str], marked_games: set[str],
                 ) -> tuple[list[str], list[str], int, list[str]]:
    """(the three cells as the workbook would hold them, the teams, the
    points, the teams counted that the graded copy left unmarked).

    A cell holds "1" once its team has come in, as the commissioner marks
    it. A "1" on his graded copy counts. Otherwise the final score decides
    (beaten by 15 or more), and failing a score, the standings export's
    list of the week's Big Losers. A graded copy can be a snapshot taken
    part-way through grading, so a team the score says came in counts even
    where the copy hasn't marked it yet - and is named, so it can be checked.
    """
    from . import results
    cells, teams, hits, unmarked = [], [], 0, []
    for slot in range(3):
        name = _mark(card[slot]) if slot < len(card) and card[slot] else ""
        mark = _mark(marks[slot]) if marks and slot < len(marks) and marks[slot] else ""
        hit = name == "1" or mark == "1"
        if team_code(name):
            teams.append(name)
            if not hit:
                game = next((g for g in week.games
                             if team_code(name) in (team_code(g.home), team_code(g.away))), None)
                by = results.lost_by(game, name) if game is not None else None
                if by is not None:
                    hit = _decided(game) and by >= results.BLOWOUT
                    if hit and marks is not None and team_code(name) in marked_games:
                        unmarked.append(team_code(name))
                elif team_code(name) in listed:
                    hit = True
        if hit:
            hits += 1
        cells.append("1" if hit else name)
    return cells, teams, hits, unmarked


def attach(
    season: Season,
    sheets: dict[int, WeekPicks] | None = None,
    *,
    key_for=None,
    results_for=None,
) -> Graded:
    """Put the pool's pick sheets into a season: every week the workbook
    hasn't scored is built and graded from them.

    * A week the workbook has nothing for - or has, but with no results in
      it - is built from the sheet: the slate, the pool's lines from the
      notes above the games, everyone's picks, Big Losers and suicide picks,
      graded against the results. That is the week nobody types in by hand.
    * The results come from the commissioner's graded copy of the sheet when
      it has been downloaded, and otherwise from the final scores the
      dashboard fetched (results.py), with the pool's line applied.
    * A week the workbook has scored is left exactly as it is, and checked:
      against the graded copy, and against the final scores. Anywhere they
      disagree is reported rather than corrected.

    A sheet that is not this pool's, or not this season's, is left alone.
    `results_for(week)` gives a week's fetched scores; by default the ones
    already kept (nothing is fetched here).
    """
    out = Graded()
    key_of = _coach_key(season, key_for)
    everything = sheets or load_all(every_season=True)
    if sheets is None:
        _OTHER_SEASON.clear()
    scores_for = results_for or _stored_results(season)
    for number, sheet in sorted(everything.items()):
        if from_other_season(season, number, sheet):
            out.other_season.append(number)
            if sheets is None:
                _OTHER_SEASON.add(number)
            out.notes.append(
                f"week {number}: {Path(sheet.source).name or 'the stored sheet'} "
                f"is from another season - its games are not this season's "
                f"week {number} - so it was left out"
            )
            continue
        if sheet.coaches and not belongs_here(season, sheet):
            out.notes.append(
                f"week {number}: {Path(sheet.source).name or 'that sheet'} is "
                f"another pool's - hardly any of its names are in this one, so "
                f"it was left out"
            )
            out.strangers.append(number)
            continue
        # The card is the sheet as it went out; a graded copy handed in on
        # its own is the card only when no other sheet for the week is kept.
        card = sheet
        if sheets is not None and sheet.graded:
            kept = load(number)
            if kept is not None and not kept.graded and _pairs(kept.games) & _pairs(sheet.games):
                card = kept
        graded = _graded_copy(season, number, card, sheet)

        week = season.weeks.get(number)
        fresh = week is None
        if fresh:
            week = Week(number=number)
            season.weeks[number] = week
        # The workbook owns a week once it has a result in it. One with only
        # the matchups typed in is still the sheet's to score.
        owned = getattr(week, "source", None) is not None and any(_decided(g) for g in week.games)
        if not week.games:
            week.games = [
                Game(index=index, home=game.home, away=game.away)
                for index, game in enumerate(card.games, start=1)
            ]
            out.notes.append(
                f"week {number}: the slate came from {Path(card.source).name or 'a pick sheet'}"
            )
        _attach_lines(week, card, out)
        try:
            scores = list(scores_for(number) or [])
        except Exception:  # noqa: BLE001 - a week with no results is still a week
            scores = []
        _apply_results(week, card, graded, scores, out, owned=owned)
        if owned:
            out.checked.append(number)
            continue
        if not card.coaches:
            # A blank sheet: the games, the line games and the results, and
            # nobody's picks yet - whatever the week already has stays.
            if not week.lines:
                season.fixture_weeks.append(number)
            (out.built if fresh or getattr(week, "source", None) is None else out.filled).append(number)
            continue

        listed = {team_code(t) for t in season.big_losers_by_week.get(number, []) if team_code(t)}
        # The games the graded copy has graded, and whether that is all of them
        # (a copy downloaded before Monday night hasn't the late game in it).
        marked_games: set[str] = set()
        unanimous = False
        if graded is not None:
            for found in graded.games:
                if _unanimous(card.game_for(found.away, found.home), found):
                    unanimous = True        # can't tell whether the copy graded it
                elif found.credited:
                    marked_games |= {team_code(found.home), team_code(found.away)}
        complete = (graded is not None and not unanimous
                    and all(found.credited for found in graded.games))
        week.lines.clear()
        outsiders = 0
        differ = []
        overruled: Counter = Counter()
        for coach in card.coaches:
            key = key_of(coach)
            if not key:
                # On this sheet but not in this pool. Their picks still
                # count towards what the crowd did - that is read from the
                # sheet itself - but they are not given a coach's row.
                outsiders += 1
                continue
            line = PlayerWeek(player=coach, week=number)
            for game in week.games:
                found = card.game_for(game.away, game.home)
                team = found.picks.get(coach, "") if found is not None else ""
                if team and _decided(game) and team_code(team) == team_code(game.winner):
                    line.correct_picks[game.index] = team
            marks = graded.losers.get(coach) if graded is not None else None
            cells, teams, hits, unmarked = _loser_slots(
                week, card.losers.get(coach, []), marks, listed, marked_games)
            line.big_loser_raw, line.big_loser_picks, line.big_loser_wins = cells, teams, hits
            overruled.update(unmarked)
            line.regular_wins = len(line.correct_picks)
            line.total_wins = line.regular_wins + hits
            pick = str(card.suicide.get(coach, "") or "").strip()
            if pick.casefold() in OUT_MARKERS:
                line.suicide_out = True
            else:
                line.suicide_pick = pick
            if complete and coach in graded.scores:
                total, regular = graded.scores[coach]
                if (total, regular) != (line.total_wins, line.regular_wins):
                    differ.append(f"{coach} {total} on the sheet, {line.total_wins} here")
            week.lines[key] = line
        for code, cards in sorted(overruled.items()):
            game = next((g for g in week.games if code in (team_code(g.home), team_code(g.away))), None)
            from . import results
            by = results.lost_by(game, code) if game is not None else None
            out.disagreements.append(
                f"week {number}: {DISPLAY.get(code, code)} lost by {by} - a Big Loser by the "
                f"pool's rule - but the graded sheet doesn't mark it on "
                f"{cards} card{'' if cards == 1 else 's'}. Counted here; the standings "
                f"export has the final word"
            )
        if differ:
            out.disagreements.append(
                f"week {number}: the graded sheet's totals differ from the dashboard's "
                f"for {len(differ)} coach{'' if len(differ) == 1 else 'es'} ("
                + "; ".join(differ[:3]) + (" ..." if len(differ) > 3 else "") + ")"
            )
        _sync_players(season, week, Path(card.source).name)
        if outsiders:
            one = outsiders == 1
            out.notes.append(
                f"week {number}: {outsiders} name{'' if one else 's'} on "
                f"{Path(card.source).name or 'the sheet'} "
                f"{'is' if one else 'are'} not in this pool and "
                f"{'was' if one else 'were'} left out of the standings"
            )
        (out.built if fresh or getattr(week, "source", None) is None else out.filled).append(number)
    _rank_players(season, sheet_only(season))
    return out


# ---- keeping them between runs ---------------------------------------------
# Kept in the app's database (store.py), which caches what it has read until
# something is written - pages ask for the pick sheets on every refresh, and
# re-parsing them each time was measurable.

def _store_stamp() -> tuple:
    from . import store
    return store.stamp("picks")


def _read_store() -> dict:
    from . import store
    return store.items("picks")


def forget_store() -> None:
    """Kept for older tests; the store's own cache follows every write."""


def _write_store(mapping: dict) -> None:
    from . import store
    try:
        store.replace("picks", mapping)
    except store.StoreError:
        pass


#: The graded copy of each week, kept beside the sheet as it went out.
GRADED_KIND = "graded"
#: Which pick sheet files have been read, and as of when (path -> mtime).
FILES_KIND = "picks-files"
#: Stored sheets carry the side each graded game was credited to from this
#: format on; an older stored graded sheet is read again from its file.
FORMAT = 2


def _entry(week: WeekPicks) -> dict:
    return {
        "source": week.source,
        "imported": time.time(),
        "games": [asdict(game) for game in week.games],
        "losers": week.losers,
        "suicide": week.suicide,
        "graded": week.graded,
        "scores": week.scores,
        "v": FORMAT,
    }


def _same_slate(entry: object, week: WeekPicks) -> bool:
    """Whether a stored sheet is the same week of the same season as this one."""
    if not isinstance(entry, dict):
        return False
    theirs = _pairs(
        GamePicks(away=str(row.get("away", "")), home=str(row.get("home", "")), note="")
        for row in entry.get("games") or [] if isinstance(row, dict)
    )
    ours = _pairs(week.games)
    return bool(theirs) and len(theirs & ours) >= SAME_SEASON * max(len(ours), 1)


def _roster(besides: int) -> set[str]:
    """Every coach named on the kept cards of the other weeks - the pool, as
    far as its sheets say."""
    names: set[str] = set()
    for key, entry in _read_store().items():
        if key == str(besides) or not isinstance(entry, dict):
            continue
        for row in entry.get("games") or []:
            if isinstance(row, dict):
                names |= {_plain(n) for n in (row.get("picks") or {})}
    return names


def _ours(sheet: WeekPicks, besides: int) -> float | None:
    """How much of a sheet is this pool, by the other weeks' cards (None: too
    few kept to say)."""
    roster = _roster(besides)
    names = {_plain(n) for n in sheet.coaches}
    if len(roster) < MIN_ROSTER or not names:
        return None
    return len(names & roster) / len(names)


def _has_picks(entry: object) -> bool:
    """Whether a kept sheet has anybody's picks on it."""
    return isinstance(entry, dict) and any(
        isinstance(row, dict) and row.get("picks") for row in entry.get("games") or [])


def _keeps_place(held: object, week: WeekPicks) -> bool:
    """Whether a kept sheet for the same week should stay where it is rather
    than give way to this one: same season, and this one is another pool's
    (its coaches aren't the ones on the other weeks' cards) while it isn't."""
    if not isinstance(held, dict) or not _same_slate(held, week):
        return False
    new = _ours(week, week.week)
    old = _from_entry(week.week, held)
    old_share = _ours(old, week.week) if old is not None else None
    return new is not None and old_share is not None and new < POOL_SHARE <= old_share


def save(week: WeekPicks) -> None:
    """Keep a sheet. The sheet as it went out is the week's card; a graded
    copy is kept beside it rather than over it, because grading writes "1"
    over every Big Loser that came in and the team is gone. A graded copy is
    the card only for a week with no other sheet of this season's. Another
    pool's sheet for the same week - Downloads holds them too - never takes
    the place of this pool's."""
    from . import store
    entry = _entry(week)
    try:
        held = _read_store().get(str(week.week))
        if week.graded:
            card = (_from_entry(week.week, held)
                    if isinstance(held, dict) and not held.get("graded") and _same_slate(held, week)
                    else None)
            if card is not None and _shared(card, week) < POOL_SHARE:
                pass            # another pool's grading of the same NFL week
            elif not _keeps_place(store.get(GRADED_KIND, str(week.week), None), week):
                store.put(GRADED_KIND, str(week.week), entry)
            if card is None and not _keeps_place(held, week):
                store.put("picks", str(week.week), entry)
        elif not week.coaches and _has_picks(held) and _same_slate(held, week):
            pass                # the blank sheet came after the full one
        elif not _keeps_place(held, week):
            store.put("picks", str(week.week), entry)
    except store.StoreError:
        pass
    # A sheet saved now is this season's until the next load says otherwise.
    _OTHER_SEASON.discard(week.week)


def imported_at(week: int) -> float | None:
    """When the stored sheet for a week was read in (seconds since 1970)."""
    entry = _read_store().get(str(week))
    try:
        return float(entry["imported"]) if isinstance(entry, dict) and entry.get("imported") else None
    except (TypeError, ValueError):
        return None


def _from_entry(week: int, entry: object) -> WeekPicks | None:
    if not isinstance(entry, dict):
        return None
    games = []
    for row in entry.get("games", []):
        if isinstance(row, dict):
            games.append(GamePicks(
                away=str(row.get("away", "")), home=str(row.get("home", "")),
                note=str(row.get("note", "")), picks=dict(row.get("picks") or {}),
                credited=str(row.get("credited", "") or ""),
            ))
    if not games:
        return None
    scores = {}
    for name, pair in (entry.get("scores") or {}).items():
        if isinstance(pair, (list, tuple)) and pair and all(isinstance(v, int) for v in pair):
            scores[str(name)] = list(pair)
    return WeekPicks(
        week=week, games=games,
        losers={k: list(v) for k, v in (entry.get("losers") or {}).items()},
        suicide=dict(entry.get("suicide") or {}),
        source=str(entry.get("source", "")),
        graded=bool(entry.get("graded", False)),
        scores=scores,
    )


def load(week: int, *, every_season: bool = False) -> WeekPicks | None:
    """The week's card: the sheet as it went out, or the graded copy when
    that is all there is."""
    if week in _OTHER_SEASON and not every_season:
        return None
    return _from_entry(week, _read_store().get(str(week)))


def load_graded(week: int) -> WeekPicks | None:
    """The commissioner's graded copy of a week, if one has been read."""
    from . import store
    try:
        entry = store.get(GRADED_KIND, str(week), None)
    except store.StoreError:
        return None
    return _from_entry(week, entry)


def search_folders(workbook: str | Path | None = None, extra=None) -> list[Path]:
    """Where a downloaded pick sheet is likely to be: the usual places, the
    workbook's folder, and the folders the dashboard watches."""
    from .predictions import search_folders as folders
    out = folders(workbook)
    seen = set(out)
    for folder in extra or ():
        try:
            resolved = Path(folder).expanduser().resolve()
        except OSError:
            continue
        if resolved not in seen and resolved.is_dir():
            seen.add(resolved)
            out.append(resolved)
    return out


def find_files(workbook: str | Path | None = None, extra=None) -> list[Path]:
    """Every pick sheet in the usual folders, newest first.

    The file name carries a number that is not the week, so the week has to
    be read out of each one; they are all offered and the caller sorts it out.
    """
    found: list[tuple[float, Path]] = []
    for folder in search_folders(workbook, extra):
        for pattern in ("picks*.xls", "picks*.xlsx", "*picks*.xls", "*picks*.xlsx"):
            try:
                hits = list(folder.glob(pattern))
            except OSError:
                continue
            for path in hits:
                if path.name.startswith(("~$", ".")):
                    continue
                try:
                    found.append((path.stat().st_mtime, path))
                except OSError:
                    continue
    seen, out = set(), []
    for _, path in sorted(found, reverse=True):
        if path not in seen:
            seen.add(path)
            out.append(path)
    return out


def _files_read() -> dict[str, float]:
    from . import store
    try:
        raw = store.items(FILES_KIND)
    except store.StoreError:
        return {}
    return {k: float(v) for k, v in raw.items() if isinstance(v, (int, float))}


def _mark_read(path: Path, when: float) -> None:
    from . import store
    try:
        store.put(FILES_KIND, str(path), when)
    except store.StoreError:
        pass


def refresh_from_disk(workbook: str | Path | None = None, extra=None) -> list[int]:
    """Read any pick sheet not read before (or changed since). Returns the weeks read.

    Oldest first, so when two files are the same week - a second download,
    or the graded copy - the newest is the last word. Each file is read once:
    the next load only looks at what is new or has changed.
    """
    if extra is not None:
        from . import predictions
        predictions.WATCHED[:] = [Path(folder) for folder in extra]
    done = _files_read()
    added = []
    for path in reversed(find_files(workbook, extra)):
        try:
            when = path.stat().st_mtime
        except OSError:
            continue
        if done.get(str(path), -1.0) >= when:
            continue
        try:
            week = read_file(path)
        except Exception:  # noqa: BLE001 - one bad file must not stop the rest
            # Not a pick sheet, not a spreadsheet at all, or not yet a whole
            # one (still downloading): noted, so it isn't opened again on
            # every load until it changes - and the other files still are.
            _mark_read(path, when)
            continue
        save(week)
        _mark_read(path, when)
        added.append(week.week)
    return sorted(set(added))
